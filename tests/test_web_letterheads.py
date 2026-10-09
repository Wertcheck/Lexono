"""Mehrere Briefkoepfe: Verwaltung, Chat-Darstellung, Editor-Wechsel, Export."""

from __future__ import annotations

import io
from collections.abc import Iterator

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.session import get_db
from app.firm_profile import get_firm_profile
from app.firm_profile.letterheads import FIRM_LETTERHEAD_REF, create_letterhead
from app.main import app
from app.models import AuditEvent, ChatConversation, ChatMessage, Client, Draft, Matter
from app.models.base import Base
from tests.auth_test_utils import extract_csrf, login_as_admin


@pytest.fixture()
def db_session() -> Iterator[Session]:
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def client(db_session: Session) -> Iterator[TestClient]:
    def _override() -> Iterator[Session]:
        yield db_session

    app.dependency_overrides[get_db] = _override
    try:
        test_client = TestClient(app)
        login_as_admin(db_session, test_client)
        yield test_client
    finally:
        app.dependency_overrides.clear()


@pytest.fixture()
def two_letterheads(db_session: Session) -> dict:
    profile = get_firm_profile(db_session)
    profile.firm_name = "Kanzlei Allgemein GbR"
    profile.street = "Allgemeinweg 1"
    profile.postal_code = "10000"
    profile.city = "Alt-Stadt"
    profile.signatory_name = "RA Allgemein"
    db_session.commit()
    immo = create_letterhead(
        db_session, name="Kanzlei Immobilienrecht", actor="t", firm_name="Immobilien Kanzlei Nord",
        street="Immoweg 2", postal_code="20000", city="Neu-Stadt", signatory_name="RA Immobilien",
    )
    client_ = Client(name="Test Mandant")
    matter = Matter(client=client_, title="Akte")
    db_session.add_all([client_, matter])
    db_session.commit()
    return {"immo": immo, "matter": matter}


def _csrf(client: TestClient) -> str:
    return extract_csrf(client.get("/dashboard/settings/letterheads").text)


# --- Verwaltung -----------------------------------------------------------


def test_letterheads_page_lists_profile_letterhead_first_with_default_badge(
    client: TestClient, two_letterheads: dict
) -> None:
    page = client.get("/dashboard/settings/letterheads").text
    assert page.index("Kanzlei allgemein") < page.index("Kanzlei Immobilienrecht")
    assert page.count("Standard</span>") == 1
    assert "Allgemeinweg 1" in page and "Immoweg 2" in page


def test_create_edit_default_and_delete_through_the_page(client: TestClient, db_session: Session) -> None:
    token = _csrf(client)
    created = client.post(
        "/dashboard/settings/letterheads/create",
        data={"csrf_token": token, "name": "Arbeitsrecht", "firm_name": "Kanzlei Arbeit", "street": "Arbeitsweg 3"},
        follow_redirects=False,
    )
    assert created.status_code == 303
    from app.models import Letterhead

    row = db_session.query(Letterhead).filter_by(name="Arbeitsrecht").one()
    assert row.firm_name == "Kanzlei Arbeit"

    client.post(f"/dashboard/settings/letterheads/{row.id}/default", data={"csrf_token": token})
    page = client.get("/dashboard/settings/letterheads").text
    assert "Standard</span>" in page
    db_session.expire_all()
    assert get_firm_profile(db_session).default_letterhead_id == row.id

    client.post(
        f"/dashboard/settings/letterheads/{row.id}/update",
        data={"csrf_token": token, "name": "Arbeitsrecht neu", "firm_name": "Kanzlei Arbeit 2"},
    )
    db_session.expire_all()
    assert db_session.get(Letterhead, row.id).name == "Arbeitsrecht neu"

    client.post(f"/dashboard/settings/letterheads/{row.id}/delete", data={"csrf_token": token})
    assert db_session.get(Letterhead, row.id) is None
    db_session.expire_all()
    assert get_firm_profile(db_session).default_letterhead_id is None


def test_duplicate_name_is_rejected_with_a_clear_message(client: TestClient, two_letterheads: dict) -> None:
    token = _csrf(client)
    response = client.post(
        "/dashboard/settings/letterheads/create",
        data={"csrf_token": token, "name": "kanzlei immobilienrecht", "firm_name": "X"},
        follow_redirects=True,
    )
    assert "bereits einen Briefkopf" in response.text


# --- Chat -------------------------------------------------------------------


def _chat_with_draft(db_session: Session, matter: Matter, *, letterhead_ref: str | None, body: str = "Sehr geehrte Damen und Herren,") -> tuple[ChatConversation, Draft]:
    conversation = ChatConversation(matter_id=matter.id, title="C", owner_id=None) if False else None
    from app.models import User

    user = db_session.query(User).first()
    conversation = ChatConversation(matter_id=matter.id, title="C", user_id=user.id)
    db_session.add(conversation)
    draft = Draft(
        matter_id=matter.id, content=f"<p>{body}</p>", version=1, status="draft", content_format="html",
        letterhead_ref=letterhead_ref,
    )
    db_session.add(draft)
    db_session.commit()
    db_session.add(ChatMessage(conversation_id=conversation.id, role="assistant", content=body, draft_id=draft.id))
    db_session.commit()
    return conversation, draft


def test_chat_panel_shows_the_letterhead_of_the_draft_exactly_once_in_the_copy_target(
    client: TestClient, db_session: Session, two_letterheads: dict
) -> None:
    conversation, _draft = _chat_with_draft(db_session, two_letterheads["matter"], letterhead_ref=two_letterheads["immo"].id)
    page = client.get(f"/dashboard/chat/{conversation.id}").text
    assert page.count("Immobilien Kanzlei Nord") == 1
    assert "Kanzlei Allgemein GbR" not in page.split("chat-document-panel")[1]
    assert page.count("RA Immobilien") == 1
    assert 'class="chat-letterhead"' in page and 'class="chat-signatory"' in page


def test_old_draft_without_ref_keeps_showing_the_profile_letterhead(
    client: TestClient, db_session: Session, two_letterheads: dict
) -> None:
    conversation, _draft = _chat_with_draft(db_session, two_letterheads["matter"], letterhead_ref=None)
    page = client.get(f"/dashboard/chat/{conversation.id}").text
    assert "Kanzlei Allgemein GbR" in page and "Immobilien Kanzlei Nord" not in page.split("chat-document-panel")[1]


def test_composer_offers_the_letterhead_choice_only_with_more_than_one(
    client: TestClient, db_session: Session
) -> None:
    assert 'name="letterhead_ref"' not in client.get("/dashboard/chat").text
    profile = get_firm_profile(db_session)
    profile.firm_name = "Eine Kanzlei"
    db_session.commit()
    create_letterhead(db_session, name="Zweiter", actor="t", firm_name="Zweite Kanzlei")
    page = client.get("/dashboard/chat").text
    assert 'name="letterhead_ref"' in page
    assert "Kanzlei allgemein (Standard)" in page and "Zweiter" in page


# --- Editor-Wechsel + Export --------------------------------------------------


def test_editor_shows_the_drafts_letterhead_and_switch_creates_a_traceable_version(
    client: TestClient, db_session: Session, two_letterheads: dict
) -> None:
    _conv, draft = _chat_with_draft(db_session, two_letterheads["matter"], letterhead_ref=None)
    page = client.get(f"/dashboard/drafts/{draft.id}/edit")
    assert "Kanzlei Allgemein GbR" in page.text and "Immobilien Kanzlei Nord" not in page.text
    token = extract_csrf(page.text)

    switched = client.post(
        f"/dashboard/drafts/{draft.id}/letterhead",
        data={"csrf_token": token, "letterhead_ref": two_letterheads["immo"].id},
        follow_redirects=False,
    )

    assert switched.status_code == 303
    new = db_session.query(Draft).filter(Draft.previous_version_id == draft.id).one()
    assert new.letterhead_ref == two_letterheads["immo"].id and new.version == 2
    assert new.content == draft.content, "Text bleibt unveraendert"
    db_session.refresh(draft)
    assert draft.letterhead_ref is None, "Vorgaengerfassung behaelt ihren Briefkopf"
    event = db_session.query(AuditEvent).filter_by(entity_id=new.id, event_type="draft_letterhead_changed").one()
    assert "Kanzlei allgemein" in event.details and "Kanzlei Immobilienrecht" in event.details
    page2 = client.get(switched.headers["location"])
    assert "Immobilien Kanzlei Nord" in page2.text and "Kanzlei Allgemein GbR" not in page2.text.split("document-page")[1]
    # Alte Fassung weist auf die neuere hin.
    assert "Es gibt eine neuere Fassung" in client.get(f"/dashboard/drafts/{draft.id}/edit").text


def test_switch_keeps_autosaved_manual_edits_and_follow_up_versions_keep_the_letterhead(
    client: TestClient, db_session: Session, two_letterheads: dict
) -> None:
    from app.drafting.versioning import create_manual_edit_version

    _conv, draft = _chat_with_draft(db_session, two_letterheads["matter"], letterhead_ref=None)
    draft.content = "<p>Handschriftlich geaendert.</p>"
    db_session.commit()
    token = extract_csrf(client.get(f"/dashboard/drafts/{draft.id}/edit").text)
    client.post(f"/dashboard/drafts/{draft.id}/letterhead", data={"csrf_token": token, "letterhead_ref": two_letterheads["immo"].id})
    v2 = db_session.query(Draft).filter(Draft.previous_version_id == draft.id).one()
    assert "Handschriftlich geaendert" in v2.content
    v3 = create_manual_edit_version(db_session, previous_draft=v2, new_content="<p>Noch eine Aenderung.</p>", actor="t")
    assert v3.letterhead_ref == two_letterheads["immo"].id


def test_docx_export_uses_the_drafts_letterhead_not_the_default(
    client: TestClient, db_session: Session, two_letterheads: dict
) -> None:
    _conv, draft = _chat_with_draft(db_session, two_letterheads["matter"], letterhead_ref=two_letterheads["immo"].id)
    response = client.get(f"/dashboard/drafts/{draft.id}/export.docx")
    assert response.status_code == 200
    doc = DocxDocument(io.BytesIO(response.content))
    header_text = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "Immobilien Kanzlei Nord" in header_text and "Kanzlei Allgemein GbR" not in header_text
    assert header_text.count("Immobilien Kanzlei Nord") == 1
    body = "\n".join(p.text for p in doc.paragraphs)
    assert body.count("RA Immobilien") == 1


def test_switching_to_the_profile_letterhead_works_back_and_forth(
    client: TestClient, db_session: Session, two_letterheads: dict
) -> None:
    _conv, draft = _chat_with_draft(db_session, two_letterheads["matter"], letterhead_ref=two_letterheads["immo"].id)
    token = extract_csrf(client.get(f"/dashboard/drafts/{draft.id}/edit").text)
    client.post(f"/dashboard/drafts/{draft.id}/letterhead", data={"csrf_token": token, "letterhead_ref": FIRM_LETTERHEAD_REF})
    new = db_session.query(Draft).filter(Draft.previous_version_id == draft.id).one()
    assert new.letterhead_ref == FIRM_LETTERHEAD_REF
