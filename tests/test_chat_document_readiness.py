"""Dokumentstatus und Chat-Verarbeitung (13.10.2026, app/documents/readiness.py).

ECHTER FUND (installierter Build, OCR_ENABLED=false): ein hochgeladener Scan blieb "pending" ohne Text, die Chat-Frage
ging trotzdem an die Cloud und wurde beantwortet, als waere das Dokument ausgewertet worden. Diese Tests laufen ueber
den echten Pfad (Upload -> DocumentProcessingService -> Chat-Router -> ChatService -> DraftingService -> Privacy
Gateway); gefaelscht ist nur die Cloud-Grenze selbst (`FakeClaudeWritingProvider` zaehlt jeden Cloud-Aufruf und haelt
den tatsaechlich gesendeten Payload fest). Ausschliesslich synthetische Daten.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pymupdf
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

import app.web.chat_router as chat_router_module
from app.ai_providers.local_ai_provider import RuleBasedLocalAIProvider
from app.chat.service import ChatService
from app.config.settings import Settings
from app.documents import readiness
from app.documents.extraction import extract_text
from app.models import AuditEvent, ChatConversation, ChatMessage, Document
from tests.auth_test_utils import login_as_admin
from tests.test_web_chat import (  # noqa: F401 - Fixtures
    FakeClaudeWritingProvider,
    _csrf,
    _parse_sse_events,
    _working_drafting_service,
    client,
    db_session,
)

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import bench_realistic_documents as harness  # noqa: E402
import make_realistic_privacy_docs as generator  # noqa: E402

TESSERACT_EXE = harness.TESSERACT_DIR / "bin" / "tesseract.exe"
needs_tesseract = pytest.mark.skipif(not TESSERACT_EXE.is_file(), reason="Tesseract nicht vorhanden")

SIMPLE_TEXT = "Rechnung Nr. 4711 über Estricharbeiten im Erdgeschoss. Der Betrag von 1.200,00 EUR ist fällig."


# --- synthetische Dokumente -------------------------------------------------------------------------------------------


def _text_pdf(path: Path, text: str = SIMPLE_TEXT) -> Path:
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=595, height=842)
        page.insert_textbox(pymupdf.Rect(60, 60, 535, 400), text, fontsize=12)
        pdf.save(str(path))
    return path


def _scan_of(source: Path, target: Path) -> Path:
    generator.make_scan(source, target)
    return target


def _mixed_pdf(path: Path, tmp: Path) -> Path:
    """Seite 1 mit Textebene, Seite 2 nur als Bild (Scan) - typischer Teil-Scan."""
    scan = _scan_of(_text_pdf(tmp / "seite2.pdf", "Seite zwei: Abnahmeprotokoll mit festgestellten Mängeln."), tmp / "s2.pdf")
    with pymupdf.open() as out:
        out.insert_pdf(pymupdf.open(str(_text_pdf(tmp / "seite1.pdf"))))
        out.insert_pdf(pymupdf.open(str(scan)))
        out.save(str(path))
    return path


def _blank_scan(path: Path) -> Path:
    with pymupdf.open() as blank:
        blank.new_page(width=595, height=842)
        pix = blank[0].get_pixmap(dpi=100)
    with pymupdf.open() as out:
        page = out.new_page(width=595, height=842)
        page.insert_image(page.rect, pixmap=pix)
        out.save(str(path))
    return path


def _empty_docx(path: Path) -> Path:
    from docx import Document as DocxDocument

    DocxDocument().save(str(path))
    return path


# --- Hilfen fuer den echten Web-Pfad ----------------------------------------------------------------------------------


@pytest.fixture()
def writer(monkeypatch) -> FakeClaudeWritingProvider:
    fake = FakeClaudeWritingProvider("Antwort der Cloud-KI.")
    monkeypatch.setattr(chat_router_module, "get_drafting_service", lambda: _working_drafting_service(fake))
    return fake


def _use_settings(monkeypatch, tmp_path, *, ocr_enabled: bool) -> Settings:
    """Isolierte QA-Konfiguration NUR fuer diesen Test (nie die installierte .env)."""
    kwargs = {"chat_upload_storage_dir": str(tmp_path / "uploads"), "ocr_enabled": ocr_enabled}
    if ocr_enabled:
        import pytesseract

        kwargs["tesseract_cmd"] = str(TESSERACT_EXE)
        monkeypatch.setenv("TESSDATA_PREFIX", str(harness.TESSERACT_DIR / "tessdata"))
        # configure_tesseract setzt den Pfad prozessweit - nach dem Test wiederherstellen (sonst beeinflusst er andere
        # OCR-Tests ohne TESSDATA_PREFIX).
        monkeypatch.setattr(pytesseract.pytesseract, "tesseract_cmd", pytesseract.pytesseract.tesseract_cmd)
    settings = Settings(**kwargs)
    monkeypatch.setattr(chat_router_module, "get_settings", lambda: settings)
    monkeypatch.setattr("app.chat.service.get_settings", lambda: settings)
    return settings


def _send_stream(client: TestClient, content: str, files: list[Path] = (), conversation_id: str = "") -> list[dict]:
    handles = [open(path, "rb") for path in files]
    try:
        response = client.post(
            "/dashboard/chat/send-stream",
            data={"csrf_token": _csrf(client), "conversation_id": conversation_id, "content": content},
            files=[("documents", (path.name, handle, "application/octet-stream")) for path, handle in zip(files, handles)],
        )
    finally:
        for handle in handles:
            handle.close()
    assert response.status_code == 200
    return _parse_sse_events(response.text)


def _done(events: list[dict]) -> dict:
    done = [event for event in events if event["kind"] == "done"]
    assert len(done) == 1, events
    return done[0]


def _conversation_id(events: list[dict]) -> str:
    return next(event["conversation_id"] for event in events if event["kind"] == "start")


def _payload_text(writer: FakeClaudeWritingProvider) -> str:
    return "".join(payload.model_dump_json() for payload in writer.received_payloads)


# --- 1. Scan bei OCR_ENABLED=false --------------------------------------------------------------------------------------


def test_scan_with_ocr_disabled_never_reaches_the_cloud(client, db_session, writer, monkeypatch, tmp_path) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    scan = _scan_of(_text_pdf(tmp_path / "rechnung.pdf"), tmp_path / "rechnung_scan.pdf")

    events = _send_stream(client, "Bitte fasse das Dokument zusammen.", [scan])

    assert writer.received_payloads == []  # kein Cloud-Aufruf
    done = _done(events)
    assert done["blocked"] is True
    assert "nicht an die Cloud-KI gesendet" in done["content"]
    assert "„rechnung_scan.pdf“" in done["content"]
    assert "Texterkennung (OCR) ist in dieser Installation deaktiviert" in done["content"]
    assert "Estricharbeiten" not in done["content"]  # nichts erfunden/ausgewertet
    document = db_session.query(Document).one()
    assert document.ocr_status == "pending" and document.extracted_text is None
    audit = db_session.query(AuditEvent).filter_by(event_type="chat_document_not_ready").one()
    assert f"{document.id}=ocr_disabled" in audit.details and "rechnung" not in audit.details


def test_general_question_still_works_after_a_blocked_document(client, db_session, writer, monkeypatch, tmp_path) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    scan = _scan_of(_text_pdf(tmp_path / "rechnung.pdf"), tmp_path / "rechnung_scan.pdf")
    conversation_id = _conversation_id(_send_stream(client, "Bitte fasse das Dokument zusammen.", [scan]))

    general = _done(_send_stream(client, "Was versteht man unter einer Abmahnung?", conversation_id=conversation_id))
    assert general["blocked"] is False and general["content"] == "Antwort der Cloud-KI."
    assert len(writer.received_payloads) == 1
    # Der lokale Hinweis ist kein Gespraechsbeitrag und geht nicht in den Verlauf der Cloud-Anfrage ein.
    assert "nicht an die Cloud-KI gesendet" not in _payload_text(writer)

    follow_up = _done(_send_stream(client, "Was steht in dem Dokument?", conversation_id=conversation_id))
    assert follow_up["blocked"] is True and "rechnung_scan.pdf" in follow_up["content"]
    assert len(writer.received_payloads) == 1  # Folgefrage mit Dokumentbezug: weiterhin kein Cloud-Aufruf


def test_general_question_without_any_document(client, db_session, writer, monkeypatch, tmp_path) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    done = _done(_send_stream(client, "Welche Fristen gelten für eine Berufung?"))
    assert done["blocked"] is False and len(writer.received_payloads) == 1


# --- 2. Scan bei aktivierter OCR (isolierte QA-Konfiguration) ---------------------------------------------------------


@needs_tesseract
def test_scan_with_ocr_enabled_is_recognised_and_sent_through_the_privacy_pipeline(
    client, db_session, writer, monkeypatch, tmp_path
) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=True)
    source = tmp_path / "protokoll.pdf"
    generator.make_pdf(source)
    scan = _scan_of(source, tmp_path / "protokoll_scan.pdf")

    # Nicht-interaktiver Pfad (/send, Maskierung): der Inhalt geht pseudonymisiert und versiegelt an die Cloud.
    response = client.post(
        "/dashboard/chat/send",
        data={"csrf_token": _csrf(client), "conversation_id": "", "content": "Bitte fasse das Protokoll zusammen."},
        files={"documents": ("protokoll_scan.pdf", scan.read_bytes(), "application/pdf")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert db_session.query(Document).one().ocr_status == "done"
    assert len(writer.received_payloads) == 1
    payload = _payload_text(writer)
    assert "Estrich" in payload and "Nacherfüllung" in payload  # Dokumentinhalt ist tatsaechlich enthalten
    assert [value for value in generator.GOLD["personal"] if value in payload] == []  # keine Klarnamen


@needs_tesseract
def test_blank_scan_with_ocr_enabled_is_no_usable_text(client, db_session, writer, monkeypatch, tmp_path) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=True)
    done = _done(_send_stream(client, "Bitte auswerten.", [_blank_scan(tmp_path / "leer_scan.pdf")]))
    assert writer.received_payloads == []
    assert "keinen verwertbaren Text" in done["content"]


# --- 3. Verarbeitungsfehler, leerer Text, laufende Verarbeitung --------------------------------------------------------


def test_extraction_error_is_shown_transparently(client, db_session, writer, monkeypatch, tmp_path) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    broken = tmp_path / "beschaedigt.pdf"
    broken.write_bytes(b"%PDF-1.4 kein gueltiges PDF")
    done = _done(_send_stream(client, "Bitte fasse das Dokument zusammen.", [broken]))
    assert writer.received_payloads == []
    assert "Fehler aufgetreten (Textextraktion fehlgeschlagen:" in done["content"]
    assert str(tmp_path) not in done["content"]  # Fehlerangabe ohne Pfad


def test_empty_docx_is_no_usable_text(client, db_session, writer, monkeypatch, tmp_path) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    done = _done(_send_stream(client, "Bitte fasse das Dokument zusammen.", [_empty_docx(tmp_path / "leer.docx")]))
    assert writer.received_payloads == []
    assert "keinen verwertbaren Text" in done["content"]


def _service_setup(db: Session, tmp_path, document: Document, content: str = "Bitte fasse das Dokument zusammen."):
    from app.models import User

    user = User(email="qa@kanzlei.test", password_hash="x", display_name="QA")
    db.add(user)
    db.commit()
    service = ChatService(tmp_path)
    conversation = service.create_conversation(db, user=user, matter_id=None, title="Test", actor="qa")
    document.matter_id = conversation.matter_id
    db.add(document)
    db.commit()
    message = service.record_user_message(db, conversation=conversation, content=content, document_ids=[document.id])
    return service, conversation, message


def test_document_still_processing_is_not_evaluated(db_session, tmp_path) -> None:
    """Spaltenstandard "not_needed" ohne Text und ohne Abschluss-Audit = Verarbeitung laeuft noch."""
    document = Document(file_path=str(tmp_path / "x.pdf"), original_filename="eingang.pdf")
    service, conversation, message = _service_setup(db_session, tmp_path, document)
    writer = FakeClaudeWritingProvider()
    reply = service.send_message(
        db_session, conversation=conversation, content=message.content,
        drafting_service=_working_drafting_service(writer), actor="qa", current_message_id=message.id,
    )
    assert writer.received_payloads == []
    assert reply.blocked and "wird noch verarbeitet" in reply.content and "erneut senden" in reply.content


# --- 4. Erfolgreiches PDF und DOCX --------------------------------------------------------------------------------------


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_successful_documents_are_unchanged(client, db_session, writer, monkeypatch, tmp_path, kind) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    if kind == "pdf":
        path = _text_pdf(tmp_path / "rechnung.pdf")
    else:
        from docx import Document as DocxDocument

        path = tmp_path / "rechnung.docx"
        docx = DocxDocument()
        docx.add_paragraph(SIMPLE_TEXT)
        docx.save(str(path))
    done = _done(_send_stream(client, "Bitte fasse das Dokument zusammen.", [path]))
    assert done["blocked"] is False and done["content"] == "Antwort der Cloud-KI."
    assert len(writer.received_payloads) == 1
    assert "Estricharbeiten" in _payload_text(writer)  # Inhalt ist im versiegelten Payload
    assert db_session.query(AuditEvent).filter_by(event_type="chat_document_not_ready").count() == 0


@pytest.mark.parametrize("kind", ["docx", "pdf"])
def test_realistic_documents_still_pass_the_gold_standard(tmp_path, kind) -> None:
    path = tmp_path / f"protokoll.{kind}"
    (generator.make_docx if kind == "docx" else generator.make_pdf)(path)
    assert not extract_text(path).needs_ocr  # Logo-/Text-PDF gilt nicht als Teil-Scan
    result = harness.run(path, generator.GOLD)
    assert result["payload_sent"] and result["leaks"] == [] and result["keep_missing"] == []


# --- 5. Mehrere bzw. teilweise verarbeitete Dokumente ------------------------------------------------------------------


def test_mixed_attachments_block_and_name_each_state(client, db_session, writer, monkeypatch, tmp_path) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    text_pdf = _text_pdf(tmp_path / "lesbar.pdf")
    scan = _scan_of(_text_pdf(tmp_path / "quelle.pdf"), tmp_path / "gescannt.pdf")
    done = _done(_send_stream(client, "Vergleiche die beiden Dokumente.", [text_pdf, scan]))
    assert writer.received_payloads == []
    assert "„gescannt.pdf“ ist ein Scan" in done["content"]
    assert "Bereits lesbar: „lesbar.pdf“" in done["content"]


def test_partially_scanned_pdf_is_not_treated_as_complete(client, db_session, writer, monkeypatch, tmp_path) -> None:
    mixed = _mixed_pdf(tmp_path / "teilscan.pdf", tmp_path)
    result = extract_text(mixed)
    assert result.needs_ocr and "Estricharbeiten" in result.text  # Textseite vorhanden, Scan-Seite fehlt

    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    done = _done(_send_stream(client, "Bitte fasse das Dokument zusammen.", [mixed]))
    assert writer.received_payloads == []
    assert "„teilscan.pdf“" in done["content"]
    assert db_session.query(Document).one().extracted_text is None  # Teiltext nicht als Akteninhalt gespeichert


def test_unready_matter_document_is_named_below_an_answer(db_session, tmp_path) -> None:
    """Frage mit Dokumentbezug ohne Anhang in der Unterhaltung: die Antwort nennt nicht verarbeitete Aktendokumente."""
    from app.models import User

    user = User(email="qa@kanzlei.test", password_hash="x", display_name="QA")
    db_session.add(user)
    db_session.commit()
    service = ChatService(tmp_path)
    conversation = service.create_conversation(db_session, user=user, matter_id=None, title="T", actor="qa")
    db_session.add_all(
        [
            Document(matter_id=conversation.matter_id, file_path="a", original_filename="lesbar.pdf",
                     extracted_text=SIMPLE_TEXT, ocr_status="not_needed"),
            Document(matter_id=conversation.matter_id, file_path="b", original_filename="scan_akte.pdf",
                     ocr_status="pending"),
        ]
    )
    db_session.commit()
    message = service.record_user_message(db_session, conversation=conversation, content="Was steht im Dokument?")
    writer = FakeClaudeWritingProvider("Zusammenfassung.")
    reply = service.send_message(
        db_session, conversation=conversation, content=message.content,
        drafting_service=_working_drafting_service(writer), actor="qa", current_message_id=message.id,
    )
    assert len(writer.received_payloads) == 1 and not reply.blocked
    assert "Nicht ausgewertet, da ohne verarbeiteten Text: „scan_akte.pdf“" in reply.content
    assert "scan_akte" not in _payload_text(writer)


# --- 6. Statusaenderung waehrend der Anfrage ---------------------------------------------------------------------------


class _FlippingLocalAI(RuleBasedLocalAIProvider):
    """Aendert den Dokumentstatus zwischen Vorpruefung und Kontextaufbau (paralleler Request, Neuverarbeitung)."""

    def __init__(self, search_service, db, document_id, change) -> None:
        super().__init__(search_service)
        self._db, self._document_id, self._change = db, document_id, change

    def prepare_draft_context(self, matter_id, db):
        document = self._db.get(Document, self._document_id)
        self._change(document)
        self._db.commit()
        return super().prepare_draft_context(matter_id, db)


def _flip_service(db, document_id, change, writer):
    service = _working_drafting_service(writer)
    service.local_ai = _FlippingLocalAI(service.search_service, db, document_id, change)
    return service


@pytest.mark.parametrize(
    ("change", "expected"),
    [
        (lambda d: (setattr(d, "ocr_status", "pending"), setattr(d, "extracted_text", None)), "ist ein Scan"),
        (lambda d: setattr(d, "deleted_at", __import__("datetime").datetime(2026, 10, 13)), "inzwischen gelöscht"),
    ],
)
@pytest.mark.parametrize("streaming", [False, True])
def test_status_change_during_the_request_never_reaches_the_cloud(db_session, tmp_path, change, expected, streaming) -> None:
    document = Document(file_path="x", original_filename="vertrag.pdf", extracted_text=SIMPLE_TEXT, ocr_status="not_needed")
    service, conversation, message = _service_setup(db_session, tmp_path, document)
    writer = FakeClaudeWritingProvider()
    drafting = _flip_service(db_session, document.id, change, writer)
    kwargs = dict(conversation=conversation, content=message.content, drafting_service=drafting, actor="qa",
                  current_message_id=message.id)
    if streaming:
        events = list(service.send_message_stream(db_session, **kwargs))
        reply = events[-1].message
    else:
        reply = service.send_message(db_session, **kwargs)
    assert writer.received_payloads == []
    assert reply.blocked and expected in reply.content


def test_document_becoming_ready_before_the_check_is_used(db_session, tmp_path) -> None:
    """Umgekehrte Richtung: wird ein Dokument vor dem Kontextaufbau fertig, wird sein Inhalt normal verwendet."""
    document = Document(file_path="x", original_filename="nachgereicht.pdf", ocr_status="pending")
    service, conversation, message = _service_setup(db_session, tmp_path, document)
    writer = FakeClaudeWritingProvider("Antwort.")
    document.extracted_text, document.ocr_status = SIMPLE_TEXT, "done"
    db_session.commit()
    reply = service.send_message(
        db_session, conversation=conversation, content=message.content,
        drafting_service=_working_drafting_service(writer), actor="qa", current_message_id=message.id,
    )
    assert not reply.blocked and "Estricharbeiten" in _payload_text(writer)


# --- 7. Weitere Einstiege: /send, Dokumentaktion, Datenschutz-Rueckfrage ---------------------------------------------


def test_non_streaming_send_blocks_too(client, db_session, writer, monkeypatch, tmp_path) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    scan = _scan_of(_text_pdf(tmp_path / "q.pdf"), tmp_path / "bescheid_scan.pdf")
    response = client.post(
        "/dashboard/chat/send",
        data={"csrf_token": _csrf(client), "conversation_id": "", "content": "Bitte zusammenfassen."},
        files={"documents": ("bescheid_scan.pdf", scan.read_bytes(), "application/pdf")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert writer.received_payloads == []
    reply = db_session.query(ChatMessage).filter_by(role="assistant").one()
    assert reply.blocked and "bescheid_scan.pdf" in reply.content


def test_document_action_on_a_pending_matter_document(client, db_session, writer, monkeypatch, tmp_path) -> None:
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    from app.models import Client as ClientModel
    from app.models import Matter

    matter = Matter(client=ClientModel(name="Synthetischer Mandant"), title="Testakte")
    db_session.add(matter)
    db_session.flush()
    document = Document(matter_id=matter.id, file_path="x", original_filename="akte_scan.pdf", ocr_status="pending")
    db_session.add(document)
    db_session.commit()
    response = client.post(
        f"/dashboard/chat/from-document/{document.id}",
        data={"csrf_token": _csrf(client), "action": "summarize"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert writer.received_payloads == []
    reply = db_session.query(ChatMessage).filter_by(role="assistant").one()
    assert reply.blocked and "akte_scan.pdf" in reply.content


def test_privacy_review_is_not_offered_for_an_unready_document(db_session, tmp_path) -> None:
    document = Document(file_path="x", original_filename="scan.pdf", ocr_status="pending")
    service, conversation, _ = _service_setup(db_session, tmp_path, document)
    message, candidates = service.pending_privacy_review(
        db_session, conversation=conversation, drafting_service=_working_drafting_service(), actor="qa"
    )
    assert message is None and candidates == []


# --- 8. Statusmodell ---------------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("status", "text", "expected"),
    [
        ("not_needed", SIMPLE_TEXT, readiness.READY),
        ("done", SIMPLE_TEXT, readiness.READY),
        ("done", "   ", readiness.NO_TEXT),
        ("pending", None, readiness.OCR_DISABLED),
        ("failed", None, readiness.FAILED),
        ("unsupported_format", None, readiness.UNSUPPORTED),
        ("irgendwas", SIMPLE_TEXT, readiness.PROCESSING),  # unbekannt: fail-closed
    ],
)
def test_readiness_states(db_session, status, text, expected) -> None:
    document = Document(file_path="x", original_filename="d.pdf", ocr_status=status, extracted_text=text)
    db_session.add(document)
    db_session.commit()
    assert readiness.assess_document(db_session, document).state == expected


def test_readiness_respects_matter_isolation(db_session) -> None:
    from app.models import Client as ClientModel
    from app.models import Matter

    first, second = (Matter(client=ClientModel(name=f"Mandant {i}"), title=f"Akte {i}") for i in (1, 2))
    db_session.add_all([first, second])
    db_session.flush()
    document = Document(matter_id=first.id, file_path="x", extracted_text=SIMPLE_TEXT, ocr_status="not_needed")
    db_session.add(document)
    db_session.commit()
    [item] = readiness.assess_documents(db_session, [document.id], matter_id=second.id)
    assert item.state == readiness.OTHER_MATTER and not item.ready


def test_ocr_configuration_is_not_changed_by_the_application() -> None:
    assert Settings.model_fields["ocr_enabled"].default is False
    assert ChatConversation is not None


def test_multi_paragraph_notice_is_rendered_as_one_block(client, db_session, writer, monkeypatch, tmp_path) -> None:
    """Befund im installierten Build: der Blockiert-Stil ist ein Flex-Container - jeder Absatz des Hinweises wurde
    zur eigenen Spalte. Der Inhalt steht jetzt in EINEM Block neben dem Warnsymbol."""
    login_as_admin(db_session, client)
    _use_settings(monkeypatch, tmp_path, ocr_enabled=False)
    scan = _scan_of(_text_pdf(tmp_path / "q.pdf"), tmp_path / "scan.pdf")
    conversation_id = _conversation_id(_send_stream(client, "Bitte fasse das Dokument zusammen.", [scan]))
    html = client.get(f"/dashboard/chat/{conversation_id}").text
    start = html.index('class="chat-message__blocked-body"')
    block = html[start : html.index("</div>", html.index("Fragen ohne Bezug", start))]
    assert "nicht an die Cloud-KI gesendet" in block and "Fragen ohne Bezug" in block
