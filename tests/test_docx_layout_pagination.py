"""Seitenumbrueche im lokalen DOCX-Seitensatz (app/documents/docx_layout.py) - Reproduktion 26.10.2026.

Befund: `pymupdf.Story` teilt eine Tabellenzeile nicht ueber Seiten. Eine einzellige Layouttabelle (typisch fuer
Briefvorlagen) wurde komplett auf Seite 2 geschoben (Seite 1 nur die Ueberschrift) und dort abgeschnitten - 15 von 32
Absaetzen fehlten in der Seitenansicht und konnten nicht markiert werden. Ausserdem: manuelle Seitenumbrueche wurden
ignoriert, weiche Zeilenumbrueche gingen verloren, und Ligaturen ("ff") verhinderten das Auffinden von Namen.

Ausschliesslich synthetische Dokumente."""

from __future__ import annotations

import io

import docx
import pymupdf
from docx.enum.text import WD_BREAK

from app.documents import docx_layout
from app.documents.pseudonymization_review import _text_flags
from tests.auth_test_utils import login_as_admin
from tests.test_pseudonymization_center import _review, _upload
from tests.test_schriftsatz_release_resume import _doc_of_conversation, writer  # noqa: F401 - Fixture
from tests.test_web_chat import client, db_session  # noqa: F401

SENT = "Die Abnahme der Bauleistung erfolgte am 12.03.2026; Maengel wurden im Protokoll festgehalten. "
FLAGS = _text_flags(pymupdf.TEXTFLAGS_TEXT)


def _base(title: str):
    document = docx.Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Kopfzeile: Kanzlei Muster (synthetisch)"
    section.footer.paragraphs[0].text = "Fusszeile: Seite"
    document.add_heading(title, level=1)
    return document


def _render(tmp_path, document, name: str):
    path = tmp_path / f"{name}.docx"
    document.save(str(path))
    return path, pymupdf.open("pdf", docx_layout.docx_pdf_bytes(path))


def _body_text(doc) -> str:
    lines = [line for page in doc for line in page.get_text(flags=FLAGS).splitlines()
             if not line.startswith(("Kopfzeile", "Fusszeile"))]
    return " ".join(" ".join(lines).split())


def _body_bottom(page) -> float:
    blocks = [b for b in page.get_text("blocks") if b[4].strip() and not b[4].startswith(("Kopfzeile", "Fusszeile"))]
    return max((b[3] for b in blocks), default=0.0)


def test_single_cell_layout_table_starts_on_page_one_and_loses_no_text(tmp_path) -> None:
    document = _base("Briefvorlage mit Layouttabelle")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Sehr geehrte Damen und Herren,"
    for i in range(30):
        cell.add_paragraph(f"Absatz {i + 1:02d}: " + SENT * 2)
    path, doc = _render(tmp_path, document, "layout")
    # vorher: Seite 1 endete nach der Ueberschrift (y=100 von 792), der Rest lag auf Seite 2 und war abgeschnitten
    assert _body_bottom(doc[0]) > doc[0].rect.height * 0.75
    text = _body_text(doc)
    assert all(f"Absatz {i + 1:02d}:" in text for i in range(30))
    assert "Absatz 30: " + " ".join(SENT.split()) in text


def test_tall_multi_column_row_is_not_pushed_or_cut_and_is_reported(tmp_path) -> None:
    document = _base("Text, dann hohe Tabellenzeile")
    for i in range(8):
        document.add_paragraph(f"Vortext {i + 1}: " + SENT * 2)
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Sachverhalt"
    table.cell(0, 1).text = "\n".join(f"Zeile {i + 1}: " + SENT * 3 for i in range(6))
    table.cell(1, 0).text = "Ansprechpartner"
    table.cell(1, 1).text = "Thomas Berger"
    path, doc = _render(tmp_path, document, "tall")
    assert _body_bottom(doc[0]) > doc[0].rect.height * 0.75  # Seite 1 wird gefuellt, nicht halb leer
    text = _body_text(doc)
    assert all(f"Zeile {i + 1}:" in text for i in range(6)) and "Thomas Berger" in text
    assert any("untereinander" in note for note in docx_layout.layout_notes(path))


def test_short_tables_keep_their_rows_side_by_side(tmp_path) -> None:
    document = _base("Kurze Tabelle")
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, (label, value) in enumerate([("Ansprechpartner", "Thomas Berger"), ("Anschrift", "Beispielstrasse 17"),
                                          ("Steuernummer", "12/345/67890")]):
        table.cell(row, 0).text, table.cell(row, 1).text = label, value
    path, doc = _render(tmp_path, document, "short")
    words = {w[4]: w for w in doc[0].get_text("words")}
    assert abs(words["Ansprechpartner"][1] - words["Thomas"][1]) < 2  # gleiche Zeile = nebeneinander
    assert docx_layout.layout_notes(path) == []


def test_manual_page_break_starts_a_new_page(tmp_path) -> None:
    document = _base("Mit Seitenumbruch")
    document.add_paragraph("Seite eins: " + SENT)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("Seite zwei: Anlage zum Schreiben.")
    path, doc = _render(tmp_path, document, "break")
    assert doc.page_count == 2
    assert "Seite eins" in doc[0].get_text() and "Seite zwei" in doc[1].get_text()
    assert "Kopfzeile: Kanzlei Muster" in doc[1].get_text()  # Kopfzeile auch auf der neuen Seite


def test_soft_line_breaks_inside_a_paragraph_are_kept(tmp_path) -> None:
    document = _base("Anschrift")
    document.add_paragraph("Finanzamt Berlin Mitte\nNeue Jakobstrasse 6\n10179 Berlin")
    path, doc = _render(tmp_path, document, "lines")
    lines = [line.strip() for line in doc[0].get_text().splitlines()]
    assert "Finanzamt Berlin Mitte" in lines and "Neue Jakobstrasse 6" in lines and "10179 Berlin" in lines


def test_ligatures_do_not_hide_findings_from_the_marking(client, db_session, writer) -> None:
    """"ff" wird im Satz als Ligatur gesetzt ("Betreﬀ"); ohne Aufloesung fand die Suche "Hoffmann" nicht."""
    login_as_admin(db_session, client)
    document = _base("Betreff: Einspruch (synthetisch)")
    document.add_paragraph("Geschaeftsfuehrer: Steffen Hoffmann, Beispielstrasse 17, 10115 Berlin")
    buffer = io.BytesIO()
    document.save(buffer)
    conversation_id = _upload(client, "ligatur.docx", buffer.getvalue(),
                              "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    data = _review(client, conversation_id, _doc_of_conversation(db_session, conversation_id))
    originals = [m["original"] for m in data["mappings"]]
    hoffmann = [o for o in originals if "Hoffmann" in o]
    assert hoffmann, originals  # synthetischer Name wird erkannt
    marked = {mark["term"] for page in data["pages"] for mark in page["marks"]}
    assert all(term in marked for term in hoffmann)
    assert not any("Hoffmann" in value for value in data["not_located"])
