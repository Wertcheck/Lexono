"""Zentraler Pseudonymisierungsdialog (25.10.2026): Pruefansicht (Seiten, Markierungen, Rueckfragen, Ergebnis) und
Bestaetigen mit den im Dialog getroffenen Entscheidungen - fuer Dokumentanalyse UND Schriftsatzerstellung.

Ausschliesslich synthetische Dokumente; Cloud = Test-Double (`FakeClaudeWritingProvider`, SIMULATION); keine
Produktiv-Datenbank."""

from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.documents import release
from app.models import ChatMessage
from app.models.document_release import DocumentRelease
from app.models.privacy_decision import PrivacyDecision
from tests.auth_test_utils import extract_upload_token, login_as_admin
from tests.test_schriftsatz_release_resume import (  # noqa: F401 - Fixtures/Hilfen
    AUFTRAG,
    _doc_of_conversation,
    writer,
)
from tests.test_web_chat import _csrf, _parse_sse_events, client, db_session  # noqa: F401

ANALYSE = "Bitte fasse das angehängte Dokument kurz zusammen."


def _pdf_multipage() -> bytes:
    """Zwei Seiten mit Kopf-/Fusszeile; der unsichere Begriff "Richter" steht auf Seite 2."""
    import pymupdf

    doc = pymupdf.open()
    for number, body in enumerate([
        ["Sachstand (synthetisch)", "Die Abnahme der Bauleistung erfolgte am 12.03.2026 vollstaendig.",
         "Geschaeftsfuehrer: Thomas Berger, Beispielstrasse 17, 10115 Berlin"],
        ["Fortsetzung", "Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024."],
    ], start=1):
        page = doc.new_page()
        page.insert_text((50, 30), "Kanzlei Muster - Kopfzeile (synthetisch)")
        y = 90
        for line in body:
            page.insert_text((50, y), line)
            y += 18
        page.insert_text((50, 810), f"Seite {number} von 2 - Fusszeile")
    data = doc.tobytes()
    doc.close()
    return data


def _docx_with_table() -> bytes:
    import docx

    document = docx.Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Kopfzeile: Kanzlei Muster (synthetisch)"
    section.footer.paragraphs[0].text = "Fusszeile: Seite 1"
    document.add_heading("Sachstand (synthetisch)", level=1)
    document.add_paragraph("Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Ansprechpartner"
    table.cell(0, 1).text = "Thomas Berger"
    table.cell(1, 0).text = "Anschrift"
    table.cell(1, 1).text = "Beispielstrasse 17, 10115 Berlin"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, filename: str, data: bytes, mime: str, content: str = AUFTRAG) -> str:
    events = _parse_sse_events(client.post(
        "/dashboard/chat/send-stream", files={"documents": (filename, data, mime)},
        data={"upload_token": extract_upload_token(client.get("/dashboard/chat").text),
              "csrf_token": _csrf(client), "conversation_id": "", "content": content},
    ).text)
    assert [e["kind"] for e in events] == ["start", "uploaded"], events
    return events[0]["conversation_id"]


def _order(client: TestClient, conversation_id: str, content: str) -> list[dict]:
    return _parse_sse_events(client.post(
        "/dashboard/chat/send-stream",
        data={"csrf_token": _csrf(client), "conversation_id": conversation_id, "content": content},
    ).text)


def _review(client: TestClient, conversation_id: str, doc) -> dict:
    response = client.get(f"/dashboard/chat/{conversation_id}/document/{doc.id}/pseudonymization-review")
    assert response.status_code == 200
    assert response.headers.get("cache-control") == "no-store"
    return response.json()


def _confirm(client: TestClient, conversation_id: str, doc, decisions: list[dict] | None, *, text_version=None,
             require_resolved: bool = True) -> list[dict]:
    data = {"csrf_token": _csrf(client), "document_id": doc.id,
            "text_version": text_version or release.text_version_of(doc.extracted_text)}
    if decisions is not None:
        data["decisions"] = json.dumps(decisions)
    if require_resolved:
        data["require_resolved"] = "1"
    return _parse_sse_events(client.post(f"/dashboard/chat/{conversation_id}/confirm-release-and-generate",
                                         data=data).text)


@pytest.fixture()
def pdf_case(client, db_session: Session, writer):
    login_as_admin(db_session, client)
    release.set_mode(db_session, release.MODE_AUTO_REVIEW, actor="test")
    conversation_id = _upload(client, "sachstand.pdf", _pdf_multipage(), "application/pdf")
    return conversation_id, _doc_of_conversation(db_session, conversation_id)


# --- Pruefansicht -----------------------------------------------------------------------------------------------------


def test_review_shows_original_pages_with_marks_at_their_position_pdf(client, db_session, writer, pdf_case) -> None:
    conversation_id, doc = pdf_case
    data = _review(client, conversation_id, doc)
    assert data["viewer_kind"] == "pages" and len(data["pages"]) == 2  # mehrseitig, Originalseiten als Bild
    assert data["page_url_base"] == f"/dashboard/matters/{doc.matter_id}/document/{doc.id}/page/"
    assert data["text_version"] == release.text_version_of(doc.extracted_text)
    assert data["status"] == "review" and [c["term"] for c in data["candidates"]] == ["Richter"]
    # offene Rueckfrage an ihrer Fundstelle auf Seite 2 - eindeutig als "open" markiert
    open_marks = [(p["number"], m) for p in data["pages"] for m in p["marks"] if m["kind"] == "open"]
    assert [n for n, _ in open_marks] == [2] and open_marks[0][1]["term"] == "Richter"
    assert all(0 <= m["x"] <= 1 and 0 <= m["y"] <= 1 and m["w"] > 0 for _, m in open_marks)
    # sicher erkannte Werte als "masked" markiert (Seite 1), mit Platzhalter
    masked_terms = {m["term"] for p in data["pages"] for m in p["marks"] if m["kind"] == "masked"}
    assert "Thomas Berger" in masked_terms
    # pseudonymisiertes Ergebnis ist getrennt vom Original und enthaelt die Werte nicht
    assert data["result_text"] and "Thomas Berger" not in data["result_text"]
    assert writer.received_payloads == []  # lokal, kein Cloud-Aufruf


def test_review_of_docx_marks_table_cells_and_names_the_layout_limitation(client, db_session, writer) -> None:
    login_as_admin(db_session, client)
    release.set_mode(db_session, release.MODE_AUTO_REVIEW, actor="test")
    conversation_id = _upload(client, "sachstand.docx", _docx_with_table(),
                              "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    doc = _doc_of_conversation(db_session, conversation_id)
    data = _review(client, conversation_id, doc)
    assert data["viewer_kind"] == "pages" and data["pages"]
    assert any("DOCX" in note and "abweichen" in note for note in data["limitations"])  # transparent benannt
    # 26.10.2026: lokaler DOCX-Satz setzt Kopf-/Fusszeilen -> keine Luecken-Meldung mehr
    assert not any("Kopf-/Fußzeile(n)" in note for note in data["limitations"])
    marked = {m["term"] for p in data["pages"] for m in p["marks"]}
    assert any("Berger" in term for term in marked)  # Name in der Tabellenzelle an seiner Position
    assert "Richter" in {c["term"] for c in data["candidates"]}
    # nicht im Seitenbild auffindbare Werte werden ausdruecklich genannt, nie stillschweigend weggelassen
    all_terms = {m["original"] for m in data["mappings"]} | {c["term"] for c in data["candidates"]}
    assert all_terms <= marked | set(data["not_located"])


def test_review_is_isolated_to_the_own_conversation(client, db_session, writer, pdf_case) -> None:
    from app.models import Client as ClientModel, Document, Matter

    conversation_id, _doc = pdf_case
    other = Matter(client=ClientModel(name="Fremd-Mandant (synthetisch)"), title="Fremde Akte", reference_number="2027/1")
    db_session.add(other)
    db_session.commit()
    foreign = Document(matter_id=other.id, original_filename="fremd.txt", mime_type="text/plain",
                       file_path="synthetic/fremd.txt", extracted_text="Fremder Text.", ocr_status="not_needed")
    db_session.add(foreign)
    db_session.commit()
    response = client.get(f"/dashboard/chat/{conversation_id}/document/{foreign.id}/pseudonymization-review")
    assert response.status_code == 404


def test_opening_the_review_and_cancelling_persists_nothing(client, db_session, writer, pdf_case) -> None:
    """Abbruch: die Pruefansicht schreibt nichts; Entscheidungen leben bis zum Bestaetigen nur im Dialog."""
    conversation_id, doc = pdf_case
    _order(client, conversation_id, AUFTRAG)  # Gate blockt (Entscheidung erforderlich)
    _review(client, conversation_id, doc)
    db_session.expire_all()
    assert db_session.query(PrivacyDecision).count() == 0
    assert db_session.query(DocumentRelease).filter_by(document_id=doc.id).count() == 0
    assert writer.received_payloads == []


# --- Bestaetigen: serverseitige Pruefung der Dialog-Entscheidungen --------------------------------------------------


def test_confirm_requires_every_open_question_to_be_answered(client, db_session, writer, pdf_case) -> None:
    conversation_id, doc = pdf_case
    _order(client, conversation_id, AUFTRAG)
    events = _confirm(client, conversation_id, doc, [])
    assert events[0]["message_id"] is None and "offenen Rückfragen" in events[0]["content"]
    db_session.expire_all()
    assert db_session.query(DocumentRelease).filter_by(document_id=doc.id).count() == 0
    assert db_session.query(PrivacyDecision).count() == 0
    assert writer.received_payloads == []
    # Nachricht wurde wieder freigegeben: ein korrekter zweiter Versuch ist moeglich
    retry = _confirm(client, conversation_id, doc, [{"term": "Richter", "decision": "personal"}])
    assert next(e for e in retry if e["kind"] == "done").get("blocked") is not True


def test_confirm_rejects_a_decision_for_a_term_that_is_not_an_open_question(client, db_session, writer, pdf_case) -> None:
    """Keine Zuordnung ueber veraltete Treffer: nur AKTUELL offene Rueckfragen dieses Dokuments sind entscheidbar."""
    conversation_id, doc = pdf_case
    _order(client, conversation_id, AUFTRAG)
    events = _confirm(client, conversation_id, doc, [{"term": "Richter", "decision": "personal"},
                                                     {"term": "Mustermann", "decision": "not_personal"}])
    assert events[0]["message_id"] is None and "nicht mehr aktuell" in events[0]["content"]
    db_session.expire_all()
    assert db_session.query(PrivacyDecision).count() == 0 and writer.received_payloads == []


def test_confirm_rejects_a_stale_text_version_before_saving_decisions(client, db_session, writer, pdf_case) -> None:
    conversation_id, doc = pdf_case
    _order(client, conversation_id, AUFTRAG)
    events = _confirm(client, conversation_id, doc, [{"term": "Richter", "decision": "personal"}],
                      text_version="veraltet")
    assert events[0]["message_id"] is None and "geändert" in events[0]["content"]
    db_session.expire_all()
    assert db_session.query(PrivacyDecision).count() == 0 and writer.received_payloads == []


def test_schriftsatz_entry_confirm_with_dialog_decisions_drafts_pseudonymized(client, db_session, writer, pdf_case) -> None:
    """Einstieg Schriftsatzerstellung: Rueckfrage im Dialog beantwortet -> bestaetigen -> Entwurf; der Begriff ist
    pseudonymisiert und es gibt KEINE zweite Rueckfrage im Chat."""
    conversation_id, doc = pdf_case
    _order(client, conversation_id, AUFTRAG)
    listing = client.get(f"/dashboard/chat/{conversation_id}/pseudonymization-documents").json()
    assert listing["order_label"] == "Schriftsatz erstellen" and listing["has_open_order"] is True
    events = _confirm(client, conversation_id, doc, [{"term": "Richter", "decision": "personal"}])
    assert not [e for e in events if e["kind"] == "review"]
    assert next(e for e in events if e["kind"] == "done").get("blocked") is not True
    assert len(writer.received_payloads) == 1
    sent = writer.received_payloads[0].model_dump_json()
    assert "Richter" not in sent and "Thomas Berger" not in sent
    db_session.expire_all()
    decision = db_session.query(PrivacyDecision).filter_by(term="Richter").one()
    assert decision.matter_id == doc.matter_id and decision.decision == "personal"
    # Nach der aktenweiten Entscheidung ist der Dokumentstand unauffaellig -> freigegeben (automatisch oder manuell)
    assert release.release_status(db_session, doc).state in release.RELEASED_STATES
    assert db_session.query(ChatMessage).filter(ChatMessage.conversation_id == conversation_id,
                                                ChatMessage.draft_id.isnot(None)).count() == 1


def test_analysis_entry_uses_the_same_dialog_and_continues_the_analysis(client, db_session, writer) -> None:
    """Einstieg Dokumentanalyse: derselbe Dialog, Button "... und Analyse fortsetzen"; nach dem Bestaetigen wird die
    Analyse (kein Schriftsatz) fortgesetzt."""
    login_as_admin(db_session, client)
    release.set_mode(db_session, release.MODE_AUTO_REVIEW, actor="test")
    conversation_id = _upload(client, "sachstand.pdf", _pdf_multipage(), "application/pdf", content=ANALYSE)
    doc = _doc_of_conversation(db_session, conversation_id)
    blocked = _order(client, conversation_id, ANALYSE)
    assert next(e for e in blocked if e["kind"] == "done").get("blocked") is True  # Dokument-Gate
    page = client.get(f"/dashboard/chat/{conversation_id}/document/{doc.id}").text
    assert 'id="pseudonymization-center"' in page and f'data-open-document="{doc.id}"' in page
    listing = client.get(f"/dashboard/chat/{conversation_id}/pseudonymization-documents").json()
    assert listing["order_label"] == "Analyse fortsetzen"
    events = _confirm(client, conversation_id, doc, [{"term": "Richter", "decision": "not_personal"}])
    assert not [e for e in events if e["kind"] == "review"]
    done = next(e for e in events if e["kind"] == "done")
    assert done.get("blocked") is not True
    assert len(writer.received_payloads) == 1
    assert "Thomas Berger" not in writer.received_payloads[0].model_dump_json()
    db_session.expire_all()
    # "kein personenbezogenes Datum" gilt nur fuer DIESE Unterhaltung (bestehende Regel)
    assert db_session.query(PrivacyDecision).filter_by(term="Richter").one().conversation_id == conversation_id


def test_without_an_open_order_the_dialog_is_review_only_and_confirm_refuses(client, db_session, writer, pdf_case) -> None:
    """Nur hochgeladen (§57): kein Auftrag -> keine Fortsetzung, keine Weiterverarbeitung."""
    conversation_id, doc = pdf_case
    listing = client.get(f"/dashboard/chat/{conversation_id}/pseudonymization-documents").json()
    assert listing["has_open_order"] is False and listing["order_label"] == ""  # Dialog = reine Pruefansicht
    events = _confirm(client, conversation_id, doc, [{"term": "Richter", "decision": "personal"}])
    assert events[0]["message_id"] is None and "keinen offenen Auftrag" in events[0]["content"]
    assert writer.received_payloads == []


def test_an_answered_order_is_never_reprocessed_after_a_new_upload(client, db_session, writer) -> None:
    """Schutz: nach einer erledigten Antwort und einem weiteren Upload wird die alte Nachricht NICHT erneut gesendet."""
    login_as_admin(db_session, client)
    release.set_mode(db_session, release.MODE_MANUAL, actor="test")
    conversation_id = _upload(client, "a.pdf", _pdf_multipage(), "application/pdf")
    doc = _doc_of_conversation(db_session, conversation_id)
    _order(client, conversation_id, AUFTRAG)
    first = _confirm(client, conversation_id, doc, [{"term": "Richter", "decision": "personal"}])
    assert next(e for e in first if e["kind"] == "done").get("blocked") is not True
    assert len(writer.received_payloads) == 1
    # weiterer Upload in dieselbe Unterhaltung (Rolle "upload", kein Auftrag)
    _parse_sse_events(client.post(
        "/dashboard/chat/send-stream", files={"documents": ("b.pdf", _pdf_multipage(), "application/pdf")},
        data={"upload_token": extract_upload_token(client.get(f"/dashboard/chat/{conversation_id}").text),
              "csrf_token": _csrf(client), "conversation_id": conversation_id, "content": ""},
    ).text)
    again = _confirm(client, conversation_id, doc, None, require_resolved=False)
    assert again[0]["message_id"] is None
    assert len(writer.received_payloads) == 1  # nichts erneut gesendet


# --- Chat-Oberflaeche unveraendert ------------------------------------------------------------------------------------


def test_chat_page_keeps_its_controls_and_offers_the_dialog_closed(client, db_session, writer) -> None:
    """Der Chat bleibt die Arbeitsflaeche: der zentrale Dialog ist auf jeder Chat-Seite vorhanden, aber geschlossen
    und ohne Dokument; alle bisherigen Bedienelemente bleiben."""
    login_as_admin(db_session, client)
    page = client.get("/dashboard/chat").text
    center = page[page.index('<dialog id="pseudonymization-center"'):]
    center = center[:center.index(">")]
    assert " open" not in center and 'data-open-document=""' in center
    for marker in ('id="privacy-review-dialog"', 'class="chat-privacy-corner"', "chat-quick-action",
                   'id="sidebar-collapse-toggle"', "/dashboard/static/js/pseudonymization_center.js"):
        assert marker in page
    assert 'class="chat-document-pane"' not in page and "chat-shell--document-view" not in page


def test_document_view_is_the_chat_with_the_dialog_opened_for_that_document(client, db_session, writer, pdf_case) -> None:
    """Ursache der frueheren Dokumentspalte (`viewing_document` -> `chat-document-pane`) ist beseitigt: die
    Dokumentansicht zeigt den unveraenderten Chat und oeffnet den Dialog fuer genau dieses Dokument."""
    conversation_id, doc = pdf_case
    _order(client, conversation_id, AUFTRAG)
    page = client.get(f"/dashboard/chat/{conversation_id}/document/{doc.id}").text
    assert 'class="chat-document-pane"' not in page and "chat-shell--document-view" not in page
    assert 'id="document-text"' not in page
    assert f'data-open-document="{doc.id}"' in page and f'data-conversation-id="{conversation_id}"' in page
    assert 'id="chat-messages"' in page or "chat-message--user" in page  # Chat-Verlauf ist sichtbar
    # Sperrhinweis im Chat bietet den Dialog an (offener Auftrag + freigabepflichtiges Dokument)
    assert "data-pseudonymization-center-open" in page and "Pseudonymisierung prüfen" in page
    # Anhang-Chip oeffnet denselben Dialog statt einer Spalte
    assert f'data-pseudo-document-link data-document-id="{doc.id}"' in page


def test_dialog_document_list_covers_all_documents_of_the_conversation(client, db_session, writer) -> None:
    login_as_admin(db_session, client)
    release.set_mode(db_session, release.MODE_AUTO_REVIEW, actor="test")
    conversation_id = _upload(client, "a.pdf", _pdf_multipage(), "application/pdf")
    _parse_sse_events(client.post(
        "/dashboard/chat/send-stream", files={"documents": ("b.docx", _docx_with_table(),
                                                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"upload_token": extract_upload_token(client.get(f"/dashboard/chat/{conversation_id}").text),
              "csrf_token": _csrf(client), "conversation_id": conversation_id, "content": ""},
    ).text)
    _order(client, conversation_id, AUFTRAG)
    listing = client.get(f"/dashboard/chat/{conversation_id}/pseudonymization-documents")
    assert listing.headers.get("cache-control") == "no-store"
    docs = listing.json()["documents"]
    assert [d["filename"] for d in docs] == ["a.pdf", "b.docx"]
    assert all(d["needs_confirmation"] for d in docs) and all(d["text_version"] for d in docs)
    assert writer.received_payloads == []


def _two_document_case(client, db_session):
    login_as_admin(db_session, client)
    release.set_mode(db_session, release.MODE_AUTO_REVIEW, actor="test")
    conversation_id = _upload(client, "a.pdf", _pdf_multipage(), "application/pdf")
    _parse_sse_events(client.post(
        "/dashboard/chat/send-stream", files={"documents": ("b.docx", _docx_with_table(),
                                                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"upload_token": extract_upload_token(client.get(f"/dashboard/chat/{conversation_id}").text),
              "csrf_token": _csrf(client), "conversation_id": conversation_id, "content": ""},
    ).text)
    from app.models import Document

    docs = db_session.query(Document).order_by(Document.created_at).all()
    return conversation_id, docs


def _confirm_many(client, conversation_id, docs, decisions, *, versions=None):
    versions = versions or [release.text_version_of(d.extracted_text) for d in docs]
    return _parse_sse_events(client.post(f"/dashboard/chat/{conversation_id}/confirm-release-and-generate", data={
        "csrf_token": _csrf(client), "document_id": docs[0].id, "text_version": versions[0],
        "documents": json.dumps([{"document_id": d.id, "text_version": v} for d, v in zip(docs[1:], versions[1:])]),
        "decisions": json.dumps(decisions), "require_resolved": "1"}).text)


def test_multi_document_confirm_releases_all_documents_with_one_decision_set(client, db_session, writer) -> None:
    conversation_id, docs = _two_document_case(client, db_session)
    _order(client, conversation_id, AUFTRAG)
    events = _confirm_many(client, conversation_id, docs, [{"term": "Richter", "decision": "personal"}])
    assert not [e for e in events if e["kind"] == "review"]
    assert next(e for e in events if e["kind"] == "done").get("blocked") is not True
    assert len(writer.received_payloads) == 1
    sent = writer.received_payloads[0].model_dump_json()
    assert "Richter" not in sent and "Thomas Berger" not in sent
    db_session.expire_all()
    assert all(release.release_status(db_session, d).state in release.RELEASED_STATES for d in docs)


def test_multi_document_confirm_rejects_a_stale_second_document(client, db_session, writer) -> None:
    conversation_id, docs = _two_document_case(client, db_session)
    _order(client, conversation_id, AUFTRAG)
    events = _confirm_many(client, conversation_id, docs, [{"term": "Richter", "decision": "personal"}],
                           versions=[release.text_version_of(docs[0].extracted_text), "veraltet"])
    assert events[0]["message_id"] is None and "geändert" in events[0]["content"]
    db_session.expire_all()
    assert db_session.query(PrivacyDecision).count() == 0 and writer.received_payloads == []


def test_multi_document_confirm_rejects_a_foreign_document(client, db_session, writer) -> None:
    from app.models import Client as ClientModel, Document, Matter

    conversation_id, docs = _two_document_case(client, db_session)
    _order(client, conversation_id, AUFTRAG)
    other = Matter(client=ClientModel(name="Fremd (synthetisch)"), title="Fremd", reference_number="2027/9")
    db_session.add(other)
    db_session.commit()
    foreign = Document(matter_id=other.id, original_filename="fremd.txt", mime_type="text/plain",
                       file_path="synthetic/fremd.txt", extracted_text="Fremder Text.", ocr_status="not_needed")
    db_session.add(foreign)
    db_session.commit()
    events = _confirm_many(client, conversation_id, [docs[0], foreign], [{"term": "Richter", "decision": "personal"}])
    assert events[0]["message_id"] is None and "nicht zu dieser Unterhaltung" in events[0]["content"]
    assert writer.received_payloads == []


def test_review_marks_chat_question_terms_at_their_position(client, db_session, writer, pdf_case) -> None:
    """HITL aus dem Chat im selben Dialog: die vom Server gestellten Rueckfrage-Begriffe werden im Dokument an ihrer
    Fundstelle als offene Rueckfrage markiert (`terms`); unbekannte Begriffe erzeugen keine falschen Markierungen."""
    conversation_id, doc = pdf_case
    data = client.get(f"/dashboard/chat/{conversation_id}/document/{doc.id}/pseudonymization-review",
                      params=[("terms", "Kopfzeile"), ("terms", "Gibtesnicht")]).json()
    marks = [m for p in data["pages"] for m in p["marks"] if m["kind"] == "open"]
    assert {m["term"] for m in marks} >= {"Kopfzeile"} and len([m for m in marks if m["term"] == "Kopfzeile"]) == 2
    assert not any(m["term"] == "Gibtesnicht" for m in marks)
    # Begriffe aus der Chat-Nachricht, die im Dokument nicht vorkommen, sind KEIN Darstellungsproblem des Dokuments
    assert "Gibtesnicht" not in data["not_located"]


def test_docx_layout_renders_header_footer_and_table_borders(tmp_path) -> None:
    import pymupdf

    from app.documents.docx_layout import docx_pdf_bytes

    import docx

    plain = tmp_path / "plain.docx"
    plain.write_bytes(_docx_with_table())  # Tabelle OHNE Rahmen
    doc = pymupdf.open("pdf", docx_pdf_bytes(plain))
    text = doc[0].get_text()
    assert "Kopfzeile: Kanzlei Muster" in text and "Fusszeile: Seite 1" in text and "Thomas Berger" in text
    assert doc[0].get_drawings() == []  # keine erfundenen Rahmen
    grid_doc = docx.Document(str(plain))
    grid_doc.tables[0].style = "Table Grid"
    grid = tmp_path / "grid.docx"
    grid_doc.save(str(grid))
    drawn = pymupdf.open("pdf", docx_pdf_bytes(grid))[0].get_drawings()
    assert drawn  # Tabellenrahmen ("Table Grid") werden gezeichnet


def test_no_findings_document_has_no_questions_and_is_ok(client, db_session, writer) -> None:
    import pymupdf

    login_as_admin(db_session, client)
    pdf = pymupdf.open()
    pdf.new_page().insert_text((50, 80), "Die Abnahme der Bauleistung erfolgte vollstaendig.")
    conversation_id = _upload(client, "leer.pdf", pdf.tobytes(), "application/pdf")
    doc = _doc_of_conversation(db_session, conversation_id)
    data = _review(client, conversation_id, doc)
    assert data["status"] == "ok" and data["candidates"] == [] and data["result_text"]
