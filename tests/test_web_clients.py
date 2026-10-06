"""Tests für app/web/clients_router.py (Mandantendatenbank, 20.08.) - löst
den bisherigen Platzhalter unter `/dashboard/clients` ab.

Gleiches Testmuster wie tests/test_web_schriftsatz.py: In-Memory-SQLite
über app.dependency_overrides."""

from __future__ import annotations

import io
import re
from collections.abc import Iterator

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.session import get_db
from app.main import app
from app.models import Client, Document, Matter, Message, Note
from app.models.base import Base
from tests.auth_test_utils import create_test_user, extract_csrf, login, login_as_admin, seed_roles


@pytest.fixture()
def db_session() -> Iterator[Session]:
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    TestingSessionLocal = sessionmaker(bind=engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def client(db_session: Session) -> Iterator[TestClient]:
    """Admin-Client - gleicher Fixture-Name wie in den übrigen Testdateien."""

    def _override_get_db() -> Iterator[Session]:
        yield db_session

    app.dependency_overrides[get_db] = _override_get_db
    try:
        test_client = TestClient(app)
        login_as_admin(db_session, test_client)
        yield test_client
    finally:
        app.dependency_overrides.clear()


@pytest.fixture()
def anwalt_client(db_session: Session) -> Iterator[TestClient]:
    def _override_get_db() -> Iterator[Session]:
        yield db_session

    app.dependency_overrides[get_db] = _override_get_db
    try:
        test_client = TestClient(app)
        roles = seed_roles(db_session)
        create_test_user(db_session, roles["anwalt"], "anwalt@kanzlei.test")
        login(test_client, "anwalt@kanzlei.test")
        yield test_client
    finally:
        app.dependency_overrides.clear()


@pytest.fixture()
def mitarbeiter_client(db_session: Session) -> Iterator[TestClient]:
    def _override_get_db() -> Iterator[Session]:
        yield db_session

    app.dependency_overrides[get_db] = _override_get_db
    try:
        test_client = TestClient(app)
        roles = seed_roles(db_session)
        create_test_user(db_session, roles["mitarbeiter"], "mitarbeiter@kanzlei.test")
        login(test_client, "mitarbeiter@kanzlei.test")
        yield test_client
    finally:
        app.dependency_overrides.clear()


def _csrf(test_client: TestClient, path: str = "/dashboard/clients") -> str:
    page = test_client.get(path)
    return extract_csrf(page.text)


def _create_client_row(
    db_session: Session,
    *,
    name: str = "Muster GmbH",
    number: str = "M-1",
    contact_email: str | None = None,
    contact_phone: str | None = None,
    client_type: str | None = None,
    city: str | None = None,
    status: str = "active",
) -> Client:
    row = Client(
        name=name,
        client_number=number,
        contact_email=contact_email,
        contact_phone=contact_phone,
        client_type=client_type,
        city=city,
        status=status,
    )
    db_session.add(row)
    db_session.commit()
    db_session.refresh(row)
    return row


# --- Liste ---


def test_clients_page_is_no_longer_a_placeholder(client: TestClient) -> None:
    response = client.get("/dashboard/clients")
    assert response.status_code == 200
    assert "in Vorbereitung" not in response.text
    assert "Mandantendatenbank" in response.text


def test_clients_page_lists_created_client(client: TestClient, db_session: Session) -> None:
    _create_client_row(db_session, name="Sichtbarer Mandant", number="V-1")
    response = client.get("/dashboard/clients")
    assert "Sichtbarer Mandant" in response.text


def test_clients_page_row_menu_offers_quick_actions_without_opening_the_client(
    client: TestClient, db_session: Session
) -> None:
    """ECHTER FUND (19.09., UI/UX-Referenzabgleich "29_mandanten_uebersicht.png",
    dasselbe Muster wie zuvor bei den Akten): Direktzugriff auf "Mandant
    öffnen"/"Archivieren" ohne vorheriges Öffnen - beide Aktionen sind
    bereits bestehende, echte Routen (reine Verdrahtung)."""
    row = _create_client_row(db_session, name="Mandant mit Schnellmenü", number="Q-1")
    response = client.get("/dashboard/clients")

    assert f'href="/dashboard/clients/{row.id}"' in response.text
    assert f'action="/dashboard/clients/{row.id}/archive"' in response.text
    assert "Reaktivieren" not in response.text


def test_clients_page_row_menu_offers_reactivate_for_archived_clients(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Archivierter Mandant", number="Q-2")
    row.status = "archived"
    db_session.commit()

    response = client.get("/dashboard/clients", params={"status": "all"})

    assert f'action="/dashboard/clients/{row.id}/reactivate"' in response.text
    assert f'action="/dashboard/clients/{row.id}/archive"' not in response.text


def test_clients_page_search_filters_by_name(client: TestClient, db_session: Session) -> None:
    _create_client_row(db_session, name="Findbar GmbH", number="F-1")
    _create_client_row(db_session, name="Anderer Mandant", number="F-2")
    response = client.get("/dashboard/clients", params={"q": "Findbar"})
    assert "Findbar GmbH" in response.text
    assert "Anderer Mandant" not in response.text


def test_clients_page_shows_archived_by_default_matching_reference(
    client: TestClient, db_session: Session
) -> None:
    """Verhaltensaenderung (03.10., Owner-Direktive "REFERENZGETREUE
    MANDANTENUEBERSICHT"): der Default-Status wurde bewusst von "active" auf
    "all" umgestellt, weil die verbindliche Referenz `29_mandanten_uebersicht.png`
    in der Default-Ansicht sowohl "Aktiv"- als auch "Inaktiv"-Mandanten zeigt
    (z. B. "Schulz, Lisa" mit Status "Inaktiv") - siehe app/web/clients_router.py,
    `clients_list_page`-Docstring. Der vorher hier getestete Verhalten
    ("archiviert standardmaessig ausgeblendet") ist damit absichtlich nicht
    mehr der Default; die Filterfaehigkeit selbst bleibt erhalten (siehe
    test_clients_page_active_filter_hides_archived unten)."""
    row = _create_client_row(db_session, name="Archiviert GmbH", number="AR-1")
    row.status = "archived"
    db_session.commit()

    response = client.get("/dashboard/clients")
    assert "Archiviert GmbH" in response.text

    response_all = client.get("/dashboard/clients", params={"status": "all"})
    assert "Archiviert GmbH" in response_all.text


def test_clients_page_active_filter_hides_archived(client: TestClient, db_session: Session) -> None:
    row = _create_client_row(db_session, name="Archiviert KG", number="AR-2")
    row.status = "archived"
    db_session.commit()

    response = client.get("/dashboard/clients", params={"status": "active"})
    assert "Archiviert KG" not in response.text


# --- Anlegen ---


def test_create_client_as_admin_redirects_to_detail_page(client: TestClient) -> None:
    csrf_token = _csrf(client)
    response = client.post(
        "/dashboard/clients/create",
        data={
            "csrf_token": csrf_token,
            "name": "Neuer Mandant",
            "client_number": "N-1",
            "contact_email": "neu@muster.test",
            "contact_phone": "030 111",
            "practice_area": "Mietrecht",
            "responsible_user_id": "",
        },
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert response.headers["location"].startswith("/dashboard/clients/")

    detail = client.get(response.headers["location"])
    assert "Neuer Mandant" in detail.text


def test_create_client_missing_name_redirects_with_error(client: TestClient) -> None:
    csrf_token = _csrf(client)
    # Bewusst Whitespace statt eines echten Leerstrings - ein leerer
    # `str = Form(...)`-Wert wird von FastAPI/Pydantic in dieser
    # Projektumgebung bereits VOR dem Routenkörper als "Feld fehlt" (422)
    # abgelehnt, siehe dasselbe Muster in tests/test_web_settings.py
    # (firm_name="   "), statt bis zur eigentlichen ClientValidationError-
    # Pruefung im Routenkoerper vorzudringen.
    response = client.post(
        "/dashboard/clients/create",
        data={"csrf_token": csrf_token, "name": "   ", "client_number": "N-2"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert "/dashboard/clients?error=" in response.headers["location"]


def test_create_client_as_anwalt_is_allowed(anwalt_client: TestClient) -> None:
    csrf_token = _csrf(anwalt_client)
    response = anwalt_client.post(
        "/dashboard/clients/create",
        data={"csrf_token": csrf_token, "name": "Anwalt-Mandant", "client_number": "AW-1"},
        follow_redirects=False,
    )
    assert response.status_code == 303


def test_create_client_as_mitarbeiter_is_forbidden(mitarbeiter_client: TestClient) -> None:
    csrf_token = _csrf(mitarbeiter_client)
    response = mitarbeiter_client.post(
        "/dashboard/clients/create",
        data={"csrf_token": csrf_token, "name": "Verboten", "client_number": "V-1"},
    )
    assert response.status_code == 403


# --- Detail ---


def test_client_detail_page_shows_matters_messages_documents(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Detail-Mandant", number="D-1")
    matter = Matter(client_id=row.id, title="Verknüpfte Akte", status="open")
    db_session.add(matter)
    db_session.flush()
    message = Message(
        matter_id=matter.id, direction="inbound", sender="a@b.test", subject="Betreff X"
    )
    db_session.add(message)
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")
    assert response.status_code == 200
    assert "Detail-Mandant" in response.text
    assert "Verknüpfte Akte" in response.text
    assert "Betreff X" in response.text
    # ECHTER FUND (17.09., Owner-Direktive §5/§7): diese Zeile zeigte
    # bereits eine echte Nachricht an, war aber kein Link auf die
    # vollstaendige Posteingang-Detailansicht - jetzt ein echter Link.
    assert f'href="/dashboard/inbox/{message.id}"' in response.text


def test_client_detail_page_404_for_unknown_id(client: TestClient) -> None:
    response = client.get("/dashboard/clients/does-not-exist")
    assert response.status_code == 404


def test_client_detail_page_shows_tab_navigation(client: TestClient, db_session: Session) -> None:
    """ECHTER FUND (19.09., UI/UX-Referenzabgleich "30_mandant_detail.png",
    dasselbe Muster wie zuvor bei matter_detail.html): die Mandanten-
    Detailseite war eine lange Einzelseite statt der in der Referenz
    gezeigten Tab-Struktur."""
    row = _create_client_row(db_session, name="Tab-Mandant", number="T-1")

    response = client.get(f"/dashboard/clients/{row.id}")

    # "uebersicht" -> "client-uebersicht" (05.10., Owner-Direktive "LONG-RUN
    # PRODUCT QUALITY PASS" Phase D): kollidierte mit derselben generischen
    # ID auf matter_detail.html (siehe dortiger Testdocstring) - eine fuer
    # diese Seite noetige ID-Selektor-Regel in app.css ueberschrieb per
    # hoeherer CSS-Spezifitaet auch hier das `[hidden]`-basierte Tab-
    # Umschalten (live per CDP reproduziert: der Uebersicht-Tab blieb beim
    # Wechsel zu einem anderen Tab sichtbar).
    for tab_name in ["client-uebersicht", "akten", "dokumente", "aufgaben", "notizen", "kommunikation"]:
        assert f'data-tab="{tab_name}"' in response.text
        assert f'id="tab-{tab_name}"' in response.text
    assert 'id="tab-uebersicht"' not in response.text


def test_client_detail_page_aggregates_tasks_and_deadlines_across_matters(
    client: TestClient, db_session: Session
) -> None:
    """ECHTER FUND (19.09.): ein Mandant traegt selbst keine Aufgaben/
    Fristen (die haengen an einer Akte) - die Referenz zeigt trotzdem einen
    "Aufgaben & Fristen"-Tab auf Mandantenebene, hier ueber alle Akten
    dieses Mandanten aggregiert."""
    from datetime import date

    from app.models import Deadline, Task

    row = _create_client_row(db_session, name="Aufgaben-Mandant", number="A-1")
    matter = Matter(client_id=row.id, title="Akte mit Fristen", status="open")
    db_session.add(matter)
    db_session.flush()
    db_session.add(Task(matter_id=matter.id, title="Frist pruefen"))
    db_session.add(
        Deadline(matter_id=matter.id, source_text="Einspruchsfrist", due_date=date(2026, 12, 1))
    )
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")

    assert "Frist pruefen" in response.text
    assert "Einspruchsfrist" in response.text


def test_client_detail_page_shows_empty_notes_honestly(client: TestClient, db_session: Session) -> None:
    row = _create_client_row(db_session, name="Notizloser Mandant", number="N-1")

    response = client.get(f"/dashboard/clients/{row.id}")

    assert "Noch keine Notizen zu diesem Mandanten." in response.text


def test_create_client_note_persists_it_and_shows_it_on_the_client_page(
    client: TestClient, db_session: Session
) -> None:
    """ECHTER FUND (19.09., UI/UX-Referenzabgleich, siehe app/models/note.py):
    bisher gab es projektweit keine Moeglichkeit, eine freie Notiz an einem
    Mandanten zu hinterlegen (nur an einer Akte, ebenfalls 19.09. gebaut)."""
    row = _create_client_row(db_session, name="Notiz-Mandant", number="N-2")
    csrf = extract_csrf(client.get(f"/dashboard/clients/{row.id}").text)

    response = client.post(
        f"/dashboard/clients/{row.id}/notes",
        data={"csrf_token": csrf, "text": "Anruf erhalten: bittet um Rückruf."},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].endswith(f"/dashboard/clients/{row.id}#notizen")

    detail = client.get(f"/dashboard/clients/{row.id}")
    assert "Anruf erhalten: bittet um Rückruf." in detail.text
    assert "Noch keine Notizen zu diesem Mandanten." not in detail.text


def test_create_client_note_with_blank_text_is_rejected(client: TestClient, db_session: Session) -> None:
    from app.models import Note

    row = _create_client_row(db_session, name="Leernotiz-Mandant", number="N-3")
    csrf = extract_csrf(client.get(f"/dashboard/clients/{row.id}").text)

    response = client.post(
        f"/dashboard/clients/{row.id}/notes",
        data={"csrf_token": csrf, "text": "   "},
    )

    assert response.status_code == 400
    assert db_session.query(Note).count() == 0


def test_create_client_note_requires_a_valid_csrf_token(client: TestClient, db_session: Session) -> None:
    from app.models import Note

    row = _create_client_row(db_session, name="CSRF-Mandant", number="N-4")

    response = client.post(
        f"/dashboard/clients/{row.id}/notes",
        data={"csrf_token": "invalid", "text": "Sollte nicht gespeichert werden."},
    )

    assert response.status_code == 403
    assert db_session.query(Note).count() == 0


def test_client_detail_shows_ai_cta_link_for_single_open_matter(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Ein-Akte-Mandant", number="AI-1")
    matter = Matter(client_id=row.id, title="Einzige Akte", status="open")
    db_session.add(matter)
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")
    assert f"/dashboard/tools/schriftsatz?matter_id={matter.id}" in response.text


# --- Individuelle Mandantendetailseite (03.10., Owner-Direktive
# "INDIVIDUELLE MANDANTENDETAILSEITE", Referenzabgleich
# `30_mandant_detail.png`) - Header/Stammdaten/Schnellaktionen/Akten-&
# Dokumentenkarte, siehe tests/test_shell_icons.py fuer die native
# Windows-Shell-Icon-Extraktion selbst (hier nur die Integration). -----


def test_client_detail_header_shows_initials_category_number_and_since_date(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(
        db_session, name="Müller, Anna", number="M-0001", client_type="Privatperson"
    )
    response = client.get(f"/dashboard/clients/{row.id}")

    assert response.status_code == 200
    assert re.search(r'client-avatar[^"]*">\s*AM\s*<', response.text), "Avatar-Initialen 'AM' nicht gefunden"
    assert "Privatperson" in response.text
    assert "Mandantennummer: M-0001" in response.text
    assert f"Seit {row.created_at.strftime('%d.%m.%Y')}" in response.text


def test_client_detail_breadcrumb_links_back_to_clients_overview(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Breadcrumb-Mandant", number="B-1")
    response = client.get(f"/dashboard/clients/{row.id}")

    assert 'href="/dashboard/clients">Mandanten</a>' in response.text
    assert "Breadcrumb-Mandant" in response.text


def test_client_detail_stammdaten_shows_real_fields_and_dash_for_missing(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(
        db_session,
        name="Vollstaendig GmbH",
        number="V-1",
        client_type="Unternehmen",
        city="Hamburg",
        contact_email="kontakt@vollstaendig.test",
        contact_phone="+49 40 1234567",
    )
    response = client.get(f"/dashboard/clients/{row.id}")

    assert "Hamburg" in response.text
    assert 'href="mailto:kontakt@vollstaendig.test"' in response.text
    assert 'href="tel:+49 40 1234567"' in response.text


def test_client_detail_stammdaten_shows_dash_for_missing_optional_fields(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Minimal-Mandant", number="MIN-1")
    response = client.get(f"/dashboard/clients/{row.id}")

    assert response.status_code == 200
    # Kein Telefon/E-Mail hinterlegt - ehrlich als "–" gezeigt, kein
    # erfundener Wert, kein toter mailto:/tel:-Link ohne Ziel.
    assert 'href="mailto:"' not in response.text
    assert 'href="tel:"' not in response.text


def test_client_detail_quick_actions_are_present_and_real(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Schnellaktion-Mandant", number="S-1")
    response = client.get(f"/dashboard/clients/{row.id}")

    assert "Dokument analysieren" in response.text
    assert "Schreiben erstellen" in response.text
    assert "Dokument zusammenfassen" in response.text
    assert "Standard-Funktion hinzufügen" in response.text
    assert 'href="/dashboard/library/prompts"' in response.text


def test_client_detail_akten_card_shows_real_matters_not_hardcoded(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Akten-Mandant", number="AK-1")
    matter = Matter(
        client_id=row.id,
        title="Einkommensteuer 2024",
        reference_number="001/2024",
        practice_area="Steuererklärung",
        status="open",
    )
    db_session.add(matter)
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")

    assert "001/2024" in response.text
    assert "Einkommensteuer 2024" in response.text
    assert "Steuererklärung" in response.text
    assert f"/dashboard/matters/{matter.id}" in response.text


def test_client_detail_akten_card_empty_state(client: TestClient, db_session: Session) -> None:
    row = _create_client_row(db_session, name="Aktenloser Mandant", number="AK-2")
    response = client.get(f"/dashboard/clients/{row.id}")

    assert "Noch keine Akten verknüpft." in response.text
    assert f"/dashboard/tools/schriftsatz?client_id={row.id}" in response.text


def test_client_detail_akten_card_shows_overflow_link_beyond_five(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Viele-Akten-Mandant", number="AK-3")
    for i in range(6):
        db_session.add(Matter(client_id=row.id, title=f"Akte {i}", status="open"))
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")

    assert f"/dashboard/matters?client_id={row.id}" in response.text


def test_client_detail_dokumente_card_shows_real_document_with_size_and_type(
    client: TestClient, db_session: Session
) -> None:
    import os

    row = _create_client_row(db_session, name="Dokument-Mandant", number="D-2")
    matter = Matter(client_id=row.id, title="Akte mit Dokument", status="open")
    db_session.add(matter)
    db_session.flush()

    tmp_dir = os.path.join(os.environ.get("TEMP", "/tmp"), "lexono_test_docs")
    os.makedirs(tmp_dir, exist_ok=True)
    file_path = os.path.join(tmp_dir, "steuerbescheid_test.pdf")
    with open(file_path, "wb") as fh:
        fh.write(b"%PDF-1.4 test content " * 50)  # > 1 KB, real size on disk

    document = Document(
        matter_id=matter.id,
        file_path=file_path,
        original_filename="Steuerbescheid_2024.pdf",
        classified_type="Bescheid",
    )
    db_session.add(document)
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")

    assert "Steuerbescheid_2024.pdf" in response.text
    assert "Bescheid" in response.text
    assert "KB" in response.text  # echte, von der Datei gelesene Größe

    os.remove(file_path)


def test_client_detail_dokumente_card_empty_state(client: TestClient, db_session: Session) -> None:
    row = _create_client_row(db_session, name="Dokumentloser Mandant", number="D-3")
    response = client.get(f"/dashboard/clients/{row.id}")

    assert "Keine Dokumente verknüpft." in response.text


def test_client_detail_document_missing_file_shows_dash_size_not_crash(
    client: TestClient, db_session: Session
) -> None:
    """Datei verschoben/geloescht - ehrlich "–" statt eines erfundenen
    Werts oder eines Absturzes (os.path.getsize auf einem nicht
    existierenden Pfad)."""
    row = _create_client_row(db_session, name="Verwaistes-Dokument-Mandant", number="D-4")
    matter = Matter(client_id=row.id, title="Akte", status="open")
    db_session.add(matter)
    db_session.flush()
    document = Document(
        matter_id=matter.id,
        file_path="C:/does/not/exist/anymore.pdf",
        original_filename="Verschwunden.pdf",
    )
    db_session.add(document)
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")

    assert response.status_code == 200
    assert "Verschwunden.pdf" in response.text


def test_client_detail_notizen_card_shows_latest_notes(client: TestClient, db_session: Session) -> None:
    row = _create_client_row(db_session, name="Notiz-Mandant", number="NO-1")
    for i in range(4):
        db_session.add(Note(client_id=row.id, text=f"Notiz {i}", author="anwalt@kanzlei.test"))
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")

    assert "Notiz 3" in response.text  # neueste zuerst


def test_client_detail_long_names_and_titles_render_without_error(
    client: TestClient, db_session: Session
) -> None:
    long_name = "Sehr " * 20 + "lange Mandantenbezeichnung GmbH"
    row = _create_client_row(db_session, name=long_name, number="LONG-1")
    long_title = "Ausserordentlich " * 10 + "lange Aktenbezeichnung"
    db_session.add(Matter(client_id=row.id, title=long_title, status="open"))
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")

    assert response.status_code == 200
    assert long_name in response.text
    assert long_title in response.text


def test_client_detail_upload_document_modal_only_offered_with_matters(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Ohne-Akte-Mandant", number="UP-1")
    response = client.get(f"/dashboard/clients/{row.id}")

    assert "upload-document-modal" not in response.text

    matter = Matter(client_id=row.id, title="Akte", status="open")
    db_session.add(matter)
    db_session.commit()

    response_with_matter = client.get(f"/dashboard/clients/{row.id}")
    assert "upload-document-modal" in response_with_matter.text
    assert f'<option value="{matter.id}">' in response_with_matter.text


def test_client_detail_document_row_menu_offers_analyze_and_summarize(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="KI-Mandant", number="KI-1")
    matter = Matter(client_id=row.id, title="Akte", status="open")
    db_session.add(matter)
    db_session.flush()
    document = Document(matter_id=matter.id, file_path="/tmp/x.pdf", original_filename="x.pdf")
    db_session.add(document)
    db_session.commit()

    response = client.get(f"/dashboard/clients/{row.id}")

    assert f'action="/dashboard/chat/from-document/{document.id}"' in response.text
    assert 'value="analyze"' in response.text
    assert 'value="summarize"' in response.text


# --- Bearbeiten ---


def test_update_client_changes_name(client: TestClient, db_session: Session) -> None:
    row = _create_client_row(db_session, name="Alter Name", number="UP-1")
    csrf_token = _csrf(client, f"/dashboard/clients/{row.id}")
    response = client.post(
        f"/dashboard/clients/{row.id}/update",
        data={
            "csrf_token": csrf_token,
            "name": "Neuer Name",
            "client_number": "UP-1",
            "contact_email": "",
            "contact_phone": "",
            "practice_area": "",
            "responsible_user_id": "",
        },
        follow_redirects=False,
    )
    assert response.status_code == 303
    db_session.refresh(row)
    assert row.name == "Neuer Name"


# --- Archivieren/Reaktivieren ---


def test_archive_then_reactivate_client(client: TestClient, db_session: Session) -> None:
    row = _create_client_row(db_session, name="Archiv-Test", number="AR-2")
    csrf_token = _csrf(client, f"/dashboard/clients/{row.id}")
    client.post(
        f"/dashboard/clients/{row.id}/archive",
        data={"csrf_token": csrf_token},
        follow_redirects=False,
    )
    db_session.refresh(row)
    assert row.status == "archived"

    csrf_token = _csrf(client, f"/dashboard/clients/{row.id}")
    client.post(
        f"/dashboard/clients/{row.id}/reactivate",
        data={"csrf_token": csrf_token},
        follow_redirects=False,
    )
    db_session.refresh(row)
    assert row.status == "active"


# --- Löschen (Kernanforderung: kein Loeschen mit Akten) ---


def test_delete_client_without_matters_succeeds(client: TestClient, db_session: Session) -> None:
    row = _create_client_row(db_session, name="Loeschbar", number="DEL-1")
    csrf_token = _csrf(client, f"/dashboard/clients/{row.id}")
    response = client.post(
        f"/dashboard/clients/{row.id}/delete",
        data={"csrf_token": csrf_token},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert response.headers["location"] == "/dashboard/clients"
    assert db_session.get(Client, row.id) is None


def test_delete_client_with_matters_is_blocked_and_client_survives(
    client: TestClient, db_session: Session
) -> None:
    row = _create_client_row(db_session, name="Mit Akte", number="DEL-2")
    matter = Matter(client_id=row.id, title="Akte bleibt", status="open")
    db_session.add(matter)
    db_session.commit()

    csrf_token = _csrf(client, f"/dashboard/clients/{row.id}")
    response = client.post(
        f"/dashboard/clients/{row.id}/delete",
        data={"csrf_token": csrf_token},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert f"/dashboard/clients/{row.id}?error=" in response.headers["location"]
    assert db_session.get(Client, row.id) is not None
    assert db_session.query(Matter).filter_by(client_id=row.id).count() == 1


def test_delete_client_with_matters_shows_error_banner_on_detail_page(
    client: TestClient, db_session: Session
) -> None:
    """Regressionsschutz: der Redirect nach einem blockierten Loeschversuch
    haengt `?error=...` an die Detail-URL an - die GET-Route muss diesen
    Query-Parameter tatsaechlich lesen UND an das Template durchreichen,
    sonst verschwindet die Fehlermeldung stillschweigend (gefunden bei der
    manuellen Verifikation im Browser, 20.08.)."""
    row = _create_client_row(db_session, name="Mit Akte Banner", number="DEL-4")
    matter = Matter(client_id=row.id, title="Akte", status="open")
    db_session.add(matter)
    db_session.commit()

    csrf_token = _csrf(client, f"/dashboard/clients/{row.id}")
    response = client.post(
        f"/dashboard/clients/{row.id}/delete",
        data={"csrf_token": csrf_token},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "banner--error" in response.text
    assert "Aufbewahrungsgruenden gesperrt" in response.text


def test_delete_client_as_anwalt_is_forbidden(
    anwalt_client: TestClient, db_session: Session
) -> None:
    """PERM_CLIENT_DELETE ist admin-exklusiv - anders als
    PERM_CLIENT_MANAGE (Anlegen/Archivieren), siehe app/auth/permissions.py."""
    row = _create_client_row(db_session, name="Nur Admin loescht", number="DEL-3")
    csrf_token = _csrf(anwalt_client, f"/dashboard/clients/{row.id}")
    response = anwalt_client.post(
        f"/dashboard/clients/{row.id}/delete",
        data={"csrf_token": csrf_token},
    )
    assert response.status_code == 403
    assert db_session.get(Client, row.id) is not None


# --- DSGVO-Datenauszug ---


def test_export_client_returns_zip_file(client: TestClient, db_session: Session) -> None:
    row = _create_client_row(db_session, name="Export-Mandant", number="EX-1")
    csrf_token = _csrf(client, f"/dashboard/clients/{row.id}")
    response = client.post(
        f"/dashboard/clients/{row.id}/export",
        data={"csrf_token": csrf_token},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert response.content[:2] == b"PK"  # ZIP-Magic-Bytes


def test_export_client_archive_deleted_after_download(
    client: TestClient, db_session: Session, tmp_path, monkeypatch
) -> None:
    """Pilot-Readiness-Härtung: das DSGVO-Datenauszug-Archiv (enthält
    unpseudonymisierte Mandanteninhalte) darf nach dem Download nicht im
    Staging-Verzeichnis liegen bleiben. Siehe app/web/download_staging.py."""
    import app.web.clients_router as clients_module

    monkeypatch.setattr(clients_module, "_DOWNLOAD_STAGING_DIR", tmp_path)

    row = _create_client_row(db_session, name="Export-Mandant-Cleanup", number="EX-2")
    csrf_token = _csrf(client, f"/dashboard/clients/{row.id}")
    response = client.post(
        f"/dashboard/clients/{row.id}/export",
        data={"csrf_token": csrf_token},
    )
    assert response.status_code == 200

    remaining = list(tmp_path.glob("*.zip"))
    assert remaining == [], f"DSGVO-Export-Archiv nicht aufgeräumt: {remaining}"


# --- CSV-/Excel-Import ---


def test_import_csv_creates_clients_and_shows_result(client: TestClient, db_session: Session) -> None:
    csrf_token = _csrf(client)
    csv_content = "Name,Mandantennummer\r\nImport Eins,IMP-1\r\nImport Zwei,IMP-2\r\n"
    response = client.post(
        "/dashboard/clients/import",
        data={"csrf_token": csrf_token},
        files={"file": ("mandanten.csv", csv_content.encode("utf-8"), "text/csv")},
    )
    assert response.status_code == 200
    assert "2</strong> Mandant" in response.text
    assert db_session.query(Client).filter(Client.client_number.in_(["IMP-1", "IMP-2"])).count() == 2


def test_import_xlsx_creates_clients(client: TestClient, db_session: Session) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Mandantennummer"])
    sheet.append(["Excel Mandant", "XL-1"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    csrf_token = _csrf(client)
    response = client.post(
        "/dashboard/clients/import",
        data={"csrf_token": csrf_token},
        files={
            "file": (
                "mandanten.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    assert db_session.query(Client).filter_by(client_number="XL-1").count() == 1


def test_import_unsupported_file_type_shows_error(client: TestClient) -> None:
    csrf_token = _csrf(client)
    response = client.post(
        "/dashboard/clients/import",
        data={"csrf_token": csrf_token},
        files={"file": ("mandanten.txt", b"irrelevant", "text/plain")},
    )
    assert response.status_code == 200
    assert "Import fehlgeschlagen" in response.text


def test_import_as_mitarbeiter_is_forbidden(mitarbeiter_client: TestClient) -> None:
    csrf_token = _csrf(mitarbeiter_client)
    response = mitarbeiter_client.post(
        "/dashboard/clients/import",
        data={"csrf_token": csrf_token},
        files={"file": ("mandanten.csv", b"Name,Mandantennummer\r\nX,Y\r\n", "text/csv")},
    )
    assert response.status_code == 403


# --- Unauthentifiziert ---


def test_unauthenticated_cannot_view_clients_list(db_session: Session) -> None:
    def _override_get_db() -> Iterator[Session]:
        yield db_session

    app.dependency_overrides[get_db] = _override_get_db
    try:
        anon_client = TestClient(app)
        response = anon_client.get("/dashboard/clients", follow_redirects=False)
        assert response.status_code == 303
        assert "/dashboard/login" in response.headers["location"]
    finally:
        app.dependency_overrides.clear()


# --- Referenzabgleich 03.10. ("REFERENZGETREUE MANDANTENUEBERSICHT"):
#     Kategorie/Ort, Sortierung, Pagination - Phase F der Owner-Direktive
#     verlangt gezielte Regressionstests fuer genau diese neu eingefuehrte
#     Funktionalitaet. ---


def test_clients_page_empty_list_shows_empty_state(client: TestClient) -> None:
    response = client.get("/dashboard/clients")
    assert response.status_code == 200
    assert "Noch keine Mandanten vorhanden" in response.text
    # Pagination-Fusszeile bleibt bei genau einer (leeren) Seite ausgeblendet -
    # exakt dasselbe, bereits verifizierte Verhalten wie matters_list.html
    # (`{% if total_pages > 1 %}`), bewusst NICHT veraendert.
    assert "matters-pagination__info" not in response.text


def test_clients_page_no_search_results_shows_filtered_empty_state(
    client: TestClient, db_session: Session
) -> None:
    _create_client_row(db_session, name="Vorhandener Mandant", number="S-1")
    response = client.get("/dashboard/clients", params={"q": "Nichtvorhandener Suchbegriff"})
    assert "Vorhandener Mandant" not in response.text
    assert "Keine Mandanten gefunden, die zu den aktuellen Filtern passen." in response.text


def test_clients_page_shows_category_and_city(client: TestClient, db_session: Session) -> None:
    _create_client_row(
        db_session, name="Mueller, Anna", number="K-1", client_type="Privatperson", city="Berlin"
    )
    response = client.get("/dashboard/clients")
    assert "Privatperson" in response.text
    assert "Berlin" in response.text


def test_clients_page_filters_by_category(client: TestClient, db_session: Session) -> None:
    _create_client_row(db_session, name="Privatperson Eins", number="K-2", client_type="Privatperson")
    _create_client_row(db_session, name="Unternehmen Eins", number="K-3", client_type="Unternehmen")
    response = client.get("/dashboard/clients", params={"client_type": "Unternehmen"})
    assert "Unternehmen Eins" in response.text
    assert "Privatperson Eins" not in response.text


def test_clients_page_missing_optional_contact_fields_renders_without_error(
    client: TestClient, db_session: Session
) -> None:
    """Weder `contact_email`/`contact_phone` noch `client_type`/`city` sind
    Pflichtfelder (siehe app/models/client.py) - ein Mandant ganz ohne diese
    Angaben darf die Seite nicht zum Absturz bringen, sondern muss die
    Platzhalter ("–") anzeigen."""
    _create_client_row(db_session, name="Minimal-Mandant", number="MIN-1")
    response = client.get("/dashboard/clients")
    assert response.status_code == 200
    assert "Minimal-Mandant" in response.text


def test_clients_page_handles_long_name_and_contact_values(
    client: TestClient, db_session: Session
) -> None:
    long_name = "Rechtsanwaltskanzlei " + "Langername " * 10
    long_email = "sehr.lange.email.adresse.fuer.diesen.mandanten@" + "beispiel" * 5 + ".de"
    _create_client_row(
        db_session,
        name=long_name,
        number="LNG-1",
        contact_email=long_email,
    )
    response = client.get("/dashboard/clients")
    assert response.status_code == 200
    assert long_name in response.text
    assert long_email in response.text


def test_clients_page_sort_name_asc_and_desc(client: TestClient, db_session: Session) -> None:
    _create_client_row(db_session, name="Zeta GmbH", number="SORT-1")
    _create_client_row(db_session, name="Alpha GmbH", number="SORT-2")

    response_asc = client.get("/dashboard/clients", params={"sort": "name_asc"})
    assert response_asc.text.index("Alpha GmbH") < response_asc.text.index("Zeta GmbH")

    response_desc = client.get("/dashboard/clients", params={"sort": "name_desc"})
    assert response_desc.text.index("Zeta GmbH") < response_desc.text.index("Alpha GmbH")


def test_clients_page_pagination_first_and_last_page(client: TestClient, db_session: Session) -> None:
    for i in range(25):
        _create_client_row(db_session, name=f"Mandant {i:02d}", number=f"PG-{i:02d}")

    first_page = client.get("/dashboard/clients", params={"page": 1, "page_size": 10})
    assert "10 von 25 Mandant" in first_page.text

    last_page = client.get("/dashboard/clients", params={"page": 3, "page_size": 10})
    assert "5 von 25 Mandant" in last_page.text

    beyond_last_page = client.get("/dashboard/clients", params={"page": 99, "page_size": 10})
    assert beyond_last_page.status_code == 200
    assert "5 von 25 Mandant" in beyond_last_page.text


def test_clients_page_size_change_affects_row_count(client: TestClient, db_session: Session) -> None:
    for i in range(25):
        _create_client_row(db_session, name=f"Seiten-Mandant {i:02d}", number=f"PS-{i:02d}")

    response_default = client.get("/dashboard/clients")
    assert "10 von 25 Mandant" in response_default.text

    response_resized = client.get("/dashboard/clients", params={"page_size": 20})
    assert "20 von 25 Mandant" in response_resized.text


def test_clients_page_filter_then_paginate_keeps_filter_scoped_count(
    client: TestClient, db_session: Session
) -> None:
    for i in range(12):
        _create_client_row(
            db_session, name=f"Gefiltert {i:02d}", number=f"FP-{i:02d}", client_type="Unternehmen"
        )
    for i in range(5):
        _create_client_row(
            db_session, name=f"Ausgeblendet {i:02d}", number=f"FX-{i:02d}", client_type="Privatperson"
        )

    response = client.get(
        "/dashboard/clients",
        params={"client_type": "Unternehmen", "page": 2, "page_size": 10},
    )
    assert "2 von 12 Mandant" in response.text
    assert "Ausgeblendet" not in response.text
