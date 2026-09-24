"""Tests fuer app/documents/rendering.py (20.09., Owner-Direktive
"PRIORITAETSERGAENZUNG: ECHTER DOKUMENTVIEWER").

Nutzt die vorhandenen synthetischen Fixtures aus tests/fixtures/ sowie,
fuer mehrseitige Faelle, zur Laufzeit erzeugte synthetische PDF-/DOCX-
Dateien (keine echten Mandantendaten)."""

from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument

from app.documents.rendering import (
    DocumentRenderError,
    ViewerMode,
    clamp_dpi,
    determine_viewer_mode,
    get_page_count,
    render_page_png,
)

FIXTURES = Path(__file__).parent / "fixtures"

_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def _multipage_pdf(tmp_path: Path, pages: int = 3) -> Path:
    path = tmp_path / "multipage.pdf"
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Testseite {i + 1}")
    doc.save(path)
    doc.close()
    return path


def _multipage_docx(tmp_path: Path) -> Path:
    """Ein expliziter `add_page_break()` allein reicht bei PyMuPDFs
    eigener DOCX-Layout-Engine nicht immer aus, um tatsaechlich eine
    zweite Seite zu erzeugen (waehrend dieser Sitzung verifiziert) - der
    Seitenumbruch wird zusaetzlich mit genug Fuelltext kombiniert, damit
    real zwei Seiten entstehen."""
    path = tmp_path / "multipage.docx"
    docx_document = DocxDocument()
    docx_document.add_heading("Testdokument", level=1)
    docx_document.add_paragraph("Erster Absatz auf Seite 1.")
    docx_document.add_page_break()
    docx_document.add_heading("Seite 2", level=1)
    for i in range(40):
        docx_document.add_paragraph(f"Fuelltext Absatz Nummer {i} fuer Seite 2 des Testdokuments.")
    docx_document.save(str(path))
    return path


# --- get_page_count ---


def test_get_page_count_for_real_pdf() -> None:
    assert get_page_count(FIXTURES / "text_document.pdf") == 1


def test_get_page_count_for_real_docx() -> None:
    assert get_page_count(FIXTURES / "text_document.docx") >= 1


def test_get_page_count_for_multipage_pdf(tmp_path: Path) -> None:
    path = _multipage_pdf(tmp_path, pages=3)
    assert get_page_count(path) == 3


def test_get_page_count_for_multipage_docx(tmp_path: Path) -> None:
    """PyMuPDF layoutet DOCX mit seiner eigenen Engine (kein pixelgenaues
    Word-Rendering) - der Seitenumbruch fuehrt sicher zu MEHREREN Seiten,
    die exakte Anzahl haengt aber von PyMuPDFs eigener Reflow-Logik ab."""
    path = _multipage_docx(tmp_path)
    assert get_page_count(path) >= 2


def test_get_page_count_raises_for_missing_file(tmp_path: Path) -> None:
    with pytest.raises(DocumentRenderError):
        get_page_count(tmp_path / "does-not-exist.pdf")


def test_get_page_count_raises_for_corrupt_file(tmp_path: Path) -> None:
    corrupt = tmp_path / "corrupt.pdf"
    corrupt.write_bytes(b"not a real pdf")
    with pytest.raises(DocumentRenderError):
        get_page_count(corrupt)


# --- render_page_png ---


def test_render_page_png_returns_valid_png_bytes() -> None:
    png_bytes = render_page_png(FIXTURES / "text_document.pdf", 1)
    assert png_bytes.startswith(_PNG_MAGIC)


def test_render_page_png_works_for_docx() -> None:
    png_bytes = render_page_png(FIXTURES / "text_document.docx", 1)
    assert png_bytes.startswith(_PNG_MAGIC)


def test_render_page_png_renders_each_page_of_a_multipage_pdf(tmp_path: Path) -> None:
    path = _multipage_pdf(tmp_path, pages=3)
    for page_number in (1, 2, 3):
        png_bytes = render_page_png(path, page_number)
        assert png_bytes.startswith(_PNG_MAGIC)


def test_render_page_png_renders_each_page_of_a_multipage_docx(tmp_path: Path) -> None:
    path = _multipage_docx(tmp_path)
    for page_number in (1, 2):
        png_bytes = render_page_png(path, page_number)
        assert png_bytes.startswith(_PNG_MAGIC)


def test_render_page_png_raises_for_out_of_range_page(tmp_path: Path) -> None:
    path = _multipage_pdf(tmp_path, pages=2)
    with pytest.raises(DocumentRenderError):
        render_page_png(path, 5)


def test_render_page_png_raises_for_page_zero(tmp_path: Path) -> None:
    path = _multipage_pdf(tmp_path, pages=2)
    with pytest.raises(DocumentRenderError):
        render_page_png(path, 0)


def test_render_page_png_higher_dpi_produces_larger_image(tmp_path: Path) -> None:
    path = _multipage_pdf(tmp_path, pages=1)
    low = render_page_png(path, 1, dpi=50)
    high = render_page_png(path, 1, dpi=200)
    assert len(high) > len(low)


def test_clamp_dpi_stays_within_safe_range() -> None:
    assert clamp_dpi(1) == 40
    assert clamp_dpi(10_000) == 300
    assert clamp_dpi(150) == 150


# --- determine_viewer_mode ---


def test_determine_viewer_mode_pdf_is_pages() -> None:
    mode = determine_viewer_mode(FIXTURES / "text_document.pdf")
    assert mode.kind == "pages"
    assert mode.page_count == 1


def test_determine_viewer_mode_docx_is_pages() -> None:
    mode = determine_viewer_mode(FIXTURES / "text_document.docx")
    assert mode.kind == "pages"
    assert mode.page_count is not None and mode.page_count >= 1


def test_determine_viewer_mode_image_is_image() -> None:
    mode = determine_viewer_mode(FIXTURES / "scanned_image.png")
    assert mode == ViewerMode(kind="image")


def test_determine_viewer_mode_txt_is_text(tmp_path: Path) -> None:
    path = tmp_path / "notiz.txt"
    path.write_text("Ein einfacher Text.", encoding="utf-8")
    assert determine_viewer_mode(path) == ViewerMode(kind="text")


def test_determine_viewer_mode_unknown_extension_is_unsupported() -> None:
    mode = determine_viewer_mode(FIXTURES / "unbekannt.xyz")
    assert mode.kind == "unsupported"


def test_determine_viewer_mode_missing_pdf_file_is_unsupported(tmp_path: Path) -> None:
    """Ehrlicher Fallback statt Absturz, wenn die Datei laut Endung ein PDF
    waere, aber tatsaechlich nicht mehr auf der Platte liegt."""
    mode = determine_viewer_mode(tmp_path / "geloescht.pdf")
    assert mode.kind == "unsupported"


def test_determine_viewer_mode_corrupt_pdf_is_unsupported(tmp_path: Path) -> None:
    corrupt = tmp_path / "kaputt.pdf"
    corrupt.write_bytes(b"kein echtes PDF")
    mode = determine_viewer_mode(corrupt)
    assert mode.kind == "unsupported"
