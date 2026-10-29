"""Native Pruefung des Pseudonymisierungs-Pop-ups in der INSTALLIERTEN Lexono.exe (WebView2-Fenster).

- Isolierte Testablage (LEXONO_DATA_DIR im Job-Ordner, eigene SQLite-DB, eigener Port 8066 = eigener WebView2-Origin,
  dadurch kein gemeinsamer localStorage/Cookie mit der produktiven App auf 8062).
- Ausschliesslich synthetische Konten und Dokumente.
- Cloud = kontrollierter Test-Double: lokaler Aufzeichnungsserver qa_recorder.py (127.0.0.1:8765) ueber den im Code
  vorgesehenen Pruefweg LEXONO_QA_LOOPBACK_CLOUD_ENDPOINT=1, ungueltiger Testschluessel, Gateway aus, Proxy fuer alles
  andere tot. KEIN echter Cloud-E2E.
- Bedienung ueber das DevTools-Protokoll des App-Fensters, Screenshots per Win32 PrintWindow aus dem nativen Fenster.
Usage: native_popup_qa.py <vorher|nachher> <scratch_dir> <out_dir>"""
import asyncio, ctypes, ctypes.wintypes as wt, hashlib, io, json, os, re, shutil, sqlite3, subprocess, sys, time, uuid
import http.cookiejar, urllib.error, urllib.parse, urllib.request
from pathlib import Path

PHASE, S, OUT = sys.argv[1], Path(sys.argv[2]), Path(sys.argv[3])
assert PHASE in ("vorher", "nachher")
assert "jobs" in str(S), "nur isolierte Scratch-Ablage im Job-Ordner"
OUT.mkdir(parents=True, exist_ok=False)
T = Path(__file__).resolve().parent
REPO = Path(__file__).resolve().parents[4]  # Repository-Wurzel (pseudo_shots.py fuer STATE liegt im QA-Arbeitsordner)
EXE = Path(os.environ["LOCALAPPDATA"]) / "Lexono" / "Lexono.exe"
PORT, CDP_PORT, REC_PORT = 8066, 9334, 8765
BASE = f"http://127.0.0.1:{PORT}"
DB = S / "data" / "kanzlei_ai.db"
ADMIN, ADMIN_PW = "qa-native@lexono.test", "QA-native-0000"
STAFF, STAFF_PW = "qa-mitarbeiter@lexono.test", "QA-native-0000"
R = {"phase": PHASE, "exe": str(EXE), "exe_sha256": hashlib.sha256(EXE.read_bytes()).hexdigest().upper(), "aufnahmen": {},
     "funktion": {}}

user32, gdi32 = ctypes.windll.user32, ctypes.windll.gdi32
user32.SetProcessDpiAwarenessContext(ctypes.c_void_p(-4))  # Per-Monitor v2: echte Pixel

# ------------------------------------------------------------------ Testablage anlegen (installierte EXE migriert)
if S.exists():
    shutil.rmtree(S)
(S / "data").mkdir(parents=True)
(S / ".env").write_text("\n".join([
    f"DATABASE_URL=sqlite:///{DB.as_posix()}", "APP_ENV=development", "LOCAL_AI_ENABLED=false", "HOST=127.0.0.1",
    f"PORT={PORT}", "MAIL_AUTO_SYNC_ENABLED=false", "WEB_SEARCH_ENABLED=false", "LAW_UPDATE_CHECK_ENABLED=false",
    "SESSION_SECRET_KEY=qa-native-wegwerf-" + os.urandom(16).hex(), ""]), encoding="utf-8")
(S / ".setup_complete").write_text("1", encoding="utf-8")
BASE_ENV = dict(os.environ, LEXONO_DATA_DIR=str(S), PYTHONDONTWRITEBYTECODE="1")
for k in ("DATABASE_URL", "ANTHROPIC_API_KEY", "ANTHROPIC_BASE_URL"):
    BASE_ENV.pop(k, None)
subprocess.run([str(EXE), "migrate"], cwd=str(S), env=BASE_ENV, check=True, timeout=600,
               stdout=open(S / "migrate.out", "w"), stderr=subprocess.STDOUT)
sys.path.insert(0, str(REPO))
os.environ["DATABASE_URL"] = f"sqlite:///{DB.as_posix()}"
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import Session  # noqa: E402

from app.auth.security import hash_password  # noqa: E402
from app.models import Role, User  # noqa: E402

with Session(create_engine(f"sqlite:///{DB.as_posix()}")) as _db:
    roles = {r.name.lower(): r for r in _db.query(Role).all()}
    R["funktion"]["rollen"] = sorted(roles)
    for email, pw, role in ((ADMIN, ADMIN_PW, "admin"), (STAFF, STAFF_PW, "mitarbeiter")):
        _db.add(User(email=email, display_name=email.split("@")[0], role_id=roles[role].id, is_active=True,
                     password_hash=hash_password(pw), must_change_password=False))
    _db.commit()


def db_counts():
    c = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
    out = {t: c.execute(f"select count(*) from {t}").fetchone()[0]
           for t in ("privacy_decisions", "document_releases", "chat_messages", "chat_conversations")}
    c.close()
    return out


def recorder_requests():
    f = OUT / "recorder" / "requests.jsonl"
    return [json.loads(l) for l in open(f, encoding="utf-8")] if f.exists() else []


# ------------------------------------------------------------------ HTTP-Sitzungen (Datenanlage, Serverpruefungen)
def session():
    cj = http.cookiejar.CookieJar()
    return urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj), urllib.request.ProxyHandler({}))


op = session()


def get(path, o=None):
    return (o or op).open(BASE + path, timeout=300).read().decode("utf-8", "replace")


def post(path, fields, files=None, o=None):
    boundary = "----QA" + uuid.uuid4().hex
    body = b""
    for k, v in fields.items():
        body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode()
    for k, (name, data, ctype) in (files or {}).items():
        body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"; filename=\"{name}\"\r\n"
                 f"Content-Type: {ctype}\r\n\r\n").encode() + data + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    req = urllib.request.Request(BASE + path, data=body, headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    return (o or op).open(req, timeout=900).read().decode("utf-8", "replace")


def login(o, email, pw):
    o.open(BASE + "/dashboard/login", timeout=30)
    o.open(urllib.request.Request(BASE + "/dashboard/login", data=urllib.parse.urlencode(
        {"email": email, "password": pw, "next": "/dashboard"}).encode()), timeout=30)


def sse(text):
    return [json.loads(l[6:]) for l in text.splitlines() if l.startswith("data: ")]


def token(name, page):
    m = re.search(rf'name="{name}" value="([^"]+)"', page)
    return m.group(1) if m else ""


def pdf_bytes(pages):
    import pymupdf
    doc = pymupdf.open()
    for n, lines in enumerate(pages, 1):
        page = doc.new_page()
        page.insert_text((50, 30), "Kanzlei Muster - Kopfzeile (synthetisch)", fontsize=9)
        y = 90
        for line in lines:
            page.insert_text((50, y), line, fontsize=11); y += 20
        page.insert_text((50, 810), f"Seite {n} von {len(pages)}", fontsize=9)
    data = doc.tobytes(); doc.close(); return data


def docx_bytes():
    import docx
    d = docx.Document()
    d.add_heading("Stellungnahme (synthetisch)", level=1)
    d.add_paragraph("Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024.")
    d.add_paragraph("Ansprechpartner: Thomas Berger, Beispielstrasse 17, 10115 Berlin.")
    d.add_paragraph("Die Nordbauholding GmbH bittet um Aussetzung der Vollziehung.")
    b = io.BytesIO(); d.save(b); return b.getvalue()


MAIN_PAGES = [
    ["Einspruch gegen den Bescheid vom 02.02.2024 (synthetisches Musterdokument)",
     "Mandant: Max Bauer, Lindenweg 4, 50667 Koeln", "Hiermit legen wir fuer die Nordbauholding GmbH Einspruch ein.",
     "Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024.",
     "Wir bestaetigen den Eingang des Schreibens von Kaufmann vom 18.03.2024."],
    ["Begruendung (Fortsetzung)", "Wir bestaetigen den Eingang des Schreibens von Meister vom 20.03.2024.",
     "Rueckfragen bitte an max.bauer@example.org."],
    ["Anlagen (Fortsetzung)", "Wir bestaetigen den Eingang des Schreibens von Gärtner vom 22.03.2024.",
     "Wir bestaetigen den Eingang des Schreibens von Kellner vom 24.03.2024."],
]
SINGLE_PAGES = [["Vermerk zur Akte (synthetisches Musterdokument)", "Mandantin: Sabine Hartmann, Lindenweg 4, 50667 Koeln",
                 "Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024."]]
ANLAGE_PAGES = [["Anlage K1 (synthetisch)", "Die Nordbauholding GmbH bittet um Aussetzung der Vollziehung."]]
AUFTRAG = "Bitte erstelle einen Schriftsatz an das Finanzamt mit dem aktuellen Sachstand."
CHAT_FRAGE = "Wir bestätigen den Eingang des Schreibens von Richter vom 16.03.2024. Bitte fasse den Stand zusammen."
PDF, DOCX = "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOC_ORIGINALS = ["Max Bauer", "Lindenweg 4", "50667 Koeln", "Nordbauholding", "max.bauer@example.org", "Richter", "Meister",
                 "Gärtner", "Kellner"]


def new_conv(files, auftrag=AUFTRAG):
    conv = ""
    for name, data, ctype in files:
        page = get(f"/dashboard/chat/{conv}" if conv else "/dashboard/chat")
        ev = sse(post("/dashboard/chat/send-stream", {"upload_token": token("upload_token", page),
                                                      "csrf_token": token("csrf_token", page),
                                                      "conversation_id": conv, "content": ""},
                      files={"documents": (name, data, ctype)}))
        conv = conv or ev[0]["conversation_id"]
    if auftrag:
        ev = sse(post("/dashboard/chat/send-stream", {"csrf_token": token("csrf_token", get(f"/dashboard/chat/{conv}")),
                                                      "conversation_id": conv, "content": auftrag}))
        R["funktion"].setdefault("auftrag_blockiert", {})[conv] = next(e for e in ev if e["kind"] == "done").get("blocked")
    return conv


# ------------------------------------------------------------------ natives Fenster
def find_window(pid):
    found = []
    WNDENUMPROC = ctypes.WINFUNCTYPE(wt.BOOL, wt.HWND, wt.LPARAM)

    def cb(hwnd, _):
        p = wt.DWORD(); user32.GetWindowThreadProcessId(hwnd, ctypes.byref(p))
        if p.value == pid and user32.IsWindowVisible(hwnd) and not user32.GetWindow(hwnd, 4):
            r = wt.RECT(); user32.GetWindowRect(hwnd, ctypes.byref(r))
            if (r.right - r.left) > 400:
                found.append(hwnd)
        return True
    user32.EnumWindows(WNDENUMPROC(cb), 0)
    return found[0] if found else None


def window_rect(hwnd):
    r = wt.RECT(); user32.GetWindowRect(hwnd, ctypes.byref(r)); return r.left, r.top, r.right, r.bottom


def capture(hwnd, path):
    """PrintWindow(PW_RENDERFULLCONTENT): das Fensterbild selbst inkl. Titelleiste (auch Teile ausserhalb des
    Bildschirms); Rueckfall Bildschirmaufnahme, falls PrintWindow ein leeres Bild liefert."""
    from PIL import Image, ImageGrab, ImageStat
    user32.SetForegroundWindow(hwnd); time.sleep(0.3)
    l, t, r, b = window_rect(hwnd); w, h = r - l, b - t
    hdc_win = user32.GetWindowDC(hwnd); hdc = gdi32.CreateCompatibleDC(hdc_win)
    bmp = gdi32.CreateCompatibleBitmap(hdc_win, w, h); gdi32.SelectObject(hdc, bmp)
    ok = user32.PrintWindow(hwnd, hdc, 2)

    class BITMAPINFOHEADER(ctypes.Structure):
        _fields_ = [("biSize", wt.DWORD), ("biWidth", wt.LONG), ("biHeight", wt.LONG), ("biPlanes", wt.WORD),
                    ("biBitCount", wt.WORD), ("biCompression", wt.DWORD), ("biSizeImage", wt.DWORD),
                    ("biXPelsPerMeter", wt.LONG), ("biYPelsPerMeter", wt.LONG), ("biClrUsed", wt.DWORD),
                    ("biClrImportant", wt.DWORD)]
    bi = BITMAPINFOHEADER(); bi.biSize = ctypes.sizeof(bi); bi.biWidth = w; bi.biHeight = -h; bi.biPlanes = 1; bi.biBitCount = 32
    buf = ctypes.create_string_buffer(w * h * 4)
    gdi32.GetDIBits(hdc, bmp, 0, h, buf, ctypes.byref(bi), 0)
    gdi32.DeleteObject(bmp); gdi32.DeleteDC(hdc); user32.ReleaseDC(hwnd, hdc_win)
    img = Image.frombuffer("RGBA", (w, h), buf, "raw", "BGRA", 0, 1).convert("RGB")
    method = "PrintWindow"
    if not ok or max(ImageStat.Stat(img).stddev) < 3:
        img = ImageGrab.grab(bbox=(l, t, r, b), all_screens=True); method = "Bildschirmaufnahme"
    img.save(path)
    return {"methode": method, "fenster_px": [w, h], "position": [l, t]}


def set_client_css(hwnd, ev_sync, target_w, target_h):
    """Fenster so bemessen, dass der WebView-Viewport target_w x target_h CSS-Pixel hat (darf groesser als der
    Bildschirm sein - PrintWindow nimmt das ganze Fenster auf)."""
    for _ in range(8):
        if user32.IsZoomed(hwnd):
            user32.ShowWindow(hwnd, 9)  # SW_RESTORE
            time.sleep(1.0)
        iw, ih, dpr = json.loads(ev_sync("JSON.stringify([innerWidth, innerHeight, devicePixelRatio])"))
        if abs(iw - target_w) <= 1 and abs(ih - target_h) <= 1:
            break
        l, t, r, b = window_rect(hwnd)
        nw = (r - l) + round((target_w - iw) * dpr); nh = (b - t) + round((target_h - ih) * dpr)
        user32.SetWindowPos(hwnd, 0, 0, 0, nw, nh, 0x0004)
        time.sleep(1.5)
    return json.loads(ev_sync("JSON.stringify([innerWidth, innerHeight, devicePixelRatio])"))


STATE = "(() => { const v = " + (T / "pseudo_shots.py").read_text(encoding="utf-8").split('STATE = r"""', 1)[1].split('"""', 1)[0] + "; return JSON.stringify(v); })()"
WAIT_READY = """new Promise(r => { const t0 = Date.now(); (function poll() { const d = document.getElementById('pseudonymization-center');
 const busy = d && d.querySelector('.pseudo-center__loading'); const done = d ? d.querySelectorAll('.pseudo-center__doc-ok, .pseudo-center__doc-badge').length : 0;
 if ((d && d.open && !busy && done >= %d && d.querySelector('.pseudo-page, .pseudo-text, [data-question]')) || Date.now() - t0 > 60000) { setTimeout(() => r(Date.now() - t0), 1800); } else { setTimeout(poll, 300); } })(); })"""
ANSWER = """((dec) => { const b = document.querySelector('#pseudonymization-center .pseudo-question.is-active [data-decision="' + dec + '"]');
 if (!b) return 'keine aktive Frage'; b.click(); const a = document.querySelector('[data-pseudo-q-apply]');
 if (a && !a.hidden && !a.disabled && !a.closest('[hidden]')) a.click(); return 'ok'; })(%s)"""


async def main():
    rec = subprocess.Popen([sys.executable, str(T / "qa_recorder.py"), str(REC_PORT), str(OUT / "recorder")])
    # QA_FORCE_DPR1=1: WebView2 mit Geraetepixelverhaeltnis 1 (CSS-Pixel = Bildschirmpixel), damit 1920/1280/960 px
    # Breite trotz 125 % Windows-Skalierung als echte Fensterbreite geprueft werden koennen
    wv_args = f"--remote-debugging-port={CDP_PORT}" + (" --force-device-scale-factor=1" if os.environ.get("QA_FORCE_DPR1") == "1" else "")
    R["funktion"]["webview2_argumente"] = wv_args
    env = dict(BASE_ENV, WEBVIEW2_ADDITIONAL_BROWSER_ARGUMENTS=wv_args,
               ANTHROPIC_BASE_URL=f"http://127.0.0.1:{REC_PORT}", LEXONO_QA_LOOPBACK_CLOUD_ENDPOINT="1",
               ANTHROPIC_API_KEY="sk-ant-qa-invalid-test-key-0000", HTTPS_PROXY="http://127.0.0.1:9",
               HTTP_PROXY="http://127.0.0.1:9", NO_PROXY="127.0.0.1,localhost", PORT=str(PORT),
               LEXONO_GATEWAY_URL="", LEXONO_GATEWAY_CLIENT_ID="", LEXONO_GATEWAY_CLIENT_SECRET="")
    proc = subprocess.Popen([str(EXE), "serve"], cwd=str(S), env=env, stdout=open(S / "native_serve.out", "w"),
                            stderr=open(S / "native_serve.err", "w"))
    R["pid"] = proc.pid
    try:
        for _ in range(180):
            try:
                op.open(BASE + "/dashboard/login", timeout=3); break
            except Exception:
                time.sleep(1)
        login(op, ADMIN, ADMIN_PW)
        # ausgelieferter Pop-up-Code der installierten EXE vs. Repository (Nachweis, welche Oberflaeche laeuft)
        for rel in ("static/js/pseudonymization_center.js", "static/css/app.css"):
            served = op.open(BASE + "/dashboard/" + rel, timeout=30).read()
            repo = (REPO / "app" / "web" / rel).read_bytes()
            R["funktion"].setdefault("ausgelieferter_code", {})[rel] = {
                "sha256": hashlib.sha256(served).hexdigest()[:16], "gleich_repo_HEAD_arbeitsstand": served == repo,
                "neues_layout_marker": b"data-pseudo-q-apply" in served or b"pseudo-question__found" in served}
        t0 = time.time()
        convs = {"main": new_conv([("QA-SYNTH-Einspruch-Bauer.pdf", pdf_bytes(MAIN_PAGES), PDF)]),
                 "multi": new_conv([("QA-SYNTH-Einspruch-Bauer.pdf", pdf_bytes(MAIN_PAGES), PDF),
                                    ("QA-SYNTH-Stellungnahme.docx", docx_bytes(), DOCX),
                                    ("QA-SYNTH-Anlage-K1.pdf", pdf_bytes(ANLAGE_PAGES), PDF)]),
                 "single": new_conv([("QA-SYNTH-Vermerk.pdf", pdf_bytes(SINGLE_PAGES), PDF)])}
        R["funktion"]["convs"] = convs
        R["funktion"]["setup_s"] = round(time.time() - t0, 1)
        R["funktion"]["recorder_nach_setup"] = len(recorder_requests())  # Dokument-Gate: nichts gesendet
        R["funktion"]["db_start"] = db_counts()

        hwnd = None
        for _ in range(90):
            hwnd = find_window(proc.pid)
            if hwnd:
                break
            time.sleep(1)
        assert hwnd, "kein natives Fenster gefunden"
        targets = []
        for _ in range(60):
            try:
                targets = json.loads(urllib.request.urlopen(f"http://127.0.0.1:{CDP_PORT}/json", timeout=2).read())
                if any(t["type"] == "page" for t in targets):
                    break
            except Exception:
                pass
            time.sleep(1)
        page = next(t for t in targets if t["type"] == "page")
        import websockets
        async with websockets.connect(page["webSocketDebuggerUrl"], max_size=200_000_000) as ws:
            n = 0

            async def cdp(method, **params):
                nonlocal n
                n += 1
                await ws.send(json.dumps({"id": n, "method": method, "params": params}))
                while True:
                    msg = json.loads(await ws.recv())
                    if msg.get("id") == n:
                        return msg.get("result", {})

            async def ev(js):
                res = await cdp("Runtime.evaluate", expression=js, returnByValue=True, awaitPromise=True)
                if "exceptionDetails" in res:
                    return {"js_fehler": str(res["exceptionDetails"].get("exception", {}).get("description", ""))[:300]}
                return res.get("result", {}).get("value")

            loop = asyncio.get_running_loop()

            def ev_sync(js):
                return asyncio.run_coroutine_threadsafe(ev(js), loop).result(30)

            async def resize(w, h):
                return await asyncio.to_thread(set_client_css, hwnd, ev_sync, w, h)

            async def shot(name):
                await asyncio.sleep(0.9)
                info = await asyncio.to_thread(capture, hwnd, OUT / f"{name}.png")
                st = await ev(STATE)
                info["zustand"] = json.loads(st) if isinstance(st, str) else st
                info["viewport"] = json.loads(await ev("JSON.stringify([innerWidth, innerHeight, devicePixelRatio])"))
                info["theme"] = await ev("document.documentElement.getAttribute('data-theme') || 'light'")
                R["aufnahmen"][name] = info
                print(name, info["methode"], info["viewport"], flush=True)

            async def go(path, wait=3.5):
                await ev(f"location.href = {json.dumps(BASE + path)}"); await asyncio.sleep(wait)

            async def theme(dark):
                await ev(f"localStorage.setItem('lexono_theme', '{'dark' if dark else 'light'}')")

            async def open_popup(conv, docs):
                await go(f"/dashboard/chat/{conv}")
                await ev("(() => { const b = document.querySelector('.chat-message [data-pseudonymization-center-open]'); if (b) b.click(); else window.LexonoPseudonymizationCenter.open({}); })()")
                return await ev(WAIT_READY % docs)

            async def click(sel):
                return await ev(f"(() => {{ const b = document.querySelector({json.dumps(sel)}); if (!b) return 'fehlt'; if (b.disabled) return 'gesperrt'; b.click(); return 'ok'; }})()")

            async def answer(dec):
                r = await ev(ANSWER % json.dumps(dec)); await asyncio.sleep(1.4); return r

            await cdp("Runtime.enable"); await cdp("Page.enable")
            await cdp("Page.addScriptToEvaluateOnNewDocument", source="window.__errs = []; addEventListener('error', e => window.__errs.push(String(e.message)));")
            await asyncio.sleep(2)
            R["funktion"]["start_url"] = await ev("location.href")
            await ev(f"(() => {{ const f = document.querySelector('form'); f.querySelector('input[name=email]').value = {json.dumps(ADMIN)}; f.querySelector('input[name=password]').value = {json.dumps(ADMIN_PW)}; f.requestSubmit ? f.requestSubmit() : f.submit(); }})()")
            await asyncio.sleep(4)
            R["funktion"]["nach_login"] = await ev("location.pathname")

            # Standardfenster wie beim Start (maximiert, falls 1400x900 nicht passt)
            R["funktion"]["standard_viewport"] = json.loads(await ev("JSON.stringify([innerWidth, innerHeight, devicePixelRatio])"))
            R["funktion"]["standard_zoomed"] = bool(user32.IsZoomed(hwnd))
            await theme(False)
            await open_popup(convs["main"], 1)
            await shot("00-maximiert-hell-erste-rueckfrage")
            await click(".pseudo-center__footer [data-pseudo-center-cancel]")

            sizes = [(1920, 1080), (1280, 720), (960, 600)]
            for dark in (False, True):
                tn = "dunkel" if dark else "hell"
                await theme(dark)
                for w, h in sizes:
                    p = f"{w}x{h}-{tn}"
                    R["funktion"].setdefault("fenstergroessen", {})[p] = await resize(w, h)
                    await open_popup(convs["main"], 1)
                    await shot(f"{p}-01-erste-rueckfrage")
                    if PHASE == "nachher":
                        await answer("personal")
                        await shot(f"{p}-02-rueckfrage-2-von-5")
                    await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(0.8)
                    await open_popup(convs["multi"], 3)
                    await shot(f"{p}-08-mehrere-dokumente")
                    await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(0.8)
                    await open_popup(convs["single"], 1)
                    await shot(f"{p}-10-einzelne-rueckfrage")
                    await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(0.8)

            await theme(False)
            await resize(1280, 720)
            if os.environ.get("QA_SHOTS_ONLY") == "1":
                raise SystemExit(0)
            if PHASE == "nachher":
                # ---- Ablauf im Dokument-Pop-up: alle drei Optionen, Navigation, Ansichten, Dokumente, Abbrechen
                await open_popup(convs["multi"], 3)
                await answer("personal")                                   # Richter (auch in der DOCX)
                await shot("1280x720-hell-20-option-pseudonymisieren-uebernommen")
                await ev(ANSWER.replace("if (a && ", "if (false && ") % json.dumps("not_personal"))  # nur waehlen
                await asyncio.sleep(1.2)
                await shot("1280x720-hell-21-option-nicht-pseudonymisieren-gewaehlt")
                await click("[data-pseudo-q-apply]"); await asyncio.sleep(1.4)
                await ev(ANSWER.replace("if (a && ", "if (false && ") % json.dumps("unsure")); await asyncio.sleep(1.2)
                await shot("1280x720-hell-22-option-unsicher-gewaehlt")
                await click("[data-pseudo-q-apply]"); await asyncio.sleep(1.4)
                await click("[data-pseudo-q-prev]"); await asyncio.sleep(1.4)
                await shot("1280x720-hell-23-vorherige-rueckfrage-auswahl-erhalten")
                await click("[data-pseudo-q-next]"); await asyncio.sleep(1.4)
                await shot("1280x720-hell-24-naechste-offene-rueckfrage")
                await ev("(() => { const t = [...document.querySelectorAll('.pseudo-center__doc-tab')].find(x => x.textContent.includes('.docx')); if (t) t.click(); })()")
                await asyncio.sleep(2.5)
                await shot("1280x720-hell-25-dokumentwechsel-docx")
                await click("[data-pseudo-view='original']"); await asyncio.sleep(1.5)
                await shot("1280x720-hell-26-ansicht-original")
                await click("[data-pseudo-view='result']"); await asyncio.sleep(1.2)
                await shot("1280x720-hell-27-ansicht-ergebnis")
                await click("[data-pseudo-view='marked']"); await asyncio.sleep(1.2)
                await ev("(() => { const t = [...document.querySelectorAll('.pseudo-center__doc-tab')].find(x => x.textContent.includes('Bauer')); if (t) t.click(); })()")
                await asyncio.sleep(2.0)
                await answer("personal"); await answer("personal")
                await shot("1280x720-hell-28-letzte-rueckfrage-beantwortet-bestaetigen-aktiv")
                before = db_counts(); rec_before = len(recorder_requests())
                await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(1.5)
                await shot("1280x720-hell-29-abgebrochen-zurueck-im-chat")
                R["funktion"]["abbrechen_dokument"] = {"vorher": before, "nachher": db_counts(),
                                                       "recorder_vorher": rec_before, "recorder_nachher": len(recorder_requests()),
                                                       "dialog_offen": await ev("document.getElementById('pseudonymization-center').open")}

            # ---- Echte Rueckfrage beim Senden aus dem Chat (Server-review-Ereignis) - Abbrechen
            await go("/dashboard/chat")
            before = db_counts(); rec_before = len(recorder_requests())
            await ev(f"(() => {{ const t = document.querySelector('textarea[name=\"content\"]'); t.value = {json.dumps(CHAT_FRAGE)}; t.dispatchEvent(new Event('input', {{bubbles: true}})); t.closest('form').requestSubmit(); }})()")
            for _ in range(40):
                await asyncio.sleep(1)
                if await ev("(() => { const d = document.getElementById('pseudonymization-center'); return !!(d && d.open && d.querySelector('[data-question]')); })()"):
                    break
            await asyncio.sleep(2.5)
            await shot("1280x720-hell-30-chat-rueckfrage-beim-senden")
            R["funktion"]["chat_rueckfrage"] = {"fragen": await ev("[...document.querySelectorAll('#pseudonymization-center [data-question]')].map(q => q.dataset.question)"),
                                                "recorder_vor_entscheidung": len(recorder_requests()) - rec_before,
                                                "db_vor_senden": before, "db_mit_offener_rueckfrage": db_counts()}
            before = db_counts()  # Abbrechen allein darf nichts aendern
            await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(1.5)
            await shot("1280x720-hell-31-chat-rueckfrage-abgebrochen")
            R["funktion"]["chat_rueckfrage"]["abbrechen"] = {"vorher": before, "nachher": db_counts(),
                                                             "recorder": len(recorder_requests()) - rec_before,
                                                             "erneut_oeffnen_button": await ev("!!document.querySelector('.privacy-review-status__open')")}
            # ---- erneut oeffnen, beantworten, bestaetigen -> Fortsetzung ueber den Test-Double
            if await click(".privacy-review-status__open") == "ok":
                for _ in range(20):
                    await asyncio.sleep(1)
                    if await ev("!!document.querySelector('#pseudonymization-center[open] [data-question]')"):
                        break
                await asyncio.sleep(1.5)
                await answer("personal")
                await shot("1280x720-hell-32-chat-rueckfrage-beantwortet")
                before = db_counts(); rec_before = len(recorder_requests())
                R["funktion"]["chat_rueckfrage"]["bestaetigen_klick"] = await click("[data-pseudo-center-confirm]")
                for _ in range(60):
                    await asyncio.sleep(1)
                    if len(recorder_requests()) > rec_before and await ev("document.readyState") == "complete":
                        break
                await asyncio.sleep(6)
                await shot("1280x720-hell-33-chat-fortgesetzt-antwort")
                reqs = recorder_requests()[rec_before:]
                bodies = [json.dumps(r["body"], ensure_ascii=False) for r in reqs]
                R["funktion"]["chat_rueckfrage"]["bestaetigen"] = {
                    "vorher": before, "nachher": db_counts(), "recorder_anfragen": len(reqs),
                    "pfade": sorted({r["path"] for r in reqs}), "richter_im_payload": sum("Richter" in b for b in bodies),
                    "platzhalter": sorted({p for b in bodies for p in re.findall(r"\[[A-Z]+_\d{2,}\]", b)}),
                    "antwort_sichtbar": await ev("document.body.innerText.includes('QA-Recorder')"),
                    "pfad": await ev("location.pathname")}

            if PHASE == "nachher":
                # ---- Dokument-Gate: alle Pflichtfragen beantworten, Bestaetigen, Auftrag fortsetzen
                await open_popup(convs["main"], 1)
                for dec in ("personal", "not_personal", "unsure", "personal", "personal"):
                    await answer(dec)
                await shot("1280x720-hell-40-alle-pflichtfragen-beantwortet")
                before = db_counts(); rec_before = len(recorder_requests())
                R["funktion"]["dokument_bestaetigen_klick"] = await click("[data-pseudo-center-confirm]")
                await asyncio.sleep(1.0)
                await shot("1280x720-hell-41-bestaetigung-laeuft")
                for _ in range(90):
                    await asyncio.sleep(1)
                    if len(recorder_requests()) > rec_before and await ev("document.readyState") == "complete":
                        break
                await asyncio.sleep(8)
                await shot("1280x720-hell-42-auftrag-fortgesetzt-antwort")
                reqs = recorder_requests()[rec_before:]
                bodies = [json.dumps(r["body"], ensure_ascii=False) for r in reqs]
                R["funktion"]["dokument_bestaetigen"] = {
                    "vorher": before, "nachher": db_counts(), "recorder_anfragen": len(reqs),
                    "pfade": sorted({r["path"] for r in reqs}),
                    "lecks": {o: sum(o in b for b in bodies) for o in DOC_ORIGINALS},
                    "kaufmann_nicht_pseudonymisiert_im_payload": sum("Kaufmann" in b for b in bodies),
                    "platzhalter": sorted({p for b in bodies for p in re.findall(r"\[[A-Z]+_\d{2,}\]", b)}),
                    "antwort_sichtbar": await ev("document.body.innerText.includes('QA-Recorder')"),
                    "pfad": await ev("location.pathname")}
            R["funktion"]["js_fehler"] = await ev("window.__errs")

        # ---- Serverseitige Pruefungen (HTTP, gleiche Instanz) an der noch offenen Mehrdokument-Unterhaltung
        if PHASE == "nachher":
            conv = convs["multi"]
            docs = json.loads(get(f"/dashboard/chat/{conv}/pseudonymization-documents"))["documents"]
            rev = json.loads(get(f"/dashboard/chat/{conv}/document/{docs[0]['id']}/pseudonymization-review"))

            def confirm(fields, o=None):
                page = get(f"/dashboard/chat/{conv}")
                base = {"csrf_token": token("csrf_token", page), "document_id": docs[0]["id"],
                        "text_version": rev["text_version"], "documents": "[]", "decisions": "[]", "require_resolved": "1"}
                base.update(fields)
                b = db_counts(); r0 = len(recorder_requests())
                try:
                    out = sse(post(f"/dashboard/chat/{conv}/confirm-release-and-generate", base, o=o))
                    status = 200
                except urllib.error.HTTPError as e:
                    out, status = [], e.code
                return {"http": status, "antwort": [(e.get("kind"), (e.get("content") or "")[:140]) for e in out],
                        "db_unveraendert": b == db_counts(), "recorder_unveraendert": r0 == len(recorder_requests())}

            sv = {"offene_rueckfragen_unbeantwortet": confirm({}),
                  "veralteter_textstand": confirm({"text_version": "veraltet-qa",
                                                   "decisions": json.dumps([{"term": c["term"], "decision": "personal"} for c in rev["candidates"]])}),
                  "fremder_begriff": confirm({"decisions": json.dumps([{"term": "Gibtesnicht", "decision": "personal"}])}),
                  "fremdes_dokument": confirm({"document_id": "00000000-0000-0000-0000-000000000000"})}
            staff = session(); login(staff, STAFF, STAFF_PW)
            page = get("/dashboard/chat", o=staff)
            sv["rolle_mitarbeiter"] = {}
            try:
                post(f"/dashboard/chat/{conv}/confirm-release-and-generate",
                     {"csrf_token": token("csrf_token", page), "document_id": docs[0]["id"], "text_version": rev["text_version"],
                      "decisions": "[]", "require_resolved": "1"}, o=staff)
                sv["rolle_mitarbeiter"]["http"] = 200
            except urllib.error.HTTPError as e:
                sv["rolle_mitarbeiter"]["http"] = e.code
            sv["rolle_mitarbeiter"]["db_nachher"] = db_counts()
            R["funktion"]["serverpruefungen"] = sv
        R["funktion"]["db_ende"] = db_counts()
        R["funktion"]["recorder_gesamt"] = len(recorder_requests())
        R["funktion"]["recorder_schluessel"] = sorted({r["api_key_prefix"] for r in recorder_requests()})
    finally:
        proc.terminate()
        try:
            proc.wait(15)
        except Exception:
            proc.kill()
        subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)], capture_output=True)
        rec.terminate()


try:
    asyncio.run(main())
except SystemExit:
    pass
(OUT / f"native_{PHASE}.json").write_text(json.dumps(R, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(R["funktion"], ensure_ascii=False, indent=1)[:6000])
