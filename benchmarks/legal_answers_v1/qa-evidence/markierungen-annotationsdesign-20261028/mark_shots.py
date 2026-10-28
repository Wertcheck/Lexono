"""Markierungs-QA des Pseudonymisierungs-Pop-ups (Quellstand, run.py serve --no-window, isolierte Instanz,
ausschliesslich synthetische Dokumente, toter Cloud-Endpunkt, headless Edge per CDP, DPR 1, mit Bildlaufleisten).
Aufnahmen: ganzes Fenster + Nahaufnahme (Ausschnitt der ersten Seite) bei Zoom 100 % und 160 %, hell/dunkel, Hover,
Klick auf eine Fundstelle (Panel folgt), Rueckfrage-Navigation (Dokument folgt).
Usage: mark_shots.py <out_dir> <phase>"""
import asyncio, base64, http.cookiejar, io, json, os, re, subprocess, sys, time, urllib.parse, urllib.request, uuid
from pathlib import Path

import websockets

T = Path(__file__).resolve().parent
REPO = Path(os.environ.get("QA_REPO", r"<REPOSITORY>"))
VENV_PY = r"<REPOSITORY>\.venv\Scripts\python.exe"
OUT = Path(sys.argv[1]); OUT.mkdir(parents=True, exist_ok=True)
PHASE = sys.argv[2]
PY = VENV_PY
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
INST, PORT, CDP = T / "mark_qa", 8077, 9364
BASE = f"http://127.0.0.1:{PORT}"
R = {"phase": PHASE, "aufnahmen": {}, "funktion": {}}

subprocess.run([PY, str(T / "pseudo_seed.py"), str(INST), str(PORT)], check=True, capture_output=True)
env = dict(os.environ, LEXONO_DATA_DIR=str(INST), PYTHONDONTWRITEBYTECODE="1",
           HTTPS_PROXY="http://127.0.0.1:9", HTTP_PROXY="http://127.0.0.1:9", NO_PROXY="127.0.0.1,localhost")
for k in ("ANTHROPIC_API_KEY", "ANTHROPIC_BASE_URL", "DATABASE_URL"):
    env.pop(k, None)
server = subprocess.Popen([PY, str(REPO / "run.py"), "serve", "--no-window"], cwd=str(INST), env=env,
                          stdout=open(INST / f"serve_{PHASE}.out", "w"), stderr=subprocess.STDOUT)
for _ in range(120):
    try:
        urllib.request.urlopen(BASE + "/dashboard/login", timeout=2); break
    except Exception:
        time.sleep(1)

cj = http.cookiejar.CookieJar()
op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj), urllib.request.ProxyHandler({}))


def get(path):
    return op.open(BASE + path, timeout=300).read().decode("utf-8", "replace")


def post(path, fields, files=None):
    boundary = "----QA" + uuid.uuid4().hex
    body = b""
    for k, v in fields.items():
        body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode()
    for k, (name, data, ctype) in (files or {}).items():
        body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"; filename=\"{name}\"\r\n"
                 f"Content-Type: {ctype}\r\n\r\n").encode() + data + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    req = urllib.request.Request(BASE + path, data=body, headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    return op.open(req, timeout=900).read().decode("utf-8", "replace")


def sse(text):
    return [json.loads(l[6:]) for l in text.splitlines() if l.startswith("data: ")]


def token(name, page):
    m = re.search(rf'name="{name}" value="([^"]+)"', page)
    return m.group(1) if m else ""


def mark_pdf():
    """Synthetisches Schreiben: mehrteilige Namen, mehrere Funde in einer Zeile, Name ueber einen Zeilenumbruch,
    Anschrift, Organisation, Datum/E-Mail, unsichere Begriffe (Rueckfragen) - auch mehrfach."""
    import pymupdf
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((60, 40), "Kanzlei Muster - Kopfzeile (synthetisch)", fontsize=9)
    lines = [
        ("Einspruch gegen den Bescheid vom 02.02.2024 (synthetisches Musterdokument)", 13),
        ("", 11),
        ("Mandant: Max Bauer, Lindenweg 4, 50667 Koeln", 11),
        ("Gegnerin: Nordbauholding GmbH, vertreten durch Sabine Hartmann", 11),
        ("Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024.", 11),
        ("Das Schreiben von Kaufmann und das Schreiben von Richter liegen vor.", 11),
        ("Rueckfragen bitte an max.bauer@example.org oder an Max Bauer direkt.", 11),
    ]
    y = 90
    for text, size in lines:
        if text:
            p.insert_text((60, y), text, fontsize=size)
        y += size * 1.9
    # Fliesstext mit Umbruch innerhalb eines mehrteiligen Namens
    box = pymupdf.Rect(60, y, 330, y + 120)
    p.insert_textbox(box, "Nach Ruecksprache mit dem Geschaeftsfuehrer Thomas Berger hat die Nordbauholding GmbH "
                          "die Unterlagen am 18.03.2024 an Max Bauer uebergeben.", fontsize=11, lineheight=1.5)
    # Name ueber einen Zeilenumbruch (zwei Textzeilen): "... von Frau Sabine" / "Hartmann am ..."
    y2 = y + 110
    p.insert_text((60, y2), "Die Vollmacht wurde unterzeichnet von Frau Sabine", fontsize=11)
    p.insert_text((60, y2 + 20), "Hartmann am 20.03.2024 in Koeln.", fontsize=11)
    data = doc.tobytes(); doc.close(); return data


def table_docx():
    import docx
    d = docx.Document()
    d.add_heading("Sachstand (synthetisch)", level=1)
    d.add_paragraph("Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024.")
    t = d.add_table(rows=3, cols=2); t.style = "Table Grid"
    for r, (a, b) in enumerate([("Ansprechpartner", "Thomas Berger"), ("Anschrift", "Beispielstrasse 17, 10115 Berlin"),
                                ("Gesellschaft", "Nordbauholding GmbH")]):
        t.cell(r, 0).text = a; t.cell(r, 1).text = b
    b = io.BytesIO(); d.save(b); return b.getvalue()


PDF, DOCX = "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def new_conv(files):
    conv = ""
    for name, data, ctype in files:
        page = get(f"/dashboard/chat/{conv}" if conv else "/dashboard/chat")
        ev = sse(post("/dashboard/chat/send-stream", {"upload_token": token("upload_token", page),
                                                      "csrf_token": token("csrf_token", page),
                                                      "conversation_id": conv, "content": ""},
                      files={"documents": (name, data, ctype)}))
        conv = conv or ev[0]["conversation_id"]
    sse(post("/dashboard/chat/send-stream", {"csrf_token": token("csrf_token", get(f"/dashboard/chat/{conv}")),
                                             "conversation_id": conv,
                                             "content": "Bitte erstelle einen Schriftsatz an das Finanzamt."}))
    return conv


MARKS = r"""(() => { const d = document.getElementById('pseudonymization-center');
 const marks = [...d.querySelectorAll('.pseudo-page .pseudo-mark')];
 const cls = m => m.className.replace('pseudo-mark ', '');
 const box = m => { const r = m.getBoundingClientRect(); return [Math.round(r.left), Math.round(r.top), Math.round(r.width), Math.round(r.height)]; };
 const active = marks.filter(m => m.classList.contains('is-active'));
 const current = marks.filter(m => m.classList.contains('is-current'));
 const st = m => { const s = getComputedStyle(m); return {bg: s.backgroundColor, border: s.borderTopWidth + ' ' + s.borderTopStyle + ' ' + s.borderBottomStyle, shadow: s.boxShadow.slice(0, 80), radius: s.borderTopLeftRadius}; };
 return {anzahl: marks.length, js_fehler: window.__errs || [],
   aktive_frage: (d.querySelector('.pseudo-question.is-active') || {dataset: {}}).dataset.question || null,
   fortschritt: (d.querySelector('[data-pseudo-progress-text]') || {}).textContent || null,
   fundstelle: (d.querySelector('.pseudo-question__occ') || {}).textContent || null,
   aktive_markierungen: active.map(m => m.dataset.term), aktuelle_markierung: current.map(m => m.dataset.term + '@' + m.closest('.pseudo-page').dataset.page),
   klassen: [...new Set(marks.map(cls))], beispiele: marks.slice(0, 40).map(m => ({term: m.dataset.term, cls: cls(m), box: box(m), title: m.title})),
   stil_beispiel: marks.length ? st(marks[0]) : null, zoom: (d.querySelector('[data-pseudo-zoom]') || {}).textContent}; })()"""

WAIT_READY = """new Promise(r => { const t0 = Date.now(); (function poll() { const d = document.getElementById('pseudonymization-center');
 const busy = d.querySelector('.pseudo-center__loading'); const done = d.querySelectorAll('.pseudo-center__doc-ok, .pseudo-center__doc-badge').length;
 if ((d.open && !busy && done >= %d && d.querySelector('.pseudo-page')) || Date.now() - t0 > 60000) { setTimeout(() => r(Date.now() - t0), 1800); } else { setTimeout(poll, 300); } })(); })"""


async def main():
    op.open(BASE + "/dashboard/login", timeout=30)
    op.open(urllib.request.Request(BASE + "/dashboard/login", data=urllib.parse.urlencode(
        {"email": "qa-pseudo@lexono.test", "password": "QA-pseudo-0000", "next": "/dashboard"}).encode()), timeout=30)
    convs = {"brief": new_conv([("QA-SYNTH-Markierungen.pdf", mark_pdf(), PDF)]),
             "tabelle": new_conv([("QA-SYNTH-Tabelle.docx", table_docx(), DOCX)])}
    for key, conv in convs.items():
        docs = json.loads(get(f"/dashboard/chat/{conv}/pseudonymization-documents"))["documents"]
        rev = json.loads(get(f"/dashboard/chat/{conv}/document/{docs[0]['id']}/pseudonymization-review"))
        R["funktion"].setdefault("server", {})[key] = {
            "rueckfragen": [c["term"] for c in rev["candidates"]],
            "funde": [(m["original"], m["category"]) for m in rev["mappings"]],
            "markierungen": [(m["term"], m["kind"], m["category"], round(m["y"], 4)) for p in rev["pages"] for m in p["marks"]]}
    cookies = [{"name": c.name, "value": c.value, "domain": "127.0.0.1", "path": "/"} for c in cj]
    edge = subprocess.Popen([EDGE, "--headless=new", f"--remote-debugging-port={CDP}",
                             f"--user-data-dir={T / ('edge_profile_marks_' + PHASE)}", "--no-first-run",
                             "--disable-extensions", "about:blank"])
    try:
        targets = []
        for _ in range(60):
            try:
                targets = json.loads(urllib.request.urlopen(f"http://127.0.0.1:{CDP}/json", timeout=2).read())
                if any(t["type"] == "page" for t in targets):
                    break
            except Exception:
                time.sleep(1)
        page = next(t for t in targets if t["type"] == "page")
        async with websockets.connect(page["webSocketDebuggerUrl"], max_size=200_000_000) as ws:
            n = 0

            async def cdp(method, **params):
                nonlocal n
                n += 1
                await ws.send(json.dumps({"id": n, "method": method, "params": params}))
                while True:
                    msg = json.loads(await ws.recv())
                    if msg.get("id") == n:
                        return msg.get("result", {})

            async def js(expr):
                res = await cdp("Runtime.evaluate", expression=expr, awaitPromise=True, returnByValue=True)
                if "exceptionDetails" in res:
                    return {"js_fehler": json.dumps(res["exceptionDetails"])[:300]}
                return res.get("result", {}).get("value")

            async def shot(name, crop=None):
                await asyncio.sleep(0.9)
                params = {"format": "png"}
                if crop:
                    params["clip"] = dict(crop, scale=1)
                data = base64.b64decode((await cdp("Page.captureScreenshot", **params))["data"])
                (OUT / f"{name}.png").write_bytes(data)
                R["aufnahmen"][name] = await js(MARKS)
                print(name, flush=True)

            async def page_crop(width_frac=0.75, height=380):
                r = await js("(() => { const p = document.querySelector('#pseudonymization-center .pseudo-page__box'); const pane = document.querySelector('[data-pseudo-document]').getBoundingClientRect(); const b = p.getBoundingClientRect(); return [b.left, Math.max(b.top, pane.top), b.width, pane.bottom - Math.max(b.top, pane.top)]; })()")
                x, y, w, h = r
                return {"x": x, "y": y, "width": w * width_frac, "height": min(h, height)}

            async def size(w, h):
                await cdp("Emulation.setDeviceMetricsOverride", width=w, height=h, deviceScaleFactor=1, mobile=False)
                await asyncio.sleep(0.4)

            async def open_popup(conv):
                await cdp("Page.navigate", url=f"{BASE}/dashboard/chat/{conv}")
                await asyncio.sleep(2.5)
                await js("(() => { const b = document.querySelector('.chat-message [data-pseudonymization-center-open]'); if (b) b.click(); else window.LexonoPseudonymizationCenter.open({}); })()")
                await js(WAIT_READY % 1)
                await js("document.querySelector('[data-pseudo-document]').scrollTo(0, 0)")
                await asyncio.sleep(0.6)

            async def click(sel):
                return await js(f"(() => {{ const b = document.querySelector({json.dumps(sel)}); if (!b) return 'fehlt'; b.click(); return 'ok'; }})()")

            await cdp("Page.enable"); await cdp("Network.enable"); await cdp("Runtime.enable")
            await cdp("Page.addScriptToEvaluateOnNewDocument", source="window.__errs = []; addEventListener('error', e => window.__errs.push(String(e.message)));")
            for c in cookies:
                await cdp("Network.setCookie", **c)
            await cdp("Page.navigate", url=BASE + "/dashboard"); await asyncio.sleep(2)

            for dark in (False, True):
                tn = "dunkel" if dark else "hell"
                await js(f"localStorage.setItem('lexono_theme', '{'dark' if dark else 'light'}')")
                for w, h in ((1920, 1080), (1280, 720)):
                    p = f"{w}x{h}-{tn}"
                    await size(w, h)
                    await open_popup(convs["brief"])
                    await shot(f"{p}-01-zoom100-fenster")
                    await shot(f"{p}-02-zoom100-nah", await page_crop())
                    for _ in range(3):
                        await click("[data-pseudo-zoom-in]"); await asyncio.sleep(0.6)
                    await asyncio.sleep(1.2)
                    await js("document.querySelector('[data-pseudo-document]').scrollTo(0, 0)")
                    await shot(f"{p}-03-zoom160-nah", await page_crop(0.6, 420))
                    await click("[data-pseudo-zoom-fit]"); await asyncio.sleep(1.2)
                    if w == 1280:
                        await click("[data-pseudo-zoom-out]"); await asyncio.sleep(1.2)
                        await js("document.querySelector('[data-pseudo-document]').scrollTo(0, 0)")
                        await shot(f"{p}-03b-zoom80-nah", await page_crop(0.9, 420))
                        await click("[data-pseudo-zoom-fit]"); await asyncio.sleep(1.2)
                    if w == 1280 and not dark:
                        # Hover ueber einer normalen (maskierten) Fundstelle
                        pos = await js("(() => { const m = [...document.querySelectorAll('.pseudo-page .pseudo-mark')].find(x => x.dataset.term === 'Max Bauer'); const r = m.getBoundingClientRect(); return [r.left + r.width / 2, r.top + r.height / 2]; })()")
                        await cdp("Input.dispatchMouseEvent", type="mouseMoved", x=pos[0], y=pos[1])
                        await shot(f"{p}-04-hover-fundstelle-nah", await page_crop())
                        await cdp("Input.dispatchMouseEvent", type="mouseMoved", x=5, y=5)
                        # Klick auf die ZWEITE Fundstelle von "Richter" im Dokument -> Panel folgt
                        R["funktion"]["klick_zweite_fundstelle"] = await js("(() => { const ms = [...document.querySelectorAll('.pseudo-page .pseudo-mark')].filter(x => x.dataset.term === 'Richter'); if (ms.length < 2) return 'nur ' + ms.length; ms[1].click(); return ms.length; })()")
                        await asyncio.sleep(1.5)
                        await shot(f"{p}-05-klick-fundstelle-panel-folgt")
                        # Klick auf eine andere Rueckfrage im Dokument (Kaufmann)
                        await js("(() => { const m = [...document.querySelectorAll('.pseudo-page .pseudo-mark')].find(x => x.dataset.term === 'Kaufmann'); if (m) m.click(); })()")
                        await asyncio.sleep(1.5)
                        await shot(f"{p}-06-klick-andere-rueckfrage")
                        # Panel -> Dokument: naechste Rueckfrage
                        await click("[data-pseudo-q-next]"); await asyncio.sleep(1.5)
                        await shot(f"{p}-07-panel-naechste-rueckfrage-dokument-folgt")
                        # Entscheidungen: pseudonymisieren / nicht pseudonymisieren -> Markierungsart
                        await js("(() => { const b = document.querySelector('.pseudo-question.is-active [data-decision=\"not_personal\"]'); if (b) b.click(); })()")
                        await asyncio.sleep(1.0)
                        await shot(f"{p}-08-entscheidung-bleibt-lesbar-nah", await page_crop())
                    await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(0.8)
                    await open_popup(convs["tabelle"])
                    await shot(f"{p}-09-docx-tabelle-nah", await page_crop())
                    await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(0.8)
            await js("localStorage.setItem('lexono_theme', 'light')")
    finally:
        edge.terminate()
        subprocess.run(["taskkill", "/F", "/T", "/PID", str(server.pid)], capture_output=True)


asyncio.run(main())
(OUT / f"messwerte_{PHASE}.json").write_text(json.dumps(R, ensure_ascii=False, indent=1), encoding="utf-8")
print("fertig", len(R["aufnahmen"]))
