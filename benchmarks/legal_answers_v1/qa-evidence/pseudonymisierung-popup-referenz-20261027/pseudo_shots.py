"""Echte Screenshots des Pseudonymisierungs-Pop-ups in der laufenden App (Quellstand, run.py serve --no-window).
Isolierte Instanz (pseudo_seed.py), ausschliesslich synthetische Dokumente, toter Cloud-Endpunkt, headless Edge per
CDP mit exakten Viewports (DPR 1), hell/dunkel. Usage: pseudo_shots.py <out_dir> <phase> [--quick]"""
import asyncio
import base64
import http.cookiejar
import io
import json
import os
import re
import sqlite3
import subprocess
import sys
import time
import urllib.parse
import urllib.request
import uuid
from pathlib import Path

import websockets

T = Path(os.environ["LEXONO_QA_WORKDIR"])  # isolierter QA-Arbeitsordner (enthaelt pseudo_seed.py)
REPO = Path(__file__).resolve().parents[4]  # Repository-Wurzel
OUT = Path(sys.argv[1]); OUT.mkdir(parents=True, exist_ok=True)
PHASE = sys.argv[2]
QUICK = "--quick" in sys.argv
PY = str(REPO / ".venv" / "Scripts" / "python.exe")
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
INST, PORT, CDP = T / "pseudo_qa", 8075, 9362
BASE = f"http://127.0.0.1:{PORT}"
R = {"phase": PHASE, "aufnahmen": {}, "funktion": {}}

subprocess.run([PY, str(T / "pseudo_seed.py"), str(INST), str(PORT)], check=True, capture_output=True)


def start():
    env = dict(os.environ, LEXONO_DATA_DIR=str(INST), PYTHONDONTWRITEBYTECODE="1",
               HTTPS_PROXY="http://127.0.0.1:9", HTTP_PROXY="http://127.0.0.1:9", NO_PROXY="127.0.0.1,localhost")
    for k in ("ANTHROPIC_API_KEY", "ANTHROPIC_BASE_URL", "DATABASE_URL"):
        env.pop(k, None)
    p = subprocess.Popen([PY, str(REPO / "run.py"), "serve", "--no-window"], cwd=str(INST), env=env,
                         stdout=open(INST / f"serve_{PHASE}.out", "w"), stderr=subprocess.STDOUT)
    for _ in range(120):
        try:
            urllib.request.urlopen(BASE + "/dashboard/login", timeout=2)
            return p
        except Exception:
            time.sleep(1)
    raise RuntimeError("Server startet nicht")


def stop(p):
    subprocess.run(["taskkill", "/F", "/T", "/PID", str(p.pid)], capture_output=True)


cj = http.cookiejar.CookieJar()
op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))


def get(path):
    return op.open(BASE + path, timeout=300).read().decode("utf-8", "replace")


def post(path, fields, files=None):
    boundary = "----QA" + uuid.uuid4().hex
    body = b""
    for k, v in fields.items():
        body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode()
    for k, (name, data, ctype) in (files or {}).items():
        body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"; filename=\"{name}\"\r\n"
                 f"Content-Type: {ctype}\r\n\r\n").encode() + data + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    req = urllib.request.Request(BASE + path, data=body,
                                 headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    return op.open(req, timeout=900).read().decode("utf-8", "replace")


def sse(text):
    return [json.loads(line[6:]) for line in text.splitlines() if line.startswith("data: ")]


def token(name, page):
    m = re.search(rf'name="{name}" value="([^"]+)"', page)
    return m.group(1) if m else ""


def pdf_bytes(pages):
    import pymupdf
    doc = pymupdf.open()
    for n, lines in enumerate(pages, 1):
        page = doc.new_page()
        page.insert_text((50, 30), "Kanzlei Muster - Kopfzeile (synthetisch)", fontsize=9)
        y = 90
        for line in lines:
            page.insert_text((50, y), line, fontsize=11)
            y += 20
        page.insert_text((50, 810), f"Seite {n} von {len(pages)}", fontsize=9)
    data = doc.tobytes(); doc.close(); return data


def docx_bytes():
    import docx
    d = docx.Document()
    d.add_heading("Stellungnahme (synthetisch)", level=1)
    d.add_paragraph("Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024.")
    d.add_paragraph("Ansprechpartner: Thomas Berger, Beispielstrasse 17, 10115 Berlin.")
    d.add_paragraph("Die Nordbauholding GmbH bittet um Aussetzung der Vollziehung.")
    b = io.BytesIO(); d.save(b); return b.getvalue()


MAIN_PAGES = [
    ["Einspruch gegen den Bescheid vom 02.02.2024 (synthetisches Musterdokument)",
     "Mandant: Max Bauer, Lindenweg 4, 50667 Koeln",
     "Hiermit legen wir fuer die Nordbauholding GmbH Einspruch ein.",
     "Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024.",
     "Wir bestaetigen den Eingang des Schreibens von Kaufmann vom 18.03.2024."],
    ["Begruendung (Fortsetzung)",
     "Wir bestaetigen den Eingang des Schreibens von Meister vom 20.03.2024.",
     "Rueckfragen bitte an max.bauer@example.org."],
    ["Anlagen (Fortsetzung)",
     "Wir bestaetigen den Eingang des Schreibens von Gärtner vom 22.03.2024.",
     "Wir bestaetigen den Eingang des Schreibens von Kellner vom 24.03.2024."],
]
SINGLE_PAGES = [["Vermerk zur Akte (synthetisches Musterdokument)", "Mandantin: Sabine Hartmann, Lindenweg 4, 50667 Koeln",
                 "Wir bestaetigen den Eingang des Schreibens von Richter vom 16.03.2024."]]
ANLAGE_PAGES = [["Anlage K1 (synthetisch)", "Die Nordbauholding GmbH bittet um Aussetzung der Vollziehung."]]
AUFTRAG = "Bitte erstelle einen Schriftsatz an das Finanzamt mit dem aktuellen Sachstand."
PDF, DOCX = "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def new_conv(files, auftrag=AUFTRAG):
    conv = ""
    for name, data, ctype in files:
        page = get(f"/dashboard/chat/{conv}" if conv else "/dashboard/chat")
        ev = sse(post("/dashboard/chat/send-stream", {"upload_token": token("upload_token", page),
                                                      "csrf_token": token("csrf_token", page),
                                                      "conversation_id": conv, "content": ""},
                      files={"documents": (name, data, ctype)}))
        conv = conv or ev[0]["conversation_id"]
    if auftrag:
        ev = sse(post("/dashboard/chat/send-stream", {"csrf_token": token("csrf_token", get(f"/dashboard/chat/{conv}")),
                                                      "conversation_id": conv, "content": auftrag}))
        R["funktion"].setdefault("auftrag_blockiert", {})[conv] = next(e for e in ev if e["kind"] == "done").get("blocked")
    return conv


def db_counts():
    c = sqlite3.connect(f"file:{INST / 'data' / 'kanzlei_ai.db'}?mode=ro", uri=True)
    out = {t: c.execute(f"select count(*) from {t}").fetchone()[0]
           for t in ("privacy_decisions", "document_releases", "chat_messages")}
    c.close()
    return out


STATE = r"""(() => { const d = document.getElementById('pseudonymization-center');
 const vis = e => { if (!e) return false; const r = e.getBoundingClientRect(); const s = getComputedStyle(e);
   return r.width > 0 && r.height > 0 && s.visibility !== 'hidden' && !e.closest('[hidden]'); };
 const inView = e => { if (!vis(e)) return false; const r = e.getBoundingClientRect(); return r.top >= -0.5 && r.left >= -0.5 && r.bottom <= innerHeight + 0.5 && r.right <= innerWidth + 0.5; };
 const box = e => { if (!e) return null; const r = e.getBoundingClientRect(); return [Math.round(r.left), Math.round(r.top), Math.round(r.width), Math.round(r.height)]; };
 const panel = d.querySelector('[data-pseudo-panel]'); const ps = d.querySelector('[data-pseudo-panel-scroll]');
 const confirm = d.querySelector('[data-pseudo-center-confirm]'); const apply = d.querySelector('[data-pseudo-q-apply]');
 const primary = [confirm, apply].filter(vis)[0] || null;
 const overflow = [...d.querySelectorAll('.pseudo-center__header, .pseudo-center__docbar, .pseudo-center__toolbar, .pseudo-center__footer, .pseudo-center__panel-head, .pseudo-question')]
   .filter(e => vis(e) && e.scrollWidth > e.clientWidth + 1).map(e => e.className.split(' ')[0]);
 return {dialog_offen: d.open, fenster: [innerWidth, innerHeight], dialog: box(d), dialog_im_fenster: inView(d),
  seite_querscroll: document.documentElement.scrollWidth > innerWidth, ueberlauf: overflow,
  fusszeile_sichtbar: inView(d.querySelector('.pseudo-center__footer')),
  abbrechen_sichtbar: inView(d.querySelector('.pseudo-center__footer [data-pseudo-center-cancel]')),
  primaer_text: primary ? primary.textContent.trim() : null, primaer_sichtbar: inView(primary), primaer_gesperrt: primary ? primary.disabled : null,
  bestaetigen_sichtbar: vis(confirm), bestaetigen_gesperrt: confirm.disabled,
  sichtbare_fragekarten: [...d.querySelectorAll('[data-question]')].filter(vis).length,
  aktive_frage: (d.querySelector('.pseudo-question.is-active') || {dataset: {}}).dataset.question || null,
  optionen_sichtbar: (() => { const r = (ps || d).getBoundingClientRect(); return [...d.querySelectorAll('.pseudo-question.is-active [data-decision]')].filter(o => { const x = o.getBoundingClientRect(); return vis(o) && x.top >= r.top - 1 && x.bottom <= r.bottom + 1; }).length; })(),
  gewaehlt: (d.querySelector('.pseudo-question.is-active [data-decision][aria-pressed="true"], .pseudo-question.is-active [data-decision][aria-checked="true"]') || {dataset: {}}).dataset.decision || null,
  fortschritt: (d.querySelector('[data-pseudo-progress-text]') || {}).textContent || null,
  panel_intern_scroll: ps ? [ps.scrollHeight, ps.clientHeight] : null, panel: box(panel),
  vorschau: box(d.querySelector('[data-pseudo-document]')), seite: box(d.querySelector('.pseudo-page')),
  miniaturen_sichtbar: [...d.querySelectorAll('.pseudo-thumb')].filter(vis).length,
  dok_tabs: [...d.querySelectorAll('.pseudo-center__doc-tab')].filter(vis).map(t => t.textContent.trim()),
  ansicht: (d.querySelector('[data-pseudo-view].is-active') || {}).textContent || null,
  seite_info: (d.querySelector('[data-pseudo-page-info]') || {}).textContent || null,
  legende: (d.querySelector('.pseudo-center__legend') || {}).textContent ? d.querySelector('.pseudo-center__legend').textContent.replace(/\s+/g, ' ').trim() : null,
  markierungen: d.querySelectorAll('.pseudo-mark').length,
  fusszeile_text: (d.querySelector('[data-pseudo-open-count]') || {}).textContent || null, js_fehler: window.__errs || []}; })()"""

WAIT_READY = """new Promise(r => { const t0 = Date.now(); (function poll() { const d = document.getElementById('pseudonymization-center');
 const busy = d.querySelector('.pseudo-center__loading'); const done = d.querySelectorAll('.pseudo-center__doc-ok, .pseudo-center__doc-badge').length;
 if ((d.open && !busy && done >= %d && d.querySelector('.pseudo-page, .pseudo-text')) || Date.now() - t0 > 60000) { setTimeout(() => r(Date.now() - t0), 1500); } else { setTimeout(poll, 300); } })(); })"""

# Antwort waehlen und - falls vorhanden (neues Layout) - "Uebernehmen und naechste Rueckfrage" ausloesen
ANSWER = """((dec) => { const b = document.querySelector('#pseudonymization-center .pseudo-question.is-active [data-decision="' + dec + '"]');
 if (!b) return 'keine aktive Frage'; b.click(); const a = document.querySelector('[data-pseudo-q-apply]');
 if (a && !a.hidden && !a.disabled && !a.closest('[hidden]')) a.click(); return 'ok'; })(%s)"""


async def main():
    proc = start()
    op.open(BASE + "/dashboard/login", timeout=30)
    op.open(urllib.request.Request(BASE + "/dashboard/login", data=urllib.parse.urlencode(
        {"email": "qa-pseudo@lexono.test", "password": "QA-pseudo-0000", "next": "/dashboard"}).encode()), timeout=30)
    t0 = time.time()
    convs = {"main": new_conv([("QA-SYNTH-Einspruch-Bauer.pdf", pdf_bytes(MAIN_PAGES), PDF)]),
             "multi": new_conv([("QA-SYNTH-Einspruch-Bauer.pdf", pdf_bytes(MAIN_PAGES), PDF),
                                ("QA-SYNTH-Stellungnahme.docx", docx_bytes(), DOCX),
                                ("QA-SYNTH-Anlage-K1.pdf", pdf_bytes(ANLAGE_PAGES), PDF)]),
             "single": new_conv([("QA-SYNTH-Vermerk.pdf", pdf_bytes(SINGLE_PAGES), PDF)]),
             "view": new_conv([("QA-SYNTH-Nur-Ansicht.pdf", pdf_bytes(ANLAGE_PAGES), PDF)], auftrag=None)}
    R["funktion"]["setup_s"] = round(time.time() - t0, 1)
    R["funktion"]["convs"] = convs
    for key, conv in convs.items():
        docs = json.loads(get(f"/dashboard/chat/{conv}/pseudonymization-documents"))["documents"]
        R["funktion"].setdefault("rueckfragen", {})[key] = {
            d["filename"]: [c["term"] for c in json.loads(get(f"/dashboard/chat/{conv}/document/{d['id']}/pseudonymization-review"))["candidates"]]
            for d in docs}
    cookies = [{"name": c.name, "value": c.value, "domain": "127.0.0.1", "path": "/"} for c in cj]
    edge = subprocess.Popen([EDGE, "--headless=new", f"--remote-debugging-port={CDP}", "--hide-scrollbars",
                             f"--user-data-dir={T / ('edge_profile_pseudo_' + PHASE)}", "--no-first-run",
                             "--disable-extensions", "about:blank"])
    try:
        targets = []
        for _ in range(60):
            try:
                targets = json.loads(urllib.request.urlopen(f"http://127.0.0.1:{CDP}/json", timeout=2).read())
                if any(t["type"] == "page" for t in targets):
                    break
            except Exception:
                time.sleep(1)
        page = next(t for t in targets if t["type"] == "page")
        async with websockets.connect(page["webSocketDebuggerUrl"], max_size=200_000_000) as ws:
            n = 0

            async def cdp(method, **params):
                nonlocal n
                n += 1
                await ws.send(json.dumps({"id": n, "method": method, "params": params}))
                while True:
                    msg = json.loads(await ws.recv())
                    if msg.get("id") == n:
                        if "error" in msg:
                            raise RuntimeError(f"{method}: {msg['error']}")
                        return msg.get("result", {})

            async def js(expr):
                res = await cdp("Runtime.evaluate", expression=expr, awaitPromise=True, returnByValue=True)
                if "exceptionDetails" in res:
                    return {"js_fehler": json.dumps(res["exceptionDetails"])[:400]}
                return res.get("result", {}).get("value")

            async def shot(name):
                await asyncio.sleep(0.8)
                data = base64.b64decode((await cdp("Page.captureScreenshot", format="png"))["data"])
                (OUT / f"{name}.png").write_bytes(data)
                R["aufnahmen"][name] = await js(STATE)
                print(name, flush=True)

            async def size(w, h):
                await cdp("Emulation.setDeviceMetricsOverride", width=w, height=h, deviceScaleFactor=1, mobile=False)
                await asyncio.sleep(0.4)

            async def theme(dark):
                await js(f"localStorage.setItem('lexono_theme', '{'dark' if dark else 'light'}')")

            async def open_popup(conv, docs):
                await cdp("Page.navigate", url=f"{BASE}/dashboard/chat/{conv}")
                await asyncio.sleep(2.5)
                await js("(() => { const b = document.querySelector('.chat-message [data-pseudonymization-center-open]'); if (b) { b.click(); return 'sperrhinweis'; } window.LexonoPseudonymizationCenter.open({}); return 'api'; })()")
                await js(WAIT_READY % docs)

            async def click(sel):
                return await js(f"(() => {{ const b = document.querySelector({json.dumps(sel)}); if (!b) return 'fehlt'; b.click(); return 'ok'; }})()")

            async def answer(dec="personal"):
                r = await js(ANSWER % json.dumps(dec))
                await asyncio.sleep(1.2)
                return r

            await cdp("Page.enable"); await cdp("Network.enable"); await cdp("Runtime.enable")
            await cdp("Page.addScriptToEvaluateOnNewDocument", source="window.__errs = []; addEventListener('error', e => window.__errs.push(String(e.message)));")
            for c in cookies:
                await cdp("Network.setCookie", **c)
            await cdp("Page.navigate", url=BASE + "/dashboard")
            await asyncio.sleep(2)

            sizes = [(1920, 1080), (1280, 720), (960, 600)]
            for dark in (() if "--review-only" in sys.argv else (False, True)):
                await theme(dark)
                tn = "dunkel" if dark else "hell"
                for w, h in sizes:
                    p = f"{w}x{h}-{tn}"
                    full = (w, h) == (1280, 720) or (w, h) == (1920, 1080) and not QUICK
                    await size(w, h)
                    await open_popup(convs["main"], 1)
                    await shot(f"{p}-01-erste-rueckfrage")
                    await answer("personal")
                    await shot(f"{p}-02-rueckfrage-2-von-5")
                    if full:
                        await click("[data-pseudo-view='original']"); await asyncio.sleep(1.2)
                        await shot(f"{p}-03-ansicht-original")
                        await click("[data-pseudo-view='result']"); await asyncio.sleep(1.0)
                        await shot(f"{p}-04-ansicht-ergebnis")
                        await click("[data-pseudo-view='marked']"); await asyncio.sleep(1.2)
                        await shot(f"{p}-04b-zurueck-markiert-entscheidung-erhalten")
                    for dec in ("not_personal", "unsure"):
                        await answer(dec)
                    await answer("personal")
                    await shot(f"{p}-05-letzte-rueckfrage")
                    if full:
                        await click("[data-pseudo-q-prev]"); await asyncio.sleep(1.2)
                        await shot(f"{p}-05b-vorherige-beantwortete-rueckfrage")
                        await click("[data-pseudo-q-next]"); await asyncio.sleep(1.2)
                    await answer("personal")
                    await shot(f"{p}-06-alle-beantwortet")
                    before = db_counts()
                    await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(1.2)
                    if full:
                        await shot(f"{p}-07-abgebrochen-zurueck-im-chat")
                    R["funktion"].setdefault("abbruch_db", {})[p] = {"vorher": before, "nachher": db_counts()}
                    await open_popup(convs["multi"], 3)
                    await shot(f"{p}-08-mehrere-dokumente")
                    if full:
                        await js("(() => { const t = [...document.querySelectorAll('.pseudo-center__doc-tab')].find(x => x.textContent.includes('.docx')); if (t) t.click(); })()")
                        await asyncio.sleep(2.0)
                        await shot(f"{p}-09-dokumentwechsel-docx")
                    await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(0.8)
                    await open_popup(convs["single"], 1)
                    await shot(f"{p}-10-einzelne-rueckfrage")
                    await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(0.8)
            # Datenschutz-Rueckfrage beim Senden aus dem Chat (Modus "review") im selben Pop-up, dann Abbrechen
            await theme(False)
            await size(1280, 720)
            await cdp("Page.navigate", url=f"{BASE}/dashboard/chat/{convs['view']}")
            await asyncio.sleep(2.5)
            # Die isolierte Instanz hat keine Cloud-KI (toter Endpunkt) - der Server liefert daher kein review-Ereignis.
            # Geoeffnet wird dieselbe Browser-Schnittstelle, die chat.html beim review-Ereignis aufruft, mit einer
            # synthetischen Rueckfrage; Callbacks werden nur protokolliert (nichts wird gesendet).
            R["funktion"]["review_geoeffnet"] = await js("""(() => { window.__review = [];
              window.LexonoPseudonymizationCenter.openReview({conversation_id: location.pathname.split('/').pop(), message_id: 'qa-synth',
                candidates: [{term: 'Ruecksicht', context_before: 'Bitte mit ', context_after: ' auf die laufende Frist knapp zusammenfassen.'}]},
                d => window.__review.push(['entscheidungen', JSON.stringify(d)]), () => window.__review.push(['abgebrochen']));
              return true; })()""")
            await asyncio.sleep(4)
            await shot("1280x720-hell-11-rueckfrage-beim-senden-im-popup")
            await js(ANSWER % json.dumps("not_personal")); await asyncio.sleep(1.0)
            await shot("1280x720-hell-11b-rueckfrage-beim-senden-beantwortet")
            before = db_counts()
            await click(".pseudo-center__footer [data-pseudo-center-cancel]"); await asyncio.sleep(1.5)
            await shot("1280x720-hell-12-rueckfrage-beim-senden-abgebrochen")
            R["funktion"]["review_callbacks"] = await js("window.__review")
            R["funktion"]["review_abbruch_db"] = {"vorher": before, "nachher": db_counts(),
                "dialoge_offen": await js("[...document.querySelectorAll('dialog[open]')].map(d => d.id)")}
            R["funktion"]["db_ende"] = db_counts()
    finally:
        edge.terminate()
        stop(proc)


asyncio.run(main())
(OUT / f"messwerte_{PHASE}.json").write_text(json.dumps(R, ensure_ascii=False, indent=1), encoding="utf-8")
print("fertig", len(R["aufnahmen"]))
