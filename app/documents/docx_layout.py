"""Lokaler Seitensatz fuer DOCX (26.10.2026) - originalnaeher als der eingebaute DOCX-Import von PyMuPDF.

Gemessener Befund (25.10.2026): `pymupdf.open(<docx>)` setzt weder Kopf-/Fusszeilen noch Tabellenrahmen. Dieser Satz
liest das Dokument mit python-docx (bereits Abhaengigkeit der Textextraktion) und setzt es mit `pymupdf.Story` (HTML/CSS)
auf Seiten in der Groesse und mit den Raendern der ersten Abschnittseinstellung:

- Kopf- und Fusszeile (Text) auf JEDER Seite,
- Absaetze in Dokumentreihenfolge mit Ueberschriften, fett/kursiv/unterstrichen, Ausrichtung, Schriftgroesse,
  Aufzaehlungen (vereinfacht),
- Tabellen mit Rahmen, wenn die Tabelle Rahmen hat (Tabellenformatvorlage "Grid"/eigene Rahmen), sonst ohne.

Grenzen (bewusst, ehrlich gemeldet): Seitenumbrueche entstehen durch den Satz und koennen von Word abweichen; Bilder,
Textfelder, Fussnoten, Felder, mehrspaltiger Satz und exakte Schriftarten werden nicht nachgebildet. Lokal, kein
externer Dienst, keine neue Abhaengigkeit.

Seitenumbruch-Korrektur (26.10.2026, reproduziert mit synthetischen Briefvorlagen): `pymupdf.Story` teilt eine
TabellenZEILE nicht ueber Seiten. Eine Zeile, die nicht mehr auf die Seite passt, wurde komplett auf die Folgeseite
geschoben (Seite 1 nahezu leer) und alles, was auch dort nicht passte, ABGESCHNITTEN (15 von 32 Absaetzen einer
einzelligen Layouttabelle fehlten in der Seitenansicht). Jetzt:
- Zeilen, die hoeher als ein Teil der Seite sind, werden als umrandete Bloecke gesetzt (Bloecke teilt der Satz
  verlustfrei ueber Seiten); mehrspaltige Zeilen stehen dann untereinander - als Hinweis gemeldet (`layout_notes`).
- Manuelle Seitenumbrueche (`w:br w:type="page"`, "Seitenumbruch oberhalb", Abschnittswechsel) beginnen eine neue Seite.
- Ligaturen werden fuer die Textsuche aufgeloest (pseudonymization_review.py), damit "ff"/"fi" in Namen gefunden werden."""

from __future__ import annotations

import html
from functools import lru_cache
from pathlib import Path

import pymupdf

_DEFAULT_FONT_PT = 11.0
_ALIGN = {0: "left", 1: "center", 2: "right", 3: "justify"}
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _run_html(run, base_pt: float) -> str:
    # Zeilenumbruch innerhalb des Absatzes (python-docx liefert <w:br/> als "\n") und Tabulator sichtbar erhalten
    text = html.escape(run.text or "").replace("\n", "<br>").replace("\t", "&emsp;")
    if not text:
        return ""
    style = []
    if run.font is not None and run.font.size is not None:
        pt = run.font.size.pt
        if abs(pt - base_pt) > 0.4:
            style.append(f"font-size:{pt:.1f}pt")
    if run.bold:
        text = f"<b>{text}</b>"
    if run.italic:
        text = f"<i>{text}</i>"
    if run.underline:
        text = f"<u>{text}</u>"
    return f'<span style="{";".join(style)}">{text}</span>' if style else text


def _paragraph_html(paragraph, base_pt: float) -> str:
    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    inner = "".join(_run_html(run, base_pt) for run in paragraph.runs) or \
        html.escape(paragraph.text or "").replace("\n", "<br>")
    align = _ALIGN.get(int(paragraph.alignment)) if paragraph.alignment is not None else None
    css = f' style="text-align:{align}"' if align else ""
    if style_name.startswith("Heading") or style_name.startswith("Überschrift") or style_name == "Title":
        level = 1 if style_name == "Title" else (int(style_name.split()[-1]) if style_name.split()[-1].isdigit() else 2)
        level = min(max(level, 1), 4)
        return f"<h{level}{css}>{inner or '&nbsp;'}</h{level}>"
    if "List" in style_name or "Liste" in style_name:
        bullet = "1." if "Number" in style_name or "Nummer" in style_name else "•"
        return f'<p class="li"{css}>{bullet}&nbsp;{inner}</p>'
    return f"<p{css}>{inner or '&nbsp;'}</p>"


def _table_has_borders(table) -> bool:
    style_name = (table.style.name if table.style is not None else "") or ""
    if "Grid" in style_name or "Gitter" in style_name:
        return True
    tbl_pr = table._tbl.tblPr  # noqa: SLF001 - python-docx stellt Rahmen nicht als API bereit
    borders = tbl_pr.find(f"{_W_NS}tblBorders") if tbl_pr is not None else None
    if borders is None:
        return False
    return any((edge.get(f"{_W_NS}val") or "nil") not in ("nil", "none") for edge in borders)


#: Anteil der Satzspiegelhoehe, ab dem eine Tabellenzeile als teilbarer Block gesetzt wird. Einzellige Zeilen
#: (Layouttabellen von Briefvorlagen) frueh, mehrspaltige erst, wenn sie sonst eine halbe Seite Leerraum erzeugen.
_SPLIT_SINGLE_CELL = 0.25
_SPLIT_MULTI_CELL = 0.5
#: Messaufwand begrenzen (jede Zeile wird einmal probeweise gesetzt).
_MAX_MEASURED_ROWS = 400


class _Layout:
    """Satzparameter + Messung fuer den Body (Breite/Hoehe des Satzspiegels in pt)."""

    def __init__(self, base_pt: float, width: float, height: float):
        self.base_pt, self.width, self.height = base_pt, width, height
        self.measured = 0
        self.notes: set[str] = set()

    def fits(self, fragment: str, fraction: float) -> bool:
        if self.measured >= _MAX_MEASURED_ROWS:
            return True
        self.measured += 1
        story = pymupdf.Story(html=fragment, user_css=_css(self.base_pt))
        more, _ = story.place(pymupdf.Rect(0, 0, self.width, self.height * fraction))
        return not more


def _cells(row, base_pt: float) -> list[str]:
    cells, seen = [], set()
    for cell in row.cells:
        if id(cell._tc) in seen:  # noqa: SLF001 - verbundene Zellen nur einmal
            continue
        seen.add(id(cell._tc))  # noqa: SLF001
        cells.append("".join(_paragraph_html(p, base_pt) for p in cell.paragraphs))
    return cells


def _table_html(table, layout: _Layout) -> str:
    cls = "grid" if _table_has_borders(table) else "plain"
    out, rows = [], []

    def flush():
        if rows:
            out.append(f'<table class="{cls}">' + "".join(rows) + "</table>")
            rows.clear()

    for row in table.rows:
        cells = _cells(row, layout.base_pt)
        row_html = "<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>"
        fraction = _SPLIT_SINGLE_CELL if len(cells) == 1 else _SPLIT_MULTI_CELL
        if layout.fits(f'<table class="{cls}">{row_html}</table>', fraction):
            rows.append(row_html)
            continue
        # zu hoch fuer eine unteilbare Zeile: als Bloecke setzen (verlustfrei ueber Seiten teilbar)
        flush()
        out.extend(f'<div class="cellbox {cls}">{c}</div>' for c in cells)
        if len(cells) > 1:
            layout.notes.add("Sehr hohe Tabellenzeilen werden in der Seitenansicht untereinander statt nebeneinander "
                             "dargestellt, damit kein Text abgeschnitten wird.")
    flush()
    return "".join(out)


def _page_break_before(child) -> bool:
    ppr = child.find(f"{_W_NS}pPr")
    flag = ppr.find(f"{_W_NS}pageBreakBefore") if ppr is not None else None
    return flag is not None and (flag.get(f"{_W_NS}val") or "true") not in ("0", "false")


def _page_break_after(child) -> bool:
    breaks = child.findall(f".//{_W_NS}br")
    if any(b.get(f"{_W_NS}type") == "page" for b in breaks):
        return True
    ppr = child.find(f"{_W_NS}pPr")
    return ppr is not None and ppr.find(f"{_W_NS}sectPr") is not None  # Abschnittswechsel (neue Seite)


def _body_segments(document, layout: _Layout) -> list[str]:
    """HTML je Seitenfolge: ein manueller Seitenumbruch beginnt ein neues Segment (= neue Seite)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    segments, parts = [], []
    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            if _page_break_before(child) and parts:
                segments.append("".join(parts)); parts = []
            parts.append(_paragraph_html(Paragraph(child, document), layout.base_pt))
            if _page_break_after(child):
                segments.append("".join(parts)); parts = []
        elif tag == "tbl":
            parts.append(_table_html(Table(child, document), layout))
    if parts or not segments:
        segments.append("".join(parts))
    return segments


def _part_text(part) -> str:
    if part is None or getattr(part, "is_linked_to_previous", False) and not part.paragraphs:
        return ""
    return "\n".join(p.text for p in part.paragraphs if p.text.strip())


def _css(base_pt: float) -> str:
    return (f"body {{margin: 0; font-family: sans-serif; font-size: {base_pt:.1f}pt; line-height: 1.25;}}"
            "p {margin: 0 0 6pt 0;} p.li {margin-left: 14pt;}"
            "h1 {font-size: 1.6em; margin: 6pt 0 8pt 0;} h2 {font-size: 1.35em; margin: 6pt 0 6pt 0;}"
            "h3 {font-size: 1.15em; margin: 4pt 0 4pt 0;} h4 {font-size: 1.05em; margin: 4pt 0 4pt 0;}"
            "table {border-collapse: collapse; margin: 4pt 0 8pt 0; width: 100%;}"
            "td {padding: 3pt 5pt; vertical-align: top;} td p {margin: 0;}"
            "table.grid td {border: 0.75pt solid #000;}"
            "div.cellbox {padding: 3pt 5pt; margin: 0 0 4pt 0;} div.cellbox p {margin: 0 0 3pt 0;}"
            "div.cellbox.grid {border: 0.75pt solid #000;}")


def _render(path: Path) -> tuple[bytes, tuple[str, ...]]:
    import docx

    document = docx.Document(str(path))
    section = document.sections[0]
    width = section.page_width.pt if section.page_width else 595.0
    height = section.page_height.pt if section.page_height else 842.0
    left = section.left_margin.pt if section.left_margin is not None else 72.0
    right = section.right_margin.pt if section.right_margin is not None else 72.0
    top = section.top_margin.pt if section.top_margin is not None else 72.0
    bottom = section.bottom_margin.pt if section.bottom_margin is not None else 72.0
    header_y = section.header_distance.pt if section.header_distance is not None else 36.0
    footer_y = section.footer_distance.pt if section.footer_distance is not None else 36.0
    normal = document.styles["Normal"].font.size if "Normal" in [s.name for s in document.styles] else None
    base_pt = normal.pt if normal is not None else _DEFAULT_FONT_PT

    mediabox = pymupdf.Rect(0, 0, width, height)
    where = pymupdf.Rect(left, top, width - right, height - bottom)
    layout = _Layout(base_pt, where.width, where.height)
    writer_buffer = __import__("io").BytesIO()
    writer = pymupdf.DocumentWriter(writer_buffer)
    for segment in _body_segments(document, layout):
        story = pymupdf.Story(html=segment or "<p>&nbsp;</p>", user_css=_css(base_pt))
        more = True
        while more:
            device = writer.begin_page(mediabox)
            more, _ = story.place(where)
            story.draw(device)
            writer.end_page()
    writer.close()
    buffer = pymupdf.open("pdf", writer_buffer.getvalue())

    header = _part_text(section.header)
    footer = _part_text(section.footer)
    for page in buffer:
        if header:
            page.insert_htmlbox(pymupdf.Rect(left, header_y, width - right, top - 4),
                                html.escape(header).replace("\n", "<br>"),
                                css=f"* {{font-family: sans-serif; font-size: {max(base_pt - 1.5, 7):.1f}pt;}}")
        if footer:
            page.insert_htmlbox(pymupdf.Rect(left, height - bottom + 4, width - right, height - footer_y + 12),
                                html.escape(footer).replace("\n", "<br>"),
                                css=f"* {{font-family: sans-serif; font-size: {max(base_pt - 1.5, 7):.1f}pt;}}")
    data = buffer.tobytes()
    buffer.close()
    return data, tuple(sorted(layout.notes))


@lru_cache(maxsize=8)
def _cached(path_str: str, mtime_ns: int, size: int) -> tuple[bytes, tuple[str, ...]]:  # noqa: ARG001 - Cache-Schluessel
    return _render(Path(path_str))


def _cached_for(path: Path) -> tuple[bytes, tuple[str, ...]]:
    stat = path.stat()
    return _cached(str(path.resolve()), stat.st_mtime_ns, stat.st_size)


def layout_notes(path: Path) -> list[str]:
    """Hinweise zum Seitensatz dieses DOCX (z. B. untereinander gesetzte hohe Tabellenzeilen); leer bei Fehlern."""
    try:
        return list(_cached_for(path)[1])
    except Exception:  # noqa: BLE001 - nur Hinweistext
        return []


def unrendered_image_count(path: Path) -> int:
    """Anzahl eingebetteter Bilder/Grafiken, die der lokale Satz NICHT darstellt (fuer einen ehrlichen Hinweis)."""
    import docx

    try:
        document = docx.Document(str(path))
    except Exception:  # noqa: BLE001
        return 0
    return len(document.element.body.findall(".//" + _W_NS + "drawing"))


def docx_pdf_bytes(path: Path) -> bytes:
    """PDF-Bytes des lokalen DOCX-Seitensatzes (zwischengespeichert je Datei, Aenderungszeit und Groesse)."""
    return _cached_for(path)[0]


def open_rendered(path: Path):
    """Oeffnet ein Dokument fuer Seitenanzeige/-suche: DOCX ueber den lokalen Satz, sonst direkt. Faellt bei einem
    Fehler des Satzes auf den eingebauten DOCX-Import zurueck (Anzeige bleibt moeglich)."""
    if path.suffix.lower() == ".docx":
        try:
            return pymupdf.open("pdf", docx_pdf_bytes(path))
        except Exception:  # noqa: BLE001
            pass
    return pymupdf.open(path)
