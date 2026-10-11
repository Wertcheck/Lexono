"""Realistisch formatierte, SYNTHETISCHE Kanzleidokumente mit Goldstandard (12.10.2026).

Derselbe Inhalt als DOCX (Kopf-/Fusszeile, Ueberschrift, Aufzaehlung, zwei Tabellen, Signatur), als PDF (Fliesstext mit
echtem Zeilenumbruch, Tabellenspalten als positionierter Text, Kopf-/Fusszeile je Seite) und als Scan (PDF-Seiten als Bild,
nur ueber OCR lesbar). Alle Namen/Adressen/Kontonummern sind erfunden.

Aufruf:  python scripts/make_realistic_privacy_docs.py <Ausgabeverzeichnis>
Erzeugt: protokoll.docx, protokoll.pdf, protokoll_scan.pdf, gold.json
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

HEADER = "Kanzlei Beispiel & Partner Rechtsanwälte · Hafenkai 3 · 24103 Kiel"
FOOTER = "Bankverbindung: IBAN DE02 1203 0000 0000 2020 51 · Seite {page}"
TITLE = "Zeugenaussage Richter"
PARAGRAPHS = [
    "In der Sache Beispielmann ./. Musterbau GmbH bestätigen wir den Eingang des Schreibens von Richter vom 16.03.2024.",
    "Die Unterlagen wurden am 12.03.2026 durch Lange geprüft und dem Auftraggeber Herrn Marquardt übergeben; die Abnahme "
    "verweigerte Frau Wessolowski unter Hinweis auf erhebliche Mängel an der Dämmung der Dachfläche.",
]
LIST_INTRO = "Als Zeugen werden benannt:"
LIST_ITEMS = ["Herr Marquardt", "Fuerst", "Dr. Henrike Stuckenschneider"]
PARTY_TABLE = [
    ["Rolle", "Name", "Anschrift"],
    ["Klägerin", "Beispielmann", "Lindenallee 5, 24103 Kiel"],
    ["Zeuge", "Lange", "Hafenstraße 9, 24105 Kiel"],
    ["Gutachter", "Dr. Pfannenstiel", "Birkenweg 2, 24106 Kiel"],
]
INVOICE_TABLE = [
    ["Position", "Menge", "Betrag"],
    ["Estrich", "3", "1.200,00 EUR"],
    ["Dachrinne", "2", "450,00 EUR"],
    ["Summe", "", "1.650,00 EUR"],
]
CLOSING = [
    "Der Streitwert beträgt 12.400,00 EUR. Wir verlangen Nacherfüllung und behalten uns Minderung sowie Schadensersatz "
    "vor (§ 634 BGB). Die Zugluft im Flachdach ist ein Baumangel.",
    "Mit freundlichen Grüßen",
    "Rechtsanwältin Dr. Svenja Falk",
]

GOLD = {
    # Darf NIE im Klartext im Cloud-Payload stehen (Personen, Anschriften, Kontodaten)
    "personal": ["Beispielmann", "Richter", "Lange", "Marquardt", "Wessolowski", "Fuerst", "Stuckenschneider", "Henrike",
                 "Pfannenstiel", "Svenja Falk", "Lindenallee", "Hafenstraße", "Birkenweg", "DE02 1203"],
    # Soll lesbar bleiben (Fachbegriffe, Betraege, Normen)
    "keep": ["12.400,00", "1.200,00", "450,00", "1.650,00", "Streitwert", "Nacherfüllung", "Minderung", "Schadensersatz",
             "Zugluft", "Flachdach", "Baumangel", "Estrich", "Dachrinne", "Dämmung", "Mängel", "§ 634 BGB"],
}


def make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = HEADER
    section.footer.paragraphs[0].text = FOOTER.format(page=1)
    doc.add_heading(TITLE, level=1)
    for paragraph in PARAGRAPHS:
        doc.add_paragraph(paragraph)
    doc.add_paragraph(LIST_INTRO)
    for item in LIST_ITEMS:
        doc.add_paragraph(item, style="List Bullet")
    for rows in (PARTY_TABLE, INVOICE_TABLE):
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        doc.add_paragraph("")
    for line in CLOSING:
        doc.add_paragraph(line)
    doc.save(str(path))


def make_pdf(path: Path) -> None:
    import fitz

    doc = fitz.open()
    page_no = 0

    def new_page():
        nonlocal page_no
        page_no += 1
        page = doc.new_page(width=595, height=842)
        page.insert_text((60, 40), HEADER, fontsize=8)
        page.insert_text((60, 815), FOOTER.format(page=page_no), fontsize=8)
        return page

    page = new_page()
    y = 80
    page.insert_text((60, y), TITLE, fontsize=14)
    y += 26

    def block(text, size=11, indent=0):
        nonlocal y, page
        rect = fitz.Rect(60 + indent, y, 535, y + 200)
        rest = page.insert_textbox(rect, text, fontsize=size)
        used = 200 - rest if rest >= 0 else 200
        y += used + 8
        if y > 760:
            page = new_page()
            y = 80

    for paragraph in PARAGRAPHS:
        block(paragraph)
    block(LIST_INTRO)
    for item in LIST_ITEMS:
        block("- " + item, indent=12)
    for rows in (PARTY_TABLE, INVOICE_TABLE):
        y += 6
        for row in rows:
            for c, value in enumerate(row):
                page.insert_text((60 + c * 150, y), value, fontsize=10)
            y += 16
        y += 10
    for line in CLOSING:
        block(line)
    doc.save(str(path))


def make_scan(pdf_path: Path, scan_path: Path, dpi: int = 200) -> None:
    import fitz

    src = fitz.open(str(pdf_path))
    out = fitz.open()
    for page in src:
        pix = page.get_pixmap(dpi=dpi)
        target = out.new_page(width=page.rect.width, height=page.rect.height)
        target.insert_image(target.rect, pixmap=pix)
    out.save(str(scan_path))


def main() -> None:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "realistic_docs")
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "protokoll.docx")
    make_pdf(out / "protokoll.pdf")
    make_scan(out / "protokoll.pdf", out / "protokoll_scan.pdf")
    (out / "gold.json").write_text(json.dumps(GOLD, ensure_ascii=False, indent=1), encoding="utf-8")
    print("erzeugt in", out.resolve())


if __name__ == "__main__":
    main()
