"""E2E-Runner des Realistic Legal Benchmark (benchmarks/legal_e2e_v1) ueber echte Lexono-Pfade - ohne Cloud-Aufruf.

Je Fall, alle Faelle in EINER Datenbank (Mandantentrennung wird so messbar):
  1. Import:      Dateien rendern (TXT/DOCX/PDF/Scan) -> IntakeService.ingest_file -> Akte zuordnen
  2. Verarbeitung: DocumentProcessingService.process_document (Extraktion, OCR, Klassifikation, Fristen)
  3. Suche:       DocumentSearchService.index_document + search_within_matter
  4. Chat:        ChatService.send_message; der Schreib-Provider ist ein Fake, der den Payload UNMITTELBAR vor dem
                  Cloud-Aufruf erfasst (nach Pseudonymisierung, Rueckfrage, Restrisiko-Scan, Siegel) - es wird nichts gesendet
  5. Schriftsatz: DraftingService.create_draft (Fake-Antwort aus Platzhaltern des Payloads -> lokale Rueckfuehrung)
  6. Bearbeitung: EditorService.autosave_draft
  7. Speicherung/Export: Draft-Datensatz, DraftDocxExportService/DraftPdfExportService

Messungen werden getrennt ausgewiesen; kritische Datenschutzbefunde (Klartext im Payload, Daten anderer Akten,
Platzhalterreste im Export) stehen einzeln im Abschnitt "kritisch" und gehen in keinen Durchschnitt ein.

Aufruf:  python -m scripts.legal_bench.e2e [--split dev|test] [--out PFAD]
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path

os.environ["HF_HUB_OFFLINE"] = "1"  # nie ein Modell aus dem Netz nachladen (auch nicht bei abweichender Umgebung)

from scripts.legal_bench.corpus import CORPUS_DIR, ROOT, doc_text, load_cases, validate_corpus  # noqa: E402
from scripts.legal_bench.render import render_case  # noqa: E402

PARTY_ROLE_LABEL = {"gegner": "Gegner", "zeuge": "Zeuge", "dritter": "Beteiligter", "anwalt_gegner": "Rechtsanwalt Gegenseite",
                    "gericht": "Gericht", "kanzlei": "Kanzlei"}
_PLACEHOLDER = re.compile(r"\[[A-ZÄÖÜ_]+_\d+\]")
_FINGERPRINT_FILES = ("app/privacy/gateway.py", "app/privacy/detectors.py", "app/privacy/pseudonymizer.py",
                      "app/documents/service.py", "app/documents/extraction.py", "app/deadlines/extractor.py",
                      "app/search/service.py", "app/chat/service.py", "app/drafting/service.py",
                      "app/drafting/editor_service.py", "app/export/docx_export_service.py", "app/export/pdf_export_service.py")


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().lower()


def _contains(value: str, text: str) -> bool:
    return _norm(value) in _norm(text)


_MONTHS = ("Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November",
           "Dezember")


def leak_variants(entry: dict) -> list[str]:
    """Formen, in denen ein geschuetzter Wert im Payload auftauchen kann (Review §34: nur der volle Wert wurde
    geprueft). Personen: zusaetzlich der Nachname ("Herr Wiegelmann"); Organisationen: das erste Wort; persoenliche
    Daten: Kurzjahr, ohne fuehrende Nullen, Monatsname."""
    value, kind = entry["value"], entry["type"]
    variants = [value]
    words = value.split()
    if kind == "PERSON" and len(words) >= 2 and len(words[-1]) >= 4:
        variants.append(words[-1])
    if kind == "ORGANISATION" and len(words) >= 2 and len(words[0]) >= 5:
        variants.append(words[0])
    match = re.fullmatch(r"(\d{2})\.(\d{2})\.(\d{4})", value)
    if kind == "PERS_DATUM" and match:
        day, month, year = match.groups()
        variants += [f"{day}.{month}.{year[2:]}", f"{int(day)}.{int(month)}.{year}", f"{int(day)}.{int(month)}.{year[2:]}",
                     f"{day}. {month}. {year}", f"{int(day)}. {_MONTHS[int(month) - 1]} {year}"]
    return variants


def _leaked(variant: str, text: str) -> bool:
    """Wortgrenzen-genau (ein Nachname soll nicht in einem laengeren Wort anschlagen; Genitiv-s zaehlt mit), Leerraum
    normalisiert."""
    pattern = r"(?<![\wÄÖÜäöüß])" + r"\s+".join(re.escape(part) for part in variant.split()) + r"s?(?![\wÄÖÜäöüß])"
    return re.search(pattern, text) is not None


def _payload_text(payload) -> str:  # noqa: ANN001
    parts: list[str] = []

    def walk(value) -> None:  # noqa: ANN001
        if isinstance(value, str):
            parts.append(value)
        elif isinstance(value, list):
            for item in value:
                walk(item)
        elif isinstance(value, dict):
            for item in value.values():
                walk(item)

    walk(payload.model_dump())
    return "\n".join(parts)


class CapturingWriter:
    """Erfasst den Payload vor der Cloud-Grenze. Antwort: optional ein Auszug des pseudonymisierten Sachverhalts (damit
    Rueckfuehrung und Export echte Platzhalter verarbeiten); es wird nichts gesendet."""

    def __init__(self, echo: bool = False) -> None:
        self.echo = echo
        self.payloads: list = []

    def write(self, payload):  # noqa: ANN001
        from app.ai_providers.claude_writing_provider import ClaudeWritingResult

        self.payloads.append(payload)
        if self.echo:
            body = payload.anonymisierter_sachverhalt[:1200]
            return ClaudeWritingResult(text=f"Sehr geehrte Damen und Herren,\n\n{body}\n\nMit freundlichen Gruessen", token_count=0)
        return ClaudeWritingResult(text="Antwort.", token_count=0)


def _tesseract() -> str | None:
    """Wie im Installer-Bundle: vendored Tesseract samt explizitem TESSDATA_PREFIX (sonst findet tesseract.exe im
    Dev-Betrieb keine Sprachdaten); System-Installation nur als Rueckfall."""
    vendor = ROOT / "windows" / "vendor" / "tesseract"
    if (vendor / "bin" / "tesseract.exe").exists() and (vendor / "tessdata").is_dir():
        os.environ.setdefault("TESSDATA_PREFIX", str(vendor / "tessdata"))
        return str(vendor / "bin" / "tesseract.exe")
    return shutil.which("tesseract")


def _embedding_provider():  # noqa: ANN202
    """Produktionsmodell (FastEmbed, lokal im Cache, offline erzwungen); sonst Test-Embedding mit Kennzeichnung."""
    try:
        from app.config import get_settings
        from app.search.embeddings import FastEmbedProvider

        provider = FastEmbedProvider(model_name=get_settings().embedding_model_name)
        provider.embed("Probe")
        return provider, "fastembed:" + get_settings().embedding_model_name
    except Exception as exc:  # noqa: BLE001
        sys.path.insert(0, str(ROOT / "tests"))
        from fake_embedding_provider import FakeEmbeddingProvider

        return FakeEmbeddingProvider(), f"fake (Produktionsmodell nicht verfuegbar: {type(exc).__name__}) - Retrieval-Qualitaet NICHT VERIFIZIERT"


def run(cases: list[dict], workdir: Path, embedding=None) -> dict:  # noqa: ANN001
    """`embedding`: (Provider, Bezeichnung) - nur fuer Tests; Standard ist das Produktionsmodell (siehe _embedding_provider)."""
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool

    from app.ai_providers.local_ai_provider import RuleBasedLocalAIProvider
    from app.attorney_instructions.service import AttorneyInstructionService
    from app.chat.service import ChatService
    from app.documents.service import DocumentProcessingService
    from app.drafting.editor_service import EditorService
    from app.drafting.service import DraftingService
    from app.export.docx_export_service import DraftDocxExportService
    from app.export.pdf_export_service import DraftPdfExportService
    from app.ingestion.intake import IntakeService
    from app.models import Client, Deadline, Draft, Matter, Party, Role, User
    from app.models.base import Base
    from app.privacy.gateway import ClaudePrivacyGateway
    from app.research.service import LegalResearchService
    from app.search.service import DocumentSearchService

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    role = Role(name="anwalt")
    db.add(role)
    db.flush()
    user = User(email="benchmark@kanzlei.test", password_hash="x", role_id=role.id)
    db.add(user)
    db.commit()

    tesseract = _tesseract()
    processing = DocumentProcessingService(ocr_enabled=tesseract is not None, tesseract_cmd=tesseract)
    intake = IntakeService(workdir / "intake")
    provider, embedding_label = embedding or _embedding_provider()
    search = DocumentSearchService(provider)
    gateway = ClaudePrivacyGateway()
    all_protect = {c["id"]: sorted({v for p in c["reference"]["privacy"]["protect"] for v in leak_variants(p)}) for c in cases}

    def drafting(writer: CapturingWriter) -> DraftingService:
        return DraftingService(RuleBasedLocalAIProvider(search), LegalResearchService(search, min_score_for_sufficient=0.0),
                               search, gateway, writer, model_name="benchmark-offline")

    # --- 1./2. Import und Verarbeitung (alle Faelle zuerst: Mandantentrennung ueber eine gemeinsame DB) -----------------
    state: dict[str, dict] = {}
    for case in cases:
        t0 = time.perf_counter()
        parties = {p["id"]: p for p in case["parties"]}
        client = Client(name=parties[case["matter"]["client_party"]]["name"])
        matter = Matter(client=client, title=case["matter"]["title"])
        db.add_all([client, matter])
        db.flush()
        for party in case["parties"]:
            if party["id"] != case["matter"]["client_party"]:
                db.add(Party(matter_id=matter.id, name=party["name"], role=PARTY_ROLE_LABEL.get(party["role"], party["role"])))
        db.commit()
        files = render_case(case, workdir / "rendered")
        documents = {}
        doc_results = {}
        for spec in case["documents"]:
            start = time.perf_counter()
            entry: dict = {"format": spec["format"], "type_reference": spec["type"]}
            try:
                document = intake.ingest_file(files[spec["id"]], db, stability_timeout_seconds=5.0)
                document.matter_id = matter.id
                db.commit()
                processing.process_document(document, db, actor="benchmark")
                db.refresh(document)
                documents[spec["id"]] = document
                text = document.extracted_text or ""
                values = _reference_values_for(case, spec["id"])
                entry.update(status=document.ocr_status, classified=document.classified_type, chars=len(text),
                             values_total=len(values), values_found=[v for v in values if _contains(v, text)],
                             values_missing=[v for v in values if not _contains(v, text)])
                search.index_document(document, db)
            except Exception as exc:  # noqa: BLE001 - Fehler werden gemessen
                entry.update(status="error", error=type(exc).__name__)
            entry["seconds"] = round(time.perf_counter() - start, 2)
            doc_results[spec["id"]] = entry
        db.commit()
        state[case["id"]] = {"case": case, "matter": matter, "documents": documents, "doc_results": doc_results,
                             "setup_seconds": round(time.perf_counter() - t0, 2)}

    results = []
    for case in cases:
        st = state[case["id"]]
        matter, documents = st["matter"], st["documents"]
        foreign = [v for cid, values in all_protect.items() if cid != case["id"] for v in values]
        own_protect = all_protect[case["id"]]
        record: dict = {"case": case["id"], "split": case["split"], "documents": st["doc_results"], "tasks": [],
                        "critical": []}

        # Fristen (Verarbeitung)
        found = sorted({d.due_date.strftime("%d.%m.%Y") for d in db.query(Deadline).filter(Deadline.matter_id == matter.id) if d.due_date})
        expected = sorted({d["date"] for d in case["reference"]["deadlines"]})
        personal = [p["value"] for p in case["reference"]["privacy"]["protect"] if p["type"] == "PERS_DATUM"]
        # bekannte Daten des Falls: Chronologie, Dokumentdaten, datierte Widerspruchswerte
        events = {e["date"] for e in case["reference"]["chronology"]} | {d["date"] for d in case["documents"]} | {
            v["value"] for w in case["reference"]["contradictions"] for v in w["values"]
            if re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", v["value"])}
        extra = sorted(set(found) - set(expected))
        # Zusaetzliche Vorschlaege getrennt nach Ursache: bekannte Ereignisdaten (Chronologie) sind Fehlalarme der
        # Fristerkennung, aber keine erfundenen Daten; alles andere ist ungeklaert und wird gesondert ausgewiesen.
        record["deadlines"] = {"expected": expected, "found": found,
                               "recall": f"{len(set(expected) & set(found))}/{len(expected)}",
                               "missing": sorted(set(expected) - set(found)),
                               "not_in_reference": extra,
                               "known_dates_as_deadline": [d for d in extra if d in events],
                               "unexplained": [d for d in extra if d not in events and d not in personal],
                               "personal_dates_as_deadline": sorted(set(found) & set(personal))}
        if record["deadlines"]["personal_dates_as_deadline"]:
            record["critical"].append({"kind": "persoenliches_datum_als_frist", "values": record["deadlines"]["personal_dates_as_deadline"]})

        for task in case["reference"]["tasks"]:
            t0 = time.perf_counter()
            item: dict = {"id": task["id"], "kind": task["kind"]}
            try:
                if task["kind"] == "suche":
                    hits = search.search_within_matter(matter.id, task["question"], db, limit=5)
                    ids = [h.entity_id for h in hits]
                    # ein nicht importiertes erwartetes Dokument zaehlt als Fehltreffer (Review §34), nicht als "entfaellt"
                    expected_ids = [documents[d].id if d in documents else None for d in task.get("expected_documents", [])]
                    item.update(rank=[ids.index(e) + 1 if e in ids else None for e in expected_ids],
                                foreign_hits=sum(1 for h in hits if h.matter_id not in (None, matter.id)))
                    if item["foreign_hits"]:
                        record["critical"].append({"kind": "suchtreffer_fremder_akte", "task": task["id"]})
                elif task["kind"] == "fristen":
                    item.update(expected=task["expected_deadlines"], found=[d for d in task["expected_deadlines"] if d in found])
                elif task["kind"] == "chat":
                    item.update(_run_chat(db, user, matter, documents, task, drafting))
                elif task["kind"] == "schriftsatz":
                    item.update(_run_draft(db, user, matter, task, drafting, case, Draft, EditorService,
                                           AttorneyInstructionService, DraftDocxExportService, DraftPdfExportService))
            except Exception as exc:  # noqa: BLE001
                item.update(status="error", error=f"{type(exc).__name__}: {str(exc)[:120]}")
            sent = item.pop("_payload_text", None)
            if sent is not None:
                leaks = [v for v in own_protect if _leaked(v, sent)]
                cross = [v for v in foreign if _leaked(v, sent) and v not in own_protect]
                item.update(payload_chars=len(sent), pii_leaks=leaks, cross_matter_values=cross,
                            needed_present=[v for v in task.get("needed_values", []) if _contains(v, sent)],
                            needed_missing=[v for v in task.get("needed_values", []) if not _contains(v, sent)])
                if leaks:
                    record["critical"].append({"kind": "klartext_im_payload", "task": task["id"], "values": leaks})
                if cross:
                    record["critical"].append({"kind": "daten_anderer_akte_im_payload", "task": task["id"], "values": cross})
            for leftover in item.get("export_placeholders", []):
                record["critical"].append({"kind": "platzhalterrest_im_export", "task": task["id"], "value": leftover})
            item["seconds"] = round(time.perf_counter() - t0, 2)
            record["tasks"].append(item)
        results.append(record)

    db.close()
    engine.dispose()
    return {"results": results, "environment": {"embedding": embedding_label, "ocr": tesseract or "nicht verfuegbar",
                                                 "writer": "CapturingWriter (offline, kein Cloud-Aufruf)"}}


def _reference_values_for(case: dict, doc_id: str) -> list[str]:
    """Alle Referenzwerte, die laut Korpus im Dokument stehen (Grundlage fuer Extraktions-/OCR-Erhalt)."""
    text = doc_text(next(d for d in case["documents"] if d["id"] == doc_id))
    ref = case["reference"]
    candidates = [v for f in ref["facts"] for v in f["values"]] + [c["date"] for c in ref["chronology"]] + \
        [d["date"] for d in ref["deadlines"]] + [c["text"] for c in ref["citations"]] + \
        [p["value"] for p in ref["privacy"]["protect"]] + list(ref["privacy"]["keep"])
    return sorted({v for v in candidates if v in text})


def _run_chat(db, user, matter, documents, task, drafting) -> dict:  # noqa: ANN001
    from app.chat.service import ChatService

    chat = ChatService(Path(tempfile.mkdtemp()) / "chat")
    conversation = chat.create_conversation(db, user=user, matter_id=matter.id, title=task["id"], actor=user.email)
    message = chat.record_user_message(db, conversation=conversation, content=task["question"],
                                       document_ids=[documents[d].id for d in task.get("attach", []) if d in documents])
    writer = CapturingWriter()
    reply = chat.send_message(db, conversation=conversation, content=task["question"], drafting_service=drafting(writer),
                              actor=user.email, current_message_id=message.id)
    result = {"status": "payload" if writer.payloads else ("blockiert" if reply.blocked else "ohne_cloud_aufruf"),
              "reply_excerpt": (reply.content or "")[:160]}
    if writer.payloads:
        result["_payload_text"] = _payload_text(writer.payloads[-1])
    return result


def _run_draft(db, user, matter, task, drafting, case, Draft, EditorService, AttorneyInstructionService,  # noqa: ANN001, N803
               DraftDocxExportService, DraftPdfExportService) -> dict:  # noqa: N803
    writer = CapturingWriter(echo=True)
    service = drafting(writer)
    outcome = service.create_draft(matter.id, "formulate_draft", db, attorney_anmerkungen=task["question"], actor=user.email)
    result: dict = {"status": "ok" if outcome.success else "blockiert", "blocked_reasons": list(outcome.blocked_reasons or [])[:3]}
    if writer.payloads:
        result["_payload_text"] = _payload_text(writer.payloads[-1])
    if not outcome.success or not outcome.draft_id:
        return result
    draft = db.get(Draft, outcome.draft_id)
    client_name = next(p["name"] for p in case["parties"] if p["id"] == case["matter"]["client_party"])
    result.update(stored=draft is not None, draft_placeholders=sorted(set(_PLACEHOLDER.findall(draft.content or ""))),
                  client_reconstructed=client_name in (draft.content or ""))
    # Bearbeitung
    editor = EditorService(AttorneyInstructionService(service))
    editor.autosave_draft(db, draft=draft, content=(draft.content or "") + "\n\nErgaenzung der Kanzlei: Frist notiert.",
                          subject=None, recipient=None, content_format="text", actor=user.email)
    db.refresh(draft)
    result["edit_saved"] = "Ergaenzung der Kanzlei" in (draft.content or "")
    # Export
    import docx
    import pymupdf

    docx_text = "\n".join(p.text for p in docx.Document(DraftDocxExportService().export_draft(draft, matter)).paragraphs)
    pdf_bytes = DraftPdfExportService().export_draft(draft, matter).getvalue()
    pdf_text = "".join(page.get_text() for page in pymupdf.open(stream=pdf_bytes, filetype="pdf"))
    result.update(export_docx_has_edit="Ergaenzung der Kanzlei" in docx_text, export_pdf_has_edit="Ergaenzung der Kanzlei" in pdf_text,
                  export_placeholders=sorted(set(_PLACEHOLDER.findall(docx_text + pdf_text))))
    return result


def summarise(results: list[dict]) -> dict:
    def per_split(split: str) -> dict:
        rows = [r for r in results if r["split"] == split]
        docs = [d for r in rows for d in r["documents"].values()]
        tasks = [t for r in rows for t in r["tasks"]]
        chats = [t for t in tasks if t["kind"] in ("chat", "schriftsatz") and "payload_chars" in t]
        searches = [t for t in tasks if t["kind"] == "suche"]
        drafts = [t for t in tasks if t["kind"] == "schriftsatz"]
        return {
            "cases": len(rows),
            "kritisch": [dict(c, case=r["case"]) for r in rows for c in r["critical"]],
            "extraktion_werte": f"{sum(len(d.get('values_found', [])) for d in docs)}/{sum(d.get('values_total', 0) for d in docs)}",
            "ocr_dokumente": sum(1 for d in docs if d["format"] == "scan"),
            "ocr_werte": f"{sum(len(d.get('values_found', [])) for d in docs if d['format'] == 'scan')}/"
                         f"{sum(d.get('values_total', 0) for d in docs if d['format'] == 'scan')}",
            "dokumentfehler": sum(1 for d in docs if d.get("status") in ("error", "failed", "pending")),
            # nur Datumsgleichheit - Art der Frist und Quelldokument werden NICHT geprueft
            "fristen_datum_recall": f"{sum(int(r['deadlines']['recall'].split('/')[0]) for r in rows)}/"
                              f"{sum(int(r['deadlines']['recall'].split('/')[1]) for r in rows)}",
            "fristen_nicht_in_referenz": sum(len(r["deadlines"]["not_in_reference"]) for r in rows),
            "davon_bekannte_falldaten": sum(len(r["deadlines"]["known_dates_as_deadline"]) for r in rows),
            "davon_ungeklaert": [f"{r['case']}:{d}" for r in rows for d in r["deadlines"]["unexplained"]],
            # nur ob ueberhaupt ein Typ vergeben wurde - die Korrektheit gegen den Referenztyp ist NICHT bewertet
            "klassifiziert_nicht_unbekannt": f"{sum(1 for d in docs if d.get('classified') not in (None, 'Unbekannt'))}/{len(docs)}",
            "payload_aufgaben": len(chats),
            "pii_leaks": sum(len(t["pii_leaks"]) for t in chats),
            "fremde_akte_im_payload": sum(len(t["cross_matter_values"]) for t in chats),
            "benoetigte_werte_im_payload": f"{sum(len(t['needed_present']) for t in chats)}/"
                                           f"{sum(len(t['needed_present']) + len(t['needed_missing']) for t in chats)}",
            "aufgaben_ohne_payload": [f"{r['case']}/{t['id']}:{t.get('status')}" for r in rows for t in r["tasks"]
                                      if t["kind"] in ("chat", "schriftsatz") and "payload_chars" not in t],
            # top1: ein erwartetes Dokument steht auf Platz 1; top5: alle erwarteten Dokumente unter den ersten fuenf.
            "suche_top1": f"{sum(1 for t in searches if 1 in (t.get('rank') or []))}/{len(searches)}",
            "suche_top5": f"{sum(1 for t in searches if t.get('rank') and all(t['rank']))}/{len(searches)}",
            "schriftsatz_ok": f"{sum(1 for t in drafts if t.get('status') == 'ok')}/{len(drafts)}",
            "export_mit_bearbeitung": f"{sum(1 for t in drafts if t.get('export_docx_has_edit') and t.get('export_pdf_has_edit'))}/{len(drafts)}",
            "aufgabenfehler": [f"{r['case']}/{t['id']}" for r in rows for t in r["tasks"] if t.get("status") == "error"],
        }

    return {"dev": per_split("dev"), "test": per_split("test")}


def _git(*args: str) -> str:
    try:
        return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, check=True).stdout.rstrip("\n")
    except Exception:  # noqa: BLE001
        return "unbekannt"


def provenance() -> dict:
    dirty = [p for p in (line[3:].strip('"') for line in _git("status", "--porcelain").splitlines())
             if not p.startswith(("assets/", "benchmarks/legal_e2e_v1/results"))]
    corpus_hash = hashlib.sha256()
    for path in sorted(CORPUS_DIR.rglob("*.json")):
        if "results" not in path.name:
            corpus_hash.update(path.read_bytes())
    sha = lambda p: hashlib.sha256((ROOT / p).read_bytes()).hexdigest()[:16]  # noqa: E731
    return {"git_commit": _git("rev-parse", "--short", "HEAD"), "git_dirty_files": dirty,
            "corpus_sha256": corpus_hash.hexdigest()[:16], "code_fingerprint": {f: sha(f) for f in _FINGERPRINT_FILES},
            "python": platform.python_version()}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--split", choices=["dev", "test"], default=None)
    parser.add_argument("--out", default=None)
    args = parser.parse_args()
    report = validate_corpus()
    if not report.ok:
        print("Korpus ungueltig:", *report.errors, sep="\n  ")
        return 1
    cases = load_cases(split=args.split)
    workdir = Path(tempfile.mkdtemp(prefix="legal_e2e_"))
    started = time.perf_counter()
    try:
        outcome = run(cases, workdir)
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    output = {"benchmark": "legal_e2e_v1", "run_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
              **provenance(), "environment": outcome["environment"], "seconds_total": round(time.perf_counter() - started, 1),
              "kosten": "keine (kein Cloud-Aufruf); echte Modellkosten NICHT VERIFIZIERT",
              "nicht_verifiziert": ["Antwortqualitaet eines echten Modells", "Kosten", "installierter Build (eigener Lauf)"],
              "summary": summarise(outcome["results"]), "results": outcome["results"]}
    out = Path(args.out) if args.out else CORPUS_DIR / "results.json"
    out.write_text(json.dumps(output, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps(output["summary"], ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
